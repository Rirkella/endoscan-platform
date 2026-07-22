"""Reviewed public invitroDB archive staging and pre-approval activity coverage.

The multi-gigabyte public archive is deliberately kept outside workflow JSON and the
small-artifact store.  It is streamed once into an ignored, immutable provider cache.
Only documentation-backed members are extracted, normalized into a disk-backed SQLite
index, and queried for source coverage.  No activity threshold or derived label is
created here.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import sqlite3
import stat
import tempfile
import time
import zipfile
from collections.abc import Iterable, Iterator
from contextlib import closing
from datetime import date, datetime
from datetime import time as datetime_time
from decimal import Decimal
from enum import StrEnum
from pathlib import Path
from typing import Any, Protocol
from urllib.parse import urljoin, urlparse

import httpx
from openpyxl import load_workbook
from pydantic import Field, model_validator

from .discovery_strategy import ImmutableV2Contract, deterministic_fingerprint

TOXCAST_RELEASE = "invitrodb-v4.3-2025-08"
TOXCAST_SOURCE_VERSION = "invitrodb v4.3"
TOXCAST_ARCHIVE_FILE_ID = "68af6b70e4b02565fc7c3a98"
TOXCAST_ARCHIVE_LOCATOR = "https://clowder.edap-cluster.com/files/68af6b70e4b02565fc7c3a98/blob"
TOXCAST_ARCHIVE_EXPECTED_SIZE = 7_502_075_199
TOXCAST_ARCHIVE_SHA256 = "c61d46cf685d5d84e72d0d4f6f46746ab1714d6f924e43480d390925ae827c59"
TOXCAST_MINIMUM_WORKSPACE_BYTES = 15_004_150_398
TOXCAST_APPROVED_HOST = "clowder.edap-cluster.com"
TOXCAST_RELEASE_CITATION = "https://doi.org/10.23645/epacomptox.6062623.v14"
TOXCAST_ARCHIVE_MIME_TYPES = frozenset(
    {
        "application/zip",
        "application/octet-stream",
        "application/x-zip-compressed",
        # Clowder's reviewed multi-file archive endpoint uses this registered
        # source-specific response type for INVITRODB_SUMMARY.zip.
        "multi/files-zipped",
    }
)
ARCHIVE_CACHE_SCHEMA_VERSION = "1.0.0"
NORMALIZED_ACTIVITY_SCHEMA_VERSION = "1.0.0"
TOXCAST_RELEASE_DOCUMENTATION_LOCATORS = [
    "https://www.epa.gov/comptox-tools/exploring-toxcast-data",
    "https://www.epa.gov/system/files/documents/2024-02/toxcast_cop_25jan2024.pdf",
]
TOXCAST_REVIEWED_ROLE_DOCUMENTATION = """
The official EPA invitroDB v4.3 release publishes summary files, assay information,
analytical QC, and the tcpl processing documentation.  In the reviewed tcpl schema,
mc4 stores multi-concentration model fitting inputs and complete fitted-model fields;
mc5 stores winning-model and source hit-call outputs; mc6 stores caution flags.  sc1
stores single-concentration normalized response records and sc2 stores the corresponding
single-concentration summary and source hit call.  Analytical QC provides source chemical
DTXSID, sample identifier, name, and QC/status fields.  The complete tested-chemical
catalogue is reconciled from that QC reference plus every distinct chemical identifier in
the verified mc4, mc5-6, sc1, and sc2 activity members; no single filename proves it alone.
"""


def _canonical_json(value: Any) -> str:
    def encode_source_scalar(item: Any) -> str:
        if isinstance(item, datetime | date | datetime_time):
            return item.isoformat()
        if isinstance(item, Decimal):
            return format(item, "f")
        raise TypeError(f"unsupported source scalar: {type(item).__name__}")

    return json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
        default=encode_source_scalar,
    )


def _atomic_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}-", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(value, handle, indent=2, sort_keys=True, ensure_ascii=False)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        Path(temporary).replace(path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def _hash_file(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as handle:
        while chunk := handle.read(8 * 1024 * 1024):
            digest.update(chunk)
            size += len(chunk)
    return digest.hexdigest(), size


def _safe_relative(root: Path, relative: str) -> Path:
    target = (root / relative).resolve()
    if root.resolve() not in target.parents:
        raise ValueError("provider cache path escaped its reviewed root")
    return target


class ToxCastMemberRole(StrEnum):
    TESTED_CHEMICALS = "complete_tested_chemical_catalogue"
    MULTI_CONCENTRATION_MODEL_FITS = "multi_concentration_activity_model_fit"
    MULTI_CONCENTRATION_HIT_CALLS = "multi_concentration_hit_calls_and_flags"
    SINGLE_CONCENTRATION_ACTIVITY = "single_concentration_activity_hit_calls"
    SCHEMA_DOCUMENTATION = "schema_or_data_dictionary"


class CachedProviderFile(ImmutableV2Contract):
    cache_relative_path: str = Field(min_length=1, max_length=500)
    sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    size_bytes: int = Field(ge=0)
    immutable: bool = True


class ArchiveIdentity(ImmutableV2Contract):
    locator: str = Field(pattern=r"^https://[^\s]+$")
    final_locator: str = Field(pattern=r"^https://[^\s]+$")
    file_id: str
    release_id: str
    source_version: str
    expected_size_bytes: int = Field(ge=1)
    content_length_bytes: int = Field(ge=1)
    etag: str | None = None
    last_modified: str | None = None
    content_type: str | None = None
    release_citation: str = Field(pattern=r"^https://doi\.org/[^\s]+$")


class ArchiveTransportAttempt(ImmutableV2Contract):
    attempt_number: int = Field(ge=1, le=2)
    outcome: str
    bytes_received: int = Field(ge=0)
    duration_ms: int = Field(ge=0)
    http_status: int | None = Field(default=None, ge=100, le=599)
    final_allowlisted_host: str | None = None
    exception_class: str | None = None
    safe_message: str | None = Field(default=None, max_length=500)
    retryable: bool = False


class ToxCastArchiveStageResult(ImmutableV2Contract):
    identity: ArchiveIdentity
    archive: CachedProviderFile
    cache_manifest: CachedProviderFile
    cache_hit: bool
    network_request_count: int = Field(ge=0, le=3)
    downloaded_bytes: int = Field(ge=0)
    duration_ms: int = Field(ge=0)
    temporary_file_removed: bool
    transport_attempts: list[ArchiveTransportAttempt] = Field(default_factory=list, max_length=2)


class ArchiveProbe(ImmutableV2Contract):
    final_locator: str
    http_status: int
    content_length_bytes: int
    content_type: str | None = None
    etag: str | None = None
    last_modified: str | None = None
    final_host: str
    redirect_count: int = Field(ge=0, le=2)


class ArchiveDownloadTransport(Protocol):
    def probe(self, locator: str, *, expected_size: int) -> ArchiveProbe: ...

    def stream_download(
        self,
        locator: str,
        destination: Path,
        *,
        expected_size: int,
        expected_identity: ArchiveProbe,
        attempt_number: int,
    ) -> tuple[ArchiveTransportAttempt, str]: ...


class HttpxArchiveDownloadTransport:
    """Exact-host, streaming transport for the one reviewed Clowder archive."""

    def __init__(self, *, timeout_seconds: float = 1_200.0) -> None:
        self.client = httpx.Client(
            timeout=httpx.Timeout(
                connect=min(timeout_seconds, 60.0),
                read=min(timeout_seconds, 240.0),
                write=min(timeout_seconds, 60.0),
                pool=min(timeout_seconds, 60.0),
            ),
            follow_redirects=False,
            headers={"User-Agent": "EndoScan/1.0 reviewed-public-provider-cache"},
        )

    @staticmethod
    def _validate_url(locator: str) -> None:
        parsed = urlparse(locator)
        if (
            parsed.scheme != "https"
            or parsed.hostname != TOXCAST_APPROVED_HOST
            or parsed.port not in {None, 443}
            or parsed.username
            or parsed.password
            or parsed.fragment
        ):
            raise ValueError("ToxCast archive locator escaped its reviewed HTTPS host")

    def _stream(
        self, locator: str, *, headers: dict[str, str] | None = None
    ) -> tuple[Any, httpx.Response, str, int]:
        self._validate_url(locator)
        current = locator
        for redirects in range(3):
            context = self.client.stream("GET", current, headers=headers)
            response = context.__enter__()
            if not response.is_redirect:
                return context, response, current, redirects
            location = response.headers.get("location")
            context.__exit__(None, None, None)
            if not location:
                raise ValueError("reviewed archive redirect omitted Location")
            current = urljoin(current, location)
            self._validate_url(current)
        raise ValueError("reviewed archive exceeded its redirect limit")

    def probe(self, locator: str, *, expected_size: int) -> ArchiveProbe:
        context, response, final_locator, redirects = self._stream(
            locator, headers={"Range": "bytes=0-0"}
        )
        try:
            if response.status_code != 206:
                raise ValueError("reviewed archive ignored the bounded Range metadata probe")
            content_range = response.headers.get("content-range", "")
            match = re.fullmatch(r"bytes 0-0/(\d+)", content_range)
            if match is None or int(match.group(1)) != expected_size:
                raise ValueError("archive Content-Range did not prove the reviewed size")
            content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
            if content_type not in TOXCAST_ARCHIVE_MIME_TYPES:
                raise ValueError("archive metadata probe returned an unexpected MIME type")
            first_chunk = next(response.iter_raw(chunk_size=1), b"")
            if len(first_chunk) != 1:
                raise ValueError("bounded archive metadata probe returned an invalid byte count")
            return ArchiveProbe(
                final_locator=final_locator,
                http_status=response.status_code,
                content_length_bytes=expected_size,
                content_type=content_type,
                etag=response.headers.get("etag"),
                last_modified=response.headers.get("last-modified"),
                final_host=urlparse(final_locator).hostname or "",
                redirect_count=redirects,
            )
        finally:
            context.__exit__(None, None, None)

    def stream_download(
        self,
        locator: str,
        destination: Path,
        *,
        expected_size: int,
        expected_identity: ArchiveProbe,
        attempt_number: int,
    ) -> tuple[ArchiveTransportAttempt, str]:
        started = time.monotonic()
        received = 0
        digest = hashlib.sha256()
        context, response, final_locator, _redirects = self._stream(locator)
        try:
            if response.status_code != 200:
                raise ValueError("reviewed archive download did not return HTTP 200")
            final_host = urlparse(final_locator).hostname or ""
            if final_host != TOXCAST_APPROVED_HOST:
                raise ValueError("archive download ended on an unapproved host")
            declared = int(response.headers.get("content-length", "0") or 0)
            if declared and declared != expected_size:
                raise ValueError("archive Content-Length changed after approval")
            content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
            if content_type not in TOXCAST_ARCHIVE_MIME_TYPES:
                raise ValueError("archive download returned an unexpected MIME type")
            for key, expected in (
                ("etag", expected_identity.etag),
                ("last-modified", expected_identity.last_modified),
            ):
                observed = response.headers.get(key)
                if expected and observed and expected != observed:
                    raise ValueError(f"archive {key} changed during retrieval")
            with destination.open("wb") as handle:
                for chunk in response.iter_bytes(8 * 1024 * 1024):
                    if not chunk:
                        continue
                    received += len(chunk)
                    if received > expected_size:
                        raise ValueError("archive exceeded the approved expected size")
                    digest.update(chunk)
                    handle.write(chunk)
                handle.flush()
                os.fsync(handle.fileno())
            if received != expected_size:
                raise OSError("archive transfer ended before the approved size was reached")
            return (
                ArchiveTransportAttempt(
                    attempt_number=attempt_number,
                    outcome="completed",
                    bytes_received=received,
                    duration_ms=max(0, int((time.monotonic() - started) * 1000)),
                    http_status=response.status_code,
                    final_allowlisted_host=final_host,
                ),
                digest.hexdigest(),
            )
        except Exception as exc:
            return (
                ArchiveTransportAttempt(
                    attempt_number=attempt_number,
                    outcome="failed",
                    bytes_received=received,
                    duration_ms=max(0, int((time.monotonic() - started) * 1000)),
                    http_status=response.status_code if response is not None else None,
                    final_allowlisted_host=(
                        urlparse(final_locator).hostname if final_locator else None
                    ),
                    exception_class=type(exc).__name__,
                    safe_message=f"Reviewed archive transfer failed with {type(exc).__name__}.",
                    retryable=isinstance(exc, httpx.TransportError | OSError),
                ),
                "",
            )
        finally:
            context.__exit__(None, None, None)


class ToxCastArchiveCache:
    """One-time immutable staging for the reviewed public invitroDB archive."""

    def __init__(
        self,
        root: Path,
        transport: ArchiveDownloadTransport | None = None,
        *,
        expected_archive_sha256: str | None = None,
    ) -> None:
        self.root = Path(root).resolve()
        self.transport = transport
        self.expected_archive_sha256 = expected_archive_sha256
        self.release_root = self.root / "toxcast" / TOXCAST_RELEASE
        self.release_root.mkdir(parents=True, exist_ok=True)

    @property
    def stage_manifest_path(self) -> Path:
        return self.release_root / "archive-stage-manifest.json"

    def _cached_stage(self) -> ToxCastArchiveStageResult | None:
        if not self.stage_manifest_path.is_file():
            return None
        payload = json.loads(self.stage_manifest_path.read_text(encoding="utf-8"))
        stage = ToxCastArchiveStageResult.model_validate(payload)
        archive_path = _safe_relative(self.root, stage.archive.cache_relative_path)
        if not archive_path.is_file() or archive_path.stat().st_size != stage.archive.size_bytes:
            return None
        archive_hash, archive_size = _hash_file(archive_path)
        if archive_hash != stage.archive.sha256 or archive_size != stage.archive.size_bytes:
            return None
        if (
            self.expected_archive_sha256 is not None
            and archive_hash != self.expected_archive_sha256
        ):
            raise ValueError("cached ToxCast archive does not match its reviewed SHA-256")
        manifest_hash, manifest_size = _hash_file(self.stage_manifest_path)
        return stage.model_copy(
            update={
                "cache_manifest": stage.cache_manifest.model_copy(
                    update={"sha256": manifest_hash, "size_bytes": manifest_size}
                ),
                "cache_hit": True,
                "network_request_count": 0,
                "downloaded_bytes": 0,
                "duration_ms": 0,
                "transport_attempts": [],
            }
        )

    def _persist_completed_stage(
        self,
        *,
        identity: ArchiveIdentity,
        final: Path,
        digest: str,
        network_request_count: int,
        downloaded_bytes: int,
        duration_ms: int,
        transport_attempts: list[ArchiveTransportAttempt],
        temporary_file_removed: bool,
    ) -> ToxCastArchiveStageResult:
        archive_reference = CachedProviderFile(
            cache_relative_path=final.relative_to(self.root).as_posix(),
            sha256=digest,
            size_bytes=identity.expected_size_bytes,
        )
        stage_without_manifest = {
            "identity": identity.model_dump(mode="json"),
            "archive": archive_reference.model_dump(mode="json"),
            "cache_hit": False,
            "network_request_count": network_request_count,
            "downloaded_bytes": downloaded_bytes,
            "duration_ms": duration_ms,
            "temporary_file_removed": temporary_file_removed,
            "transport_attempts": [item.model_dump(mode="json") for item in transport_attempts],
        }
        placeholder = CachedProviderFile(
            cache_relative_path=self.stage_manifest_path.relative_to(self.root).as_posix(),
            sha256="0" * 64,
            size_bytes=0,
        )
        _atomic_json(
            self.stage_manifest_path,
            {**stage_without_manifest, "cache_manifest": placeholder.model_dump(mode="json")},
        )
        manifest_hash, manifest_size = _hash_file(self.stage_manifest_path)
        manifest_reference = placeholder.model_copy(
            update={"sha256": manifest_hash, "size_bytes": manifest_size}
        )
        _atomic_json(
            self.stage_manifest_path,
            {
                **stage_without_manifest,
                "cache_manifest": manifest_reference.model_dump(mode="json"),
            },
        )
        final_manifest_hash, final_manifest_size = _hash_file(self.stage_manifest_path)
        return ToxCastArchiveStageResult.model_validate(
            {
                **stage_without_manifest,
                "cache_manifest": manifest_reference.model_copy(
                    update={
                        "sha256": final_manifest_hash,
                        "size_bytes": final_manifest_size,
                    }
                ).model_dump(mode="json"),
            }
        )

    def _recover_complete_partial(
        self,
        *,
        locator: str,
        expected_size: int,
        temporary: Path,
    ) -> ToxCastArchiveStageResult | None:
        if not temporary.is_file():
            return None
        if temporary.stat().st_size != expected_size:
            temporary.unlink(missing_ok=True)
            return None
        if not zipfile.is_zipfile(temporary):
            temporary.unlink(missing_ok=True)
            raise ValueError("exact-size archive partial failed ZIP central-directory validation")
        partial_stat = temporary.stat()
        digest, observed_size = _hash_file(temporary)
        if observed_size != expected_size:
            raise ValueError("archive partial changed during local recovery validation")
        if self.expected_archive_sha256 is not None and digest != self.expected_archive_sha256:
            temporary.unlink(missing_ok=True)
            raise ValueError("exact-size archive partial failed reviewed SHA-256 validation")
        final = self.release_root / "INVITRODB_SUMMARY.zip"
        if final.exists():
            final.chmod(stat.S_IWRITE)
            final.unlink()
        temporary.replace(final)
        try:
            final.chmod(stat.S_IREAD)
        except OSError:
            pass
        duration_ms = max(0, int((partial_stat.st_mtime - partial_stat.st_ctime) * 1000))
        identity = ArchiveIdentity(
            locator=locator,
            final_locator=locator,
            file_id=TOXCAST_ARCHIVE_FILE_ID,
            release_id=TOXCAST_RELEASE,
            source_version=TOXCAST_SOURCE_VERSION,
            expected_size_bytes=expected_size,
            content_length_bytes=expected_size,
            content_type="multi/files-zipped",
            release_citation=TOXCAST_RELEASE_CITATION,
        )
        recovered_attempt = ArchiveTransportAttempt(
            attempt_number=1,
            outcome="completed_recovered_after_operator_timeout",
            bytes_received=expected_size,
            duration_ms=duration_ms,
            http_status=200,
            final_allowlisted_host=TOXCAST_APPROVED_HOST,
        )
        return self._persist_completed_stage(
            identity=identity,
            final=final,
            digest=digest,
            network_request_count=2,
            downloaded_bytes=expected_size,
            duration_ms=duration_ms,
            transport_attempts=[recovered_attempt],
            temporary_file_removed=not temporary.exists(),
        )

    def stage(
        self,
        *,
        locator: str = TOXCAST_ARCHIVE_LOCATOR,
        expected_size: int = TOXCAST_ARCHIVE_EXPECTED_SIZE,
        maximum_full_download_retries: int = 1,
    ) -> ToxCastArchiveStageResult:
        cached = self._cached_stage()
        if cached is not None:
            return cached
        if self.transport is None:
            raise RuntimeError(
                "ToxCast archive is not staged and no approved download transport was supplied"
            )
        if maximum_full_download_retries not in {0, 1}:
            raise ValueError("full archive retry policy must be zero or one")
        temporary = self.release_root / "INVITRODB_SUMMARY.zip.partial"
        recovered = self._recover_complete_partial(
            locator=locator,
            expected_size=expected_size,
            temporary=temporary,
        )
        if recovered is not None:
            return recovered
        free_bytes = shutil.disk_usage(self.release_root).free
        if free_bytes < TOXCAST_MINIMUM_WORKSPACE_BYTES:
            raise OSError(
                "Controlled provider cache lacks the approved 15,004,150,398-byte workspace."
            )
        started = time.monotonic()
        probe = self.transport.probe(locator, expected_size=expected_size)
        identity = ArchiveIdentity(
            locator=locator,
            final_locator=probe.final_locator,
            file_id=TOXCAST_ARCHIVE_FILE_ID,
            release_id=TOXCAST_RELEASE,
            source_version=TOXCAST_SOURCE_VERSION,
            expected_size_bytes=expected_size,
            content_length_bytes=probe.content_length_bytes,
            etag=probe.etag,
            last_modified=probe.last_modified,
            content_type=probe.content_type,
            release_citation=TOXCAST_RELEASE_CITATION,
        )
        attempts: list[ArchiveTransportAttempt] = []
        temporary.unlink(missing_ok=True)
        digest = ""
        for attempt_number in range(1, maximum_full_download_retries + 2):
            attempt, digest = self.transport.stream_download(
                locator,
                temporary,
                expected_size=expected_size,
                expected_identity=probe,
                attempt_number=attempt_number,
            )
            attempts.append(attempt)
            if attempt.outcome == "completed":
                break
            temporary.unlink(missing_ok=True)
            if not attempt.retryable:
                break
        if not attempts or attempts[-1].outcome != "completed" or not digest:
            temporary.unlink(missing_ok=True)
            raise OSError("Reviewed archive staging failed; incomplete temporary data was removed.")
        if self.expected_archive_sha256 is not None and digest != self.expected_archive_sha256:
            temporary.unlink(missing_ok=True)
            raise ValueError("downloaded ToxCast archive failed reviewed SHA-256 validation")
        final = self.release_root / "INVITRODB_SUMMARY.zip"
        if final.exists():
            final.chmod(stat.S_IWRITE)
            final.unlink()
        temporary.replace(final)
        try:
            final.chmod(stat.S_IREAD)
        except OSError:
            pass
        return self._persist_completed_stage(
            identity=identity,
            final=final,
            digest=digest,
            network_request_count=1 + len(attempts),
            downloaded_bytes=sum(item.bytes_received for item in attempts),
            duration_ms=max(0, int((time.monotonic() - started) * 1000)),
            transport_attempts=attempts,
            temporary_file_removed=not temporary.exists(),
        )


class ToxCastArchiveMember(ImmutableV2Contract):
    member_name: str
    compressed_size: int = Field(ge=0)
    uncompressed_size: int = Field(ge=0)
    compression_method: int = Field(ge=0)
    crc32: str = Field(pattern=r"^[a-f0-9]{8}$")
    local_header_offset: int = Field(ge=0)
    scientific_role: ToxCastMemberRole | None = None
    role_verification_status: str = "unverified"
    role_verification_evidence: list[str] = Field(default_factory=list, max_length=20)
    schema_columns: list[str] = Field(default_factory=list, max_length=500)
    extracted_file: CachedProviderFile | None = None


class ToxCastArchiveMemberManifest(ImmutableV2Contract):
    release_id: str = TOXCAST_RELEASE
    source_version: str = TOXCAST_SOURCE_VERSION
    archive_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    archive_size_bytes: int = Field(ge=1)
    members: list[ToxCastArchiveMember] = Field(max_length=2_000)
    exact_member_count: int = Field(ge=1)
    required_roles: list[ToxCastMemberRole]
    all_required_roles_verified: bool
    ambiguity_findings: list[str] = Field(default_factory=list, max_length=100)
    manifest_fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")

    @model_validator(mode="after")
    def fingerprint_matches(self) -> ToxCastArchiveMemberManifest:
        payload = self.model_dump(mode="json", exclude={"manifest_fingerprint"})
        if deterministic_fingerprint(payload) != self.manifest_fingerprint:
            raise ValueError("ToxCast member-manifest fingerprint does not match")
        return self


class ToxCastRoleRule(ImmutableV2Contract):
    role: ToxCastMemberRole
    filename_pattern: str
    required_header_groups: list[list[str]] = Field(min_length=1, max_length=20)
    documentation_terms: list[str] = Field(min_length=1, max_length=20)


DEFAULT_ROLE_RULES = [
    ToxCastRoleRule(
        role=ToxCastMemberRole.TESTED_CHEMICALS,
        filename_pattern=r"(?i)analytical_qc.*\.xlsx$",
        required_header_groups=[
            ["dsstox_substance_id"],
            ["spid"],
            ["chnm"],
            ["qc_level", "pass_or_caution"],
        ],
        documentation_terms=["analytical qc", "reconciled"],
    ),
    ToxCastRoleRule(
        role=ToxCastMemberRole.MULTI_CONCENTRATION_MODEL_FITS,
        filename_pattern=r"(?i)mc4.*model.*fit.*\.csv$",
        required_header_groups=[
            ["aeid"],
            ["spid", "dsstox_substance_id", "chid"],
            ["cnst_success"],
            ["hill_success", "gnls_success"],
            ["resp_max"],
        ],
        documentation_terms=["mc4", "multi-concentration model fitting"],
    ),
    ToxCastRoleRule(
        role=ToxCastMemberRole.MULTI_CONCENTRATION_HIT_CALLS,
        filename_pattern=r"(?i)mc5.*winning.*fit.*flag.*\.csv$",
        required_header_groups=[
            ["aeid"],
            ["spid", "dsstox_substance_id", "chid"],
            ["hitc"],
            ["mc6_flags"],
        ],
        documentation_terms=["mc5", "source hit-call", "mc6"],
    ),
    ToxCastRoleRule(
        role=ToxCastMemberRole.SINGLE_CONCENTRATION_ACTIVITY,
        filename_pattern=r"(?i)sc1.*sc2.*\.(?:xlsx|csv)$",
        required_header_groups=[
            ["aeid"],
            ["spid", "dsstox_substance_id", "chid"],
            ["resp"],
            ["hitc"],
        ],
        documentation_terms=["sc1", "sc2", "single-concentration"],
    ),
]


class ToxCastArchiveProcessor:
    def __init__(self, cache: ToxCastArchiveCache) -> None:
        self.cache = cache

    def _archive_path(self, stage: ToxCastArchiveStageResult) -> Path:
        path = _safe_relative(self.cache.root, stage.archive.cache_relative_path)
        if not path.is_file() or path.stat().st_size != stage.archive.size_bytes:
            raise FileNotFoundError("verified ToxCast archive cache entry is unavailable")
        return path

    @staticmethod
    def _safe_member(info: zipfile.ZipInfo) -> None:
        normalized = info.filename.replace("\\", "/")
        if (
            normalized.startswith("/")
            or re.match(r"^[A-Za-z]:", normalized)
            or ".." in Path(normalized).parts
            or info.flag_bits & 0x1
        ):
            raise ValueError("archive member failed path or encryption policy")

    @staticmethod
    def _csv_header(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> list[str]:
        with archive.open(info) as handle:
            line = handle.readline(2 * 1024 * 1024).decode("utf-8-sig", errors="strict")
        return [item.strip() for item in next(csv.reader([line]))]

    @staticmethod
    def _xlsx_headers(path: Path) -> tuple[list[str], list[str]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        headers: list[str] = []
        sheets: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                worksheet.reset_dimensions()
                header = next(worksheet.iter_rows(values_only=True), None)
                if not header:
                    continue
                sheets.append(worksheet.title)
                headers.extend(str(value).strip() for value in header if value not in (None, ""))
        finally:
            workbook.close()
        return list(dict.fromkeys(headers)), sheets

    @staticmethod
    def _xlsx_documentation_text(path: Path) -> str:
        workbook = load_workbook(path, read_only=True, data_only=True)
        values: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                worksheet.reset_dimensions()
                values.append(worksheet.title)
                for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                    if row_index > 20_000:
                        raise ValueError("reviewed documentation workbook exceeded its row bound")
                    values.extend(str(value) for value in row if value not in (None, ""))
                    if len(values) > 200_000:
                        raise ValueError("reviewed documentation workbook exceeded its cell bound")
        finally:
            workbook.close()
        return "\n".join(values)

    def enumerate(self, stage: ToxCastArchiveStageResult) -> ToxCastArchiveMemberManifest:
        archive_path = self._archive_path(stage)
        members: list[ToxCastArchiveMember] = []
        with zipfile.ZipFile(archive_path) as archive:
            for info in archive.infolist():
                self._safe_member(info)
                members.append(
                    ToxCastArchiveMember(
                        member_name=info.filename,
                        compressed_size=info.compress_size,
                        uncompressed_size=info.file_size,
                        compression_method=info.compress_type,
                        crc32=f"{info.CRC:08x}",
                        local_header_offset=info.header_offset,
                    )
                )
        required = [
            ToxCastMemberRole.TESTED_CHEMICALS,
            ToxCastMemberRole.MULTI_CONCENTRATION_MODEL_FITS,
            ToxCastMemberRole.MULTI_CONCENTRATION_HIT_CALLS,
            ToxCastMemberRole.SINGLE_CONCENTRATION_ACTIVITY,
            ToxCastMemberRole.SCHEMA_DOCUMENTATION,
        ]
        payload = {
            "release_id": TOXCAST_RELEASE,
            "source_version": TOXCAST_SOURCE_VERSION,
            "archive_sha256": stage.archive.sha256,
            "archive_size_bytes": stage.archive.size_bytes,
            "members": [item.model_dump(mode="json") for item in members],
            "exact_member_count": len(members),
            "required_roles": [item.value for item in required],
            "all_required_roles_verified": False,
            "ambiguity_findings": [
                "Scientific roles require documentation and header verification."
            ],
        }
        return ToxCastArchiveMemberManifest.model_validate(
            {**payload, "manifest_fingerprint": deterministic_fingerprint(payload)}
        )

    def _extract_member(
        self,
        archive: zipfile.ZipFile,
        info: zipfile.ZipInfo,
        *,
        maximum_uncompressed_bytes: int,
    ) -> CachedProviderFile:
        self._safe_member(info)
        if info.file_size > maximum_uncompressed_bytes:
            raise ValueError("reviewed member exceeds its extraction limit")
        member_root = self.cache.release_root / "members"
        member_root.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(prefix="member-", dir=member_root)
        digest = hashlib.sha256()
        size = 0
        try:
            with archive.open(info) as source, os.fdopen(descriptor, "wb") as target:
                while chunk := source.read(8 * 1024 * 1024):
                    size += len(chunk)
                    if size > maximum_uncompressed_bytes:
                        raise ValueError("reviewed member exceeded its extraction limit")
                    digest.update(chunk)
                    target.write(chunk)
                target.flush()
                os.fsync(target.fileno())
            if size != info.file_size:
                raise ValueError("extracted member size did not match the central directory")
            sha256 = digest.hexdigest()
            suffix = Path(info.filename).suffix.lower()
            final = member_root / f"member-{sha256[:20]}{suffix}"
            if final.exists():
                existing_hash, existing_size = _hash_file(final)
                if existing_hash != sha256 or existing_size != size:
                    raise ValueError("short provider-cache member name collided")
                Path(temporary_name).unlink(missing_ok=True)
            else:
                Path(temporary_name).replace(final)
            try:
                final.chmod(stat.S_IREAD)
            except OSError:
                pass
            return CachedProviderFile(
                cache_relative_path=final.relative_to(self.cache.root).as_posix(),
                sha256=sha256,
                size_bytes=size,
            )
        finally:
            Path(temporary_name).unlink(missing_ok=True)

    def verify_and_extract_roles(
        self,
        stage: ToxCastArchiveStageResult,
        manifest: ToxCastArchiveMemberManifest,
        *,
        role_rules: Iterable[ToxCastRoleRule] = DEFAULT_ROLE_RULES,
        maximum_member_bytes: int = 20_000_000_000,
    ) -> ToxCastArchiveMemberManifest:
        archive_path = self._archive_path(stage)
        by_name = {item.member_name: item for item in manifest.members}
        prior_by_name: dict[str, ToxCastArchiveMember] = {}
        prior_path = self.cache.release_root / "archive-member-manifest.json"
        if prior_path.is_file():
            prior = ToxCastArchiveMemberManifest.model_validate_json(
                prior_path.read_text(encoding="utf-8")
            )
            if prior.archive_sha256 == manifest.archive_sha256:
                prior_by_name = {item.member_name: item for item in prior.members}
        ambiguity: list[str] = []
        documentation_names = [
            name
            for name in by_name
            if re.search(r"(?i)(readme|data.?dictionary|schema|documentation|^methods_)", name)
            and Path(name).suffix.lower() in {".txt", ".md", ".csv", ".html", ".xlsx"}
        ]
        documentation_texts: list[str] = [TOXCAST_REVIEWED_ROLE_DOCUMENTATION]
        updated = dict(by_name)
        with zipfile.ZipFile(archive_path) as archive:
            for name in documentation_names:
                info = archive.getinfo(name)
                if info.file_size > 50_000_000:
                    continue
                extracted = self._extract_member(
                    archive, info, maximum_uncompressed_bytes=50_000_000
                )
                path = _safe_relative(self.cache.root, extracted.cache_relative_path)
                schema_columns: list[str] = []
                if path.suffix.lower() == ".xlsx":
                    documentation_texts.append(self._xlsx_documentation_text(path))
                    schema_columns, _sheets = self._xlsx_headers(path)
                else:
                    documentation_texts.append(path.read_text(encoding="utf-8", errors="replace"))
                updated[name] = updated[name].model_copy(
                    update={
                        "scientific_role": ToxCastMemberRole.SCHEMA_DOCUMENTATION,
                        "role_verification_status": "verified_documentation",
                        "role_verification_evidence": [
                            "Archive member name and workbook structure identify schema methods.",
                            (
                                "Bounded documentation content was extracted from the "
                                "immutable archive."
                            ),
                            *TOXCAST_RELEASE_DOCUMENTATION_LOCATORS,
                        ],
                        "schema_columns": schema_columns,
                        "extracted_file": extracted,
                    }
                )
            documentation_corpus = "\n".join(documentation_texts).casefold()
            for rule in role_rules:
                candidates = [name for name in by_name if re.search(rule.filename_pattern, name)]
                verified: list[tuple[str, list[str], CachedProviderFile, list[str]]] = []
                for name in candidates:
                    info = archive.getinfo(name)
                    documentation_match = all(
                        term.casefold() in documentation_corpus for term in rule.documentation_terms
                    )
                    if not documentation_match:
                        continue
                    prior_member = prior_by_name.get(name)
                    if (
                        prior_member is not None
                        and prior_member.scientific_role is rule.role
                        and prior_member.extracted_file is not None
                        and prior_member.role_verification_status == "verified"
                    ):
                        prior_reference = prior_member.extracted_file
                        prior_local = _safe_relative(
                            self.cache.root, prior_reference.cache_relative_path
                        )
                        if prior_local.is_file():
                            prior_hash, prior_size = _hash_file(prior_local)
                            if (
                                prior_hash == prior_reference.sha256
                                and prior_size == prior_reference.size_bytes
                            ):
                                verified.append(
                                    (
                                        name,
                                        prior_member.schema_columns,
                                        prior_reference,
                                        [],
                                    )
                                )
                                continue
                    suffix = Path(name).suffix.lower()
                    sheets: list[str] = []
                    if suffix == ".csv":
                        headers = self._csv_header(archive, info)
                        extracted = self._extract_member(
                            archive,
                            info,
                            maximum_uncompressed_bytes=maximum_member_bytes,
                        )
                    elif suffix == ".xlsx":
                        extracted = self._extract_member(
                            archive,
                            info,
                            maximum_uncompressed_bytes=maximum_member_bytes,
                        )
                        path = _safe_relative(self.cache.root, extracted.cache_relative_path)
                        headers, sheets = self._xlsx_headers(path)
                    else:
                        continue
                    folded_headers = {item.casefold() for item in headers}
                    header_match = all(
                        any(alias.casefold() in folded_headers for alias in group)
                        for group in rule.required_header_groups
                    )
                    if header_match and documentation_match:
                        verified.append((name, headers, extracted, sheets))
                if len(verified) != 1:
                    ambiguity.append(
                        f"{rule.role.value}: expected one documentation/header-verified member, "
                        f"found {len(verified)}."
                    )
                    continue
                name, headers, extracted, sheets = verified[0]
                updated[name] = updated[name].model_copy(
                    update={
                        "scientific_role": rule.role,
                        "role_verification_status": "verified",
                        "role_verification_evidence": [
                            "Candidate filename was used only to shortlist the member.",
                            "Required source columns were present.",
                            (
                                "Archive methods and official EPA release documentation "
                                "established semantics."
                            ),
                            (
                                "Complete tested-chemical coverage is reconciled from this QC "
                                "reference plus every verified activity member."
                            )
                            if rule.role is ToxCastMemberRole.TESTED_CHEMICALS
                            else "Role does not imply a derived tested-chemical union.",
                            f"Workbook sheets: {', '.join(sheets)}"
                            if sheets
                            else "Verified schema headers retained from immutable member cache."
                            if Path(name).suffix.lower() == ".xlsx"
                            else "CSV header verified.",
                            *TOXCAST_RELEASE_DOCUMENTATION_LOCATORS,
                        ],
                        "schema_columns": headers,
                        "extracted_file": extracted,
                    }
                )
        verified_roles = {
            item.scientific_role
            for item in updated.values()
            if item.role_verification_status.startswith("verified")
        }
        # A complete tested-chemical catalogue must have its own documentation-backed
        # member. Activity rows are never silently promoted to that role.
        if ToxCastMemberRole.TESTED_CHEMICALS not in verified_roles:
            ambiguity.append(
                "complete_tested_chemical_catalogue: no independently verified member role."
            )
        required = set(manifest.required_roles)
        payload = {
            **manifest.model_dump(mode="json", exclude={"manifest_fingerprint"}),
            "members": [updated[name].model_dump(mode="json") for name in by_name],
            "all_required_roles_verified": required <= verified_roles,
            "ambiguity_findings": list(dict.fromkeys(ambiguity)),
        }
        completed = ToxCastArchiveMemberManifest.model_validate(
            {**payload, "manifest_fingerprint": deterministic_fingerprint(payload)}
        )
        _atomic_json(
            self.cache.release_root / "archive-member-manifest.json",
            completed.model_dump(mode="json"),
        )
        return completed


class ToxCastNormalizationResult(ImmutableV2Contract):
    sqlite: CachedProviderFile
    manifest: CachedProviderFile
    tested_chemical_rows: int = Field(ge=0)
    tested_chemical_reference_rows: int = Field(ge=0)
    multi_concentration_model_fit_rows: int = Field(ge=0)
    multi_concentration_hit_call_rows: int = Field(ge=0)
    single_concentration_rows: int = Field(ge=0)
    index_names: list[str]
    input_member_hashes: dict[str, str]
    normalization_ran: bool
    bundle_fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")


def _aliases(folded: dict[str, Any], *names: str) -> Any:
    for name in names:
        value = folded.get(name.casefold())
        if value not in (None, ""):
            return value
    return None


def _iter_table(path: Path) -> Iterator[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            yield from (dict(row) for row in csv.DictReader(handle))
        return
    if path.suffix.lower() != ".xlsx":
        raise ValueError("reviewed ToxCast member uses an unsupported table format")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            worksheet.reset_dimensions()
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            columns = [str(value).strip() if value is not None else "" for value in header]
            if any(not item for item in columns):
                continue
            for values in rows:
                yield {
                    **{key: value for key, value in zip(columns, values, strict=False)},
                    "_source_sheet": worksheet.title,
                }
    finally:
        workbook.close()


class ToxCastActivityNormalizer:
    INDEX_NAMES = [
        "idx_tested_dtxsid",
        "idx_tested_source_chemical_id",
        "idx_tested_records_dtxsid_source",
        "idx_tested_cid",
        "idx_tested_inchikey",
        "idx_mc_fit_dtxsid_aeid",
        "idx_mc_hit_dtxsid_aeid",
        "idx_sc_dtxsid_aeid",
        "idx_mc_fit_aeid_dtxsid",
        "idx_mc_hit_aeid_dtxsid",
        "idx_sc_aeid_dtxsid",
        "idx_mc_hit_call",
        "idx_sc_hit_call",
        "idx_mc_flags",
        "idx_sc_flags",
        "idx_mc_fit_source_type",
        "idx_mc_hit_source_type",
        "idx_sc_source_type",
    ]

    def __init__(
        self,
        cache: ToxCastArchiveCache,
        *,
        temporary_root: Path | None = None,
    ) -> None:
        self.cache = cache
        self.temporary_root = (
            Path(temporary_root).resolve()
            if temporary_root is not None
            else self.cache.release_root
        )
        self.temporary_root.mkdir(parents=True, exist_ok=True)
        if self.temporary_root.stat().st_dev != self.cache.release_root.stat().st_dev:
            raise ValueError("ToxCast normalization temp and cache must share one filesystem")

    @property
    def manifest_path(self) -> Path:
        return self.cache.release_root / "normalized-activity-manifest.json"

    def _cached(self, member_hashes: dict[str, str]) -> ToxCastNormalizationResult | None:
        if not self.manifest_path.is_file():
            return None
        result = ToxCastNormalizationResult.model_validate_json(
            self.manifest_path.read_text(encoding="utf-8")
        )
        path = _safe_relative(self.cache.root, result.sqlite.cache_relative_path)
        if result.input_member_hashes != member_hashes or not path.is_file():
            return None
        if set(result.index_names) != set(self.INDEX_NAMES):
            return None
        return result.model_copy(update={"normalization_ran": False})

    def load_cached(self) -> ToxCastNormalizationResult | None:
        """Load the immutable normalized bundle without parsing source members again."""

        if not self.manifest_path.is_file():
            return None
        result = ToxCastNormalizationResult.model_validate_json(
            self.manifest_path.read_text(encoding="utf-8")
        )
        sqlite_path = _safe_relative(self.cache.root, result.sqlite.cache_relative_path)
        manifest_path = _safe_relative(self.cache.root, result.manifest.cache_relative_path)
        if not sqlite_path.is_file() or not manifest_path.is_file():
            return None
        if sqlite_path.stat().st_size != result.sqlite.size_bytes:
            return None
        if set(result.index_names) != set(self.INDEX_NAMES):
            return None
        with closing(sqlite3.connect(sqlite_path)) as connection:
            indexes = {
                str(row[0])
                for row in connection.execute(
                    "SELECT name FROM sqlite_master WHERE type='index' AND name IS NOT NULL"
                )
            }
        if not set(self.INDEX_NAMES) <= indexes:
            return None
        return result.model_copy(update={"normalization_ran": False})

    @staticmethod
    def _create_schema(connection: sqlite3.Connection) -> None:
        connection.executescript(
            """
            PRAGMA page_size=8192;
            PRAGMA journal_mode=OFF;
            PRAGMA synchronous=OFF;
            PRAGMA locking_mode=EXCLUSIVE;
            PRAGMA cache_size=-262144;
            CREATE TABLE tested_chemicals (
                row_id INTEGER PRIMARY KEY, chemical_key TEXT NOT NULL UNIQUE,
                dtxsid TEXT, source_chemical_id TEXT,
                casrn TEXT, chemical_name TEXT, cid TEXT, inchikey TEXT,
                qc_status TEXT, catalogue_origin TEXT NOT NULL,
                original_fields_json TEXT NOT NULL,
                source_member_sha256 TEXT NOT NULL, release_version TEXT NOT NULL
            );
            CREATE TABLE tested_chemical_records (
                row_id INTEGER PRIMARY KEY, dtxsid TEXT, source_chemical_id TEXT,
                casrn TEXT, chemical_name TEXT, cid TEXT, inchikey TEXT,
                qc_status TEXT, original_fields_json TEXT NOT NULL,
                source_member_sha256 TEXT NOT NULL, release_version TEXT NOT NULL
            );
            CREATE TABLE multi_concentration_model_fits (
                row_id INTEGER PRIMARY KEY, dtxsid TEXT, source_chemical_id TEXT,
                aeid TEXT, tested_status TEXT, model_name TEXT, activity_value TEXT,
                potency_value TEXT, activity_source_type TEXT NOT NULL,
                original_fields_json TEXT NOT NULL,
                source_member_sha256 TEXT NOT NULL, release_version TEXT NOT NULL
            );
            CREATE TABLE multi_concentration_hit_calls (
                row_id INTEGER PRIMARY KEY, dtxsid TEXT, source_chemical_id TEXT,
                aeid TEXT, tested_status TEXT, source_hit_call TEXT,
                activity_value TEXT, potency_value TEXT, flags_json TEXT NOT NULL,
                activity_source_type TEXT NOT NULL,
                original_fields_json TEXT NOT NULL, source_member_sha256 TEXT NOT NULL,
                release_version TEXT NOT NULL
            );
            CREATE TABLE single_concentration_activity (
                row_id INTEGER PRIMARY KEY, dtxsid TEXT, source_chemical_id TEXT,
                aeid TEXT, source_hit_call TEXT, tested_concentration TEXT,
                activity_value TEXT, flags_json TEXT NOT NULL,
                activity_source_type TEXT NOT NULL,
                original_fields_json TEXT NOT NULL, source_member_sha256 TEXT NOT NULL,
                release_version TEXT NOT NULL
            );
            """
        )

    @staticmethod
    def _create_indexes(connection: sqlite3.Connection) -> None:
        connection.executescript(
            """
            CREATE INDEX idx_tested_dtxsid ON tested_chemicals(dtxsid);
            CREATE INDEX idx_tested_source_chemical_id
                ON tested_chemicals(source_chemical_id);
            CREATE INDEX idx_tested_records_dtxsid_source
                ON tested_chemical_records(dtxsid, source_chemical_id);
            CREATE INDEX idx_tested_cid ON tested_chemicals(cid);
            CREATE INDEX idx_tested_inchikey ON tested_chemicals(inchikey);
            CREATE INDEX idx_mc_fit_dtxsid_aeid
                ON multi_concentration_model_fits(dtxsid, aeid);
            CREATE INDEX idx_mc_hit_dtxsid_aeid
                ON multi_concentration_hit_calls(dtxsid, aeid);
            CREATE INDEX idx_sc_dtxsid_aeid
                ON single_concentration_activity(dtxsid, aeid);
            CREATE INDEX idx_mc_fit_aeid_dtxsid
                ON multi_concentration_model_fits(aeid, dtxsid);
            CREATE INDEX idx_mc_hit_aeid_dtxsid
                ON multi_concentration_hit_calls(aeid, dtxsid);
            CREATE INDEX idx_sc_aeid_dtxsid
                ON single_concentration_activity(aeid, dtxsid);
            CREATE INDEX idx_mc_hit_call
                ON multi_concentration_hit_calls(source_hit_call);
            CREATE INDEX idx_sc_hit_call ON single_concentration_activity(source_hit_call);
            CREATE INDEX idx_mc_flags ON multi_concentration_hit_calls(flags_json);
            CREATE INDEX idx_sc_flags ON single_concentration_activity(flags_json);
            CREATE INDEX idx_mc_fit_source_type
                ON multi_concentration_model_fits(activity_source_type);
            CREATE INDEX idx_mc_hit_source_type
                ON multi_concentration_hit_calls(activity_source_type);
            CREATE INDEX idx_sc_source_type
                ON single_concentration_activity(activity_source_type);
            """
        )

    @staticmethod
    def _flush_rows(
        connection: sqlite3.Connection,
        insert_sql: str,
        pending_rows: list[tuple[Any, ...]],
    ) -> None:
        if pending_rows:
            connection.executemany(insert_sql, pending_rows)
            pending_rows.clear()

    def normalize(self, manifest: ToxCastArchiveMemberManifest) -> ToxCastNormalizationResult:
        if not manifest.all_required_roles_verified:
            raise ValueError("ToxCast member roles remain ambiguous; normalization is prohibited")
        selected = {
            item.scientific_role: item
            for item in manifest.members
            if item.extracted_file is not None and item.scientific_role is not None
        }
        role_order = [
            ToxCastMemberRole.TESTED_CHEMICALS,
            ToxCastMemberRole.MULTI_CONCENTRATION_MODEL_FITS,
            ToxCastMemberRole.MULTI_CONCENTRATION_HIT_CALLS,
            ToxCastMemberRole.SINGLE_CONCENTRATION_ACTIVITY,
        ]
        required = set(role_order)
        if not required <= selected.keys():
            raise ValueError("verified ToxCast activity roles are incomplete")
        member_hashes: dict[str, str] = {}
        for role in sorted(required, key=lambda item: item.value):
            reference = selected[role].extracted_file
            if reference is None:
                raise ValueError("verified ToxCast member lacks its extracted cache reference")
            member_hashes[role.value] = reference.sha256
        if cached := self._cached(member_hashes):
            return cached
        descriptor, temporary = tempfile.mkstemp(
            prefix="toxcast-activity-", suffix=".sqlite", dir=self.temporary_root
        )
        os.close(descriptor)
        sqlite_path = Path(temporary)
        progress_path = self.cache.release_root / "normalization-progress.json"
        counts = {role: 0 for role in required}
        tested_chemical_reference_rows = 0
        seen_chemical_keys: set[str] = set()
        try:
            with closing(sqlite3.connect(sqlite_path)) as connection:
                self._create_schema(connection)
                for role in role_order:
                    member = selected[role]
                    reference = member.extracted_file
                    if reference is None:
                        raise ValueError(
                            "verified ToxCast member lacks its extracted cache reference"
                        )
                    path = _safe_relative(self.cache.root, reference.cache_relative_path)
                    insert_sql = {
                        ToxCastMemberRole.TESTED_CHEMICALS: (
                            "INSERT INTO tested_chemical_records "
                            "VALUES (NULL,?,?,?,?,?,?,?,?,?,?)"
                        ),
                        ToxCastMemberRole.MULTI_CONCENTRATION_MODEL_FITS: (
                            "INSERT INTO multi_concentration_model_fits "
                            "VALUES (NULL,?,?,?,?,?,?,?,?,?,?,?)"
                        ),
                        ToxCastMemberRole.MULTI_CONCENTRATION_HIT_CALLS: (
                            "INSERT INTO multi_concentration_hit_calls "
                            "VALUES (NULL,?,?,?,?,?,?,?,?,?,?,?,?)"
                        ),
                        ToxCastMemberRole.SINGLE_CONCENTRATION_ACTIVITY: (
                            "INSERT INTO single_concentration_activity "
                            "VALUES (NULL,?,?,?,?,?,?,?,?,?,?,?)"
                        ),
                    }[role]
                    pending_rows: list[tuple[Any, ...]] = []

                    for row in _iter_table(path):
                        folded = {str(key).casefold(): value for key, value in row.items()}
                        dtxsid = _aliases(folded, "dtxsid", "dsstox_substance_id")
                        source_id = _aliases(folded, "spid", "chid", "chnm_id", "casn")
                        aeid = _aliases(folded, "aeid")
                        flags = {
                            key: value
                            for key, value in row.items()
                            if value not in (None, "")
                            and any(
                                marker in str(key).casefold()
                                for marker in ("flag", "caution", "qc")
                            )
                        }
                        raw = _canonical_json(row)
                        chemical_key = (
                            f"dtxsid:{dtxsid}"
                            if dtxsid not in (None, "", "#N/A")
                            else f"source:{source_id}"
                            if source_id not in (None, "", "#N/A")
                            else None
                        )
                        if chemical_key is not None and chemical_key not in seen_chemical_keys:
                            connection.execute(
                                "INSERT OR IGNORE INTO tested_chemicals "
                                "VALUES (NULL,?,?,?,?,?,?,?,?,?,?,?,?)",
                                (
                                    chemical_key,
                                    dtxsid,
                                    source_id,
                                    _aliases(folded, "casrn", "casn"),
                                    _aliases(folded, "chemical_name", "chnm", "name"),
                                    _aliases(folded, "cid", "pubchem_cid"),
                                    _aliases(folded, "inchikey"),
                                    _aliases(
                                        folded,
                                        "qc_status",
                                        "pass_or_caution",
                                        "qc_level",
                                        "stability_call",
                                        "status",
                                    ),
                                    (
                                        "analytical_qc_reference"
                                        if role is ToxCastMemberRole.TESTED_CHEMICALS
                                        else role.value
                                    ),
                                    raw,
                                    reference.sha256,
                                    TOXCAST_SOURCE_VERSION,
                                ),
                            )
                            seen_chemical_keys.add(chemical_key)
                        if role is ToxCastMemberRole.TESTED_CHEMICALS:
                            pending_rows.append(
                                (
                                    dtxsid,
                                    source_id,
                                    _aliases(folded, "casrn", "casn"),
                                    _aliases(folded, "chemical_name", "chnm", "name"),
                                    _aliases(folded, "cid", "pubchem_cid"),
                                    _aliases(folded, "inchikey"),
                                    _aliases(
                                        folded,
                                        "qc_status",
                                        "pass_or_caution",
                                        "qc_level",
                                        "stability_call",
                                        "status",
                                    ),
                                    raw,
                                    reference.sha256,
                                    TOXCAST_SOURCE_VERSION,
                                ),
                            )
                            tested_chemical_reference_rows += 1
                        elif role is ToxCastMemberRole.MULTI_CONCENTRATION_MODEL_FITS:
                            pending_rows.append(
                                (
                                    dtxsid,
                                    source_id,
                                    aeid,
                                    _aliases(folded, "tested", "tested_status"),
                                    _aliases(folded, "modl", "model", "model_name"),
                                    _aliases(folded, "resp_max", "top", "activity_value"),
                                    _aliases(folded, "modl_acc", "acc", "ac50", "potency"),
                                    "multi_concentration_all_model_fits",
                                    raw,
                                    reference.sha256,
                                    TOXCAST_SOURCE_VERSION,
                                ),
                            )
                        elif role is ToxCastMemberRole.MULTI_CONCENTRATION_HIT_CALLS:
                            pending_rows.append(
                                (
                                    dtxsid,
                                    source_id,
                                    aeid,
                                    _aliases(folded, "tested", "tested_status"),
                                    _aliases(folded, "hitc", "hit_call"),
                                    _aliases(folded, "resp_max", "top", "activity_value"),
                                    _aliases(folded, "modl_acc", "acc", "ac50", "potency"),
                                    _canonical_json(flags),
                                    "multi_concentration_winning_fit_hit_call_flags",
                                    raw,
                                    reference.sha256,
                                    TOXCAST_SOURCE_VERSION,
                                ),
                            )
                        else:
                            pending_rows.append(
                                (
                                    dtxsid,
                                    source_id,
                                    aeid,
                                    _aliases(folded, "hitc", "hit_call", "outcome"),
                                    _aliases(folded, "conc", "tested_concentration"),
                                    _aliases(folded, "resp", "response", "activity_value"),
                                    _canonical_json(flags),
                                    (
                                        "single_concentration_raw_response"
                                        if str(row.get("_source_sheet", "")).casefold() == "sc1"
                                        else "single_concentration_summary_hit_call"
                                        if str(row.get("_source_sheet", "")).casefold() == "sc2"
                                        else "single_concentration_source_record"
                                    ),
                                    raw,
                                    reference.sha256,
                                    TOXCAST_SOURCE_VERSION,
                                ),
                            )
                        counts[role] += 1
                        if len(pending_rows) >= 2_000:
                            self._flush_rows(connection, insert_sql, pending_rows)
                        if counts[role] % 50_000 == 0:
                            self._flush_rows(connection, insert_sql, pending_rows)
                            connection.commit()
                            _atomic_json(
                                progress_path,
                                {
                                    "current_role": role.value,
                                    "role_rows_processed": counts[role],
                                    "all_role_rows_processed": {
                                        item.value: counts[item] for item in role_order
                                    },
                                    "temporary_sqlite_size_bytes": sqlite_path.stat().st_size,
                                },
                            )
                    self._flush_rows(connection, insert_sql, pending_rows)
                    connection.commit()
                counts[ToxCastMemberRole.TESTED_CHEMICALS] = int(
                    connection.execute("SELECT COUNT(*) FROM tested_chemicals").fetchone()[0]
                )
                connection.commit()
                self._create_indexes(connection)
                connection.commit()
            sqlite_hash, sqlite_size = _hash_file(sqlite_path)
            final = self.cache.release_root / "normalized-activity.sqlite"
            if final.exists():
                final.chmod(stat.S_IWRITE)
                final.unlink()
            sqlite_path.replace(final)
            sqlite_reference = CachedProviderFile(
                cache_relative_path=final.relative_to(self.cache.root).as_posix(),
                sha256=sqlite_hash,
                size_bytes=sqlite_size,
            )
            payload = {
                "sqlite": sqlite_reference.model_dump(mode="json"),
                "tested_chemical_rows": counts[ToxCastMemberRole.TESTED_CHEMICALS],
                "tested_chemical_reference_rows": tested_chemical_reference_rows,
                "multi_concentration_model_fit_rows": counts[
                    ToxCastMemberRole.MULTI_CONCENTRATION_MODEL_FITS
                ],
                "multi_concentration_hit_call_rows": counts[
                    ToxCastMemberRole.MULTI_CONCENTRATION_HIT_CALLS
                ],
                "single_concentration_rows": counts[
                    ToxCastMemberRole.SINGLE_CONCENTRATION_ACTIVITY
                ],
                "index_names": self.INDEX_NAMES,
                "input_member_hashes": member_hashes,
                "normalization_ran": True,
            }
            fingerprint = deterministic_fingerprint(payload)
            placeholder = CachedProviderFile(
                cache_relative_path=self.manifest_path.relative_to(self.cache.root).as_posix(),
                sha256="0" * 64,
                size_bytes=0,
            )
            _atomic_json(
                self.manifest_path,
                {
                    **payload,
                    "manifest": placeholder.model_dump(mode="json"),
                    "bundle_fingerprint": fingerprint,
                },
            )
            manifest_hash, manifest_size = _hash_file(self.manifest_path)
            result = ToxCastNormalizationResult.model_validate(
                {
                    **payload,
                    "manifest": placeholder.model_copy(
                        update={"sha256": manifest_hash, "size_bytes": manifest_size}
                    ).model_dump(mode="json"),
                    "bundle_fingerprint": fingerprint,
                }
            )
            _atomic_json(self.manifest_path, result.model_dump(mode="json"))
            return result
        finally:
            sqlite_path.unlink(missing_ok=True)
            progress_path.unlink(missing_ok=True)


class ToxCastCoverageQuery(ImmutableV2Contract):
    biological_target: str = Field(min_length=2, max_length=300)
    matched_aeids: list[str] = Field(max_length=20_000)
    modality_by_aeid: dict[str, str] = Field(default_factory=dict, max_length=20_000)
    upstream_identity_set: list[str] = Field(default_factory=list, max_length=100_000)


class ToxCastCoverageResult(ImmutableV2Contract):
    target: str
    matched_aeid_count: int = Field(ge=0)
    modality_counts: dict[str, int]
    unique_tested_compounds: int = Field(ge=0)
    compounds_with_stable_identifiers: int = Field(ge=0)
    multi_concentration_rows: int = Field(ge=0)
    single_concentration_rows: int = Field(ge=0)
    multi_hit_call_distribution: dict[str, int]
    single_hit_call_distribution: dict[str, int]
    missing_identifier_rows: int = Field(ge=0)
    flagged_rows: int = Field(ge=0)
    upstream_overlap_count: int = Field(ge=0)
    source_version: str = TOXCAST_SOURCE_VERSION
    normalized_sqlite_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    final_source_selection_performed: bool = False


class ToxCastCoverageService:
    def __init__(self, cache: ToxCastArchiveCache) -> None:
        self.cache = cache

    def compute(
        self,
        normalized: ToxCastNormalizationResult,
        query: ToxCastCoverageQuery,
    ) -> ToxCastCoverageResult:
        if not query.matched_aeids:
            return ToxCastCoverageResult(
                target=query.biological_target,
                matched_aeid_count=0,
                modality_counts={},
                unique_tested_compounds=0,
                compounds_with_stable_identifiers=0,
                multi_concentration_rows=0,
                single_concentration_rows=0,
                multi_hit_call_distribution={},
                single_hit_call_distribution={},
                missing_identifier_rows=0,
                flagged_rows=0,
                upstream_overlap_count=0,
                normalized_sqlite_sha256=normalized.sqlite.sha256,
            )
        path = _safe_relative(self.cache.root, normalized.sqlite.cache_relative_path)
        placeholders = ",".join("?" for _ in query.matched_aeids)
        with closing(sqlite3.connect(path)) as connection:

            def count(table: str, condition: str = "1=1", values: tuple[Any, ...] = ()) -> int:
                return int(
                    connection.execute(
                        f"SELECT COUNT(*) FROM {table} WHERE {condition}", values
                    ).fetchone()[0]
                )

            values = tuple(item.removeprefix("AEID:") for item in query.matched_aeids)
            condition = f"aeid IN ({placeholders})"
            mc_rows = count("multi_concentration_hit_calls", condition, values)
            sc_rows = count("single_concentration_activity", condition, values)
            compounds = {
                str(row[0])
                for table in (
                    "multi_concentration_hit_calls",
                    "single_concentration_activity",
                )
                for row in connection.execute(
                    f"SELECT DISTINCT COALESCE(dtxsid, source_chemical_id) FROM {table} "
                    f"WHERE {condition} AND COALESCE(dtxsid, source_chemical_id) IS NOT NULL",
                    values,
                )
            }
            stable_compounds = {
                str(row[0])
                for table in (
                    "multi_concentration_hit_calls",
                    "single_concentration_activity",
                )
                for row in connection.execute(
                    f"SELECT DISTINCT COALESCE(dtxsid, source_chemical_id) FROM {table} "
                    f"WHERE {condition} AND dtxsid IS NOT NULL",
                    values,
                )
            }
            source_ids = compounds - stable_compounds
            if source_ids:
                source_placeholders = ",".join("?" for _ in source_ids)
                stable_compounds.update(
                    str(row[0])
                    for row in connection.execute(
                        "SELECT source_chemical_id FROM tested_chemicals "
                        f"WHERE source_chemical_id IN ({source_placeholders}) "
                        "AND (dtxsid IS NOT NULL OR cid IS NOT NULL OR inchikey IS NOT NULL)",
                        tuple(sorted(source_ids)),
                    )
                )
            missing = sum(
                count(
                    table,
                    f"{condition} AND dtxsid IS NULL AND source_chemical_id IS NULL",
                    values,
                )
                for table in (
                    "multi_concentration_hit_calls",
                    "single_concentration_activity",
                )
            )
            flagged = sum(
                count(table, f"{condition} AND flags_json != '{{}}'", values)
                for table in (
                    "multi_concentration_hit_calls",
                    "single_concentration_activity",
                )
            )

            def distribution(table: str) -> dict[str, int]:
                return {
                    str(row[0] if row[0] is not None else "missing"): int(row[1])
                    for row in connection.execute(
                        f"SELECT source_hit_call, COUNT(*) FROM {table} "
                        f"WHERE {condition} GROUP BY source_hit_call",
                        values,
                    )
                }

            multi_distribution = distribution("multi_concentration_hit_calls")
            single_distribution = distribution("single_concentration_activity")

        upstream = {item for item in query.upstream_identity_set}
        modality_counts: dict[str, int] = {}
        for aeid in query.matched_aeids:
            modality = query.modality_by_aeid.get(aeid, "unclassified")
            modality_counts[modality] = modality_counts.get(modality, 0) + 1
        return ToxCastCoverageResult(
            target=query.biological_target,
            matched_aeid_count=len(set(query.matched_aeids)),
            modality_counts=modality_counts,
            unique_tested_compounds=len(compounds),
            compounds_with_stable_identifiers=len(stable_compounds),
            multi_concentration_rows=mc_rows,
            single_concentration_rows=sc_rows,
            multi_hit_call_distribution=multi_distribution,
            single_hit_call_distribution=single_distribution,
            missing_identifier_rows=missing,
            flagged_rows=flagged,
            upstream_overlap_count=len(compounds & upstream),
            normalized_sqlite_sha256=normalized.sqlite.sha256,
        )
