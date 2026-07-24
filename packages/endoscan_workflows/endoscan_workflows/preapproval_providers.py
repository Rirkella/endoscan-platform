"""Complete metadata-only provider execution for workflow semantics v2.

The contracts in this module intentionally keep complete scientific record sets out
of workflow JSON and agent context.  Reviewed requests are executed through the
provider-neutral executor, immutable raw responses are normalized into SQLite, and
tool outputs contain only content-addressed references, exact counts, fingerprints,
and bounded previews.

This is a PRE-APPROVAL layer.  It cannot select sources, construct an assembly
recipe, retrieve an expression matrix, infer activity thresholds, or standardize
chemical forms.
"""

from __future__ import annotations

import csv
import gzip
import io
import json
import math
import re
import sqlite3
import tempfile
import xml.etree.ElementTree as ET
from collections.abc import Iterator
from contextlib import closing
from datetime import date, datetime
from enum import StrEnum
from pathlib import Path
from typing import Any, Literal
from urllib.parse import quote, urlparse

from openpyxl import load_workbook
from pydantic import Field, model_validator

from .artifacts import LocalArtifactStore
from .contracts import ToolInvocation
from .discovery_strategy import (
    ArtifactReference,
    CandidateStatus,
    DiscoveryExecutionRecord,
    DiscoveryTaskStatus,
    EvidenceRole,
    HydratedSource,
    HydrationCompleteness,
    ImmutableV2Contract,
    ProviderDiscoveryTask,
    SourceCandidate,
    deterministic_fingerprint,
)
from .provider_execution import (
    ProviderExecutionOutcome,
    ProviderItemKind,
    ProviderResourceRequest,
    ProviderRetrievalTask,
    ProviderTaskExecutor,
)
from .source_governor import current_source_request_context
from .toxcast_public_activity import (
    TOXCAST_ARCHIVE_SHA256,
    ToxCastActivityNormalizer,
    ToxCastArchiveCache,
    ToxCastCoverageQuery,
    ToxCastCoverageResult,
    ToxCastCoverageService,
    ToxCastNormalizationResult,
)

PREAPPROVAL_PROVIDER_REGISTRY_VERSION: Literal["1.0.0"] = "1.0.0"
PREAPPROVAL_NORMALIZED_SCHEMA_VERSION: Literal["1.0.0"] = "1.0.0"
PREVIEW_LIMIT = 20
MAXIMUM_PREAPPROVAL_NORMALIZED_SQLITE_BYTES = 400_000_000


class ProviderRecordKind(StrEnum):
    RELEASE = "release"
    ASSAY = "assay"
    ACTIVITY_SUMMARY = "activity_summary"
    ACTIVITY = "activity"
    COMPOUND = "compound"
    RELATIONSHIP = "relationship"
    TRANSCRIPTOMIC = "transcriptomic"
    PUBLICATION = "publication"
    DATASET_LINK = "dataset_link"


class ProviderRetrievalMode(StrEnum):
    MANIFEST_FILES = "manifest_files"
    CURSOR_PAGES = "cursor_pages"
    CURSOR_HYDRATION = "cursor_hydration"
    IDENTIFIER_BATCHES = "identifier_batches"
    TOXCAST_AUTHORITATIVE = "toxcast_authoritative"


class ProviderAccessMode(StrEnum):
    PUBLIC_RELEASE = "public_release"
    AUTHENTICATED_API = "authenticated_api"


class ProviderCapabilityStatus(StrEnum):
    OPERATIONAL = "operational"
    MISSING = "missing"
    HEAVY_FILE_REQUIRES_APPROVAL = "heavy_file_requires_approval"
    OPTIONAL_ACCELERATION_UNAVAILABLE = "optional_acceleration_unavailable"


class ProviderResourceScope(StrEnum):
    RELEASE_METADATA = "release_metadata"
    SCIENTIFIC_DATA = "scientific_data"


class ProviderExpansionMode(StrEnum):
    STATIC = "static"
    CANDIDATE_SUMMARY_BATCH = "candidate_summary_batch"
    CANDIDATE_COMPOUND_IDENTIFIERS = "candidate_compound_identifiers"
    CANDIDATE_ACTIVITY_ROWS = "candidate_activity_rows"
    CANDIDATE_RELATIONSHIPS = "candidate_relationships"
    GEO_SERIES_SOFT = "geo_series_soft"
    MATCHED_AEID_SUMMARY = "matched_aeid_summary"
    MATCHED_AEID_DETAILS = "matched_aeid_details"


class ProviderApiContract(ImmutableV2Contract):
    openapi_locator: str = Field(pattern=r"^https://[^\s]+$")
    openapi_version: str = Field(min_length=1, max_length=40)
    api_version: str = Field(min_length=1, max_length=40)
    documentation_revision: str = Field(min_length=3, max_length=160)
    schema_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    official_implementation: str = Field(pattern=r"^https://[^\s]+$")
    official_client_reference: str = Field(pattern=r"^https://[^\s]+$")
    authentication_scheme: Literal["x-api-key"]


class ProviderReleaseAuthority(ImmutableV2Contract):
    release_date: str = Field(pattern=r"^\d{4}-\d{2}(?:-\d{2})?$")
    citation_doi: str = Field(pattern=r"^https://doi\.org/[^\s]+$")
    release_space_id: str = Field(pattern=r"^[a-f0-9]{24}$")
    reviewed_dataset_ids: dict[str, str] = Field(min_length=1, max_length=20)
    reviewed_file_ids: dict[str, str] = Field(min_length=1, max_length=20)


class ProviderBulkFallbackResource(ImmutableV2Contract):
    resource_id: str = Field(pattern=r"^[a-z][a-z0-9_.-]{2,159}$")
    file_id: str = Field(pattern=r"^[a-f0-9]{24}$")
    dataset_id: str = Field(pattern=r"^[a-f0-9]{24}$")
    locator: str = Field(pattern=r"^https://[^\s]+$")
    scientific_role_status: Literal["unverified"] = "unverified"
    automatic_download: Literal[False] = False
    size_bytes: int = Field(ge=1)
    sha256: str | None = Field(default=None, pattern=r"^[a-f0-9]{64}$")
    required_members: list[str] = Field(default_factory=list, max_length=50)
    finding_code: Literal["PUBLIC_PROVIDER_HEAVY_FILE_REQUIRES_APPROVAL"] = (
        "PUBLIC_PROVIDER_HEAVY_FILE_REQUIRES_APPROVAL"
    )
    verification_required: list[str] = Field(min_length=1, max_length=20)


class ProviderCapabilityDeclaration(ImmutableV2Contract):
    capability: str = Field(pattern=r"^[a-z][a-z0-9_.-]{2,159}$")
    status: ProviderCapabilityStatus
    finding_code: (
        Literal[
            "REGISTERED_PROVIDER_CAPABILITY_MISSING",
            "PUBLIC_ACCESS_PATH_MISSING",
            "PUBLIC_PROVIDER_HEAVY_FILE_REQUIRES_APPROVAL",
        ]
        | None
    ) = None
    source_resource_ids: list[str] = Field(default_factory=list, max_length=30)
    required_table_or_member: str | None = Field(default=None, max_length=240)
    scientific_consequence: str = Field(min_length=3, max_length=500)
    metadata_only_coverage_available: bool


class ProviderResourceTemplate(ImmutableV2Contract):
    resource_id: str = Field(pattern=r"^[a-z][a-z0-9_.-]{2,159}$")
    logical_role: str = Field(min_length=2, max_length=120)
    record_kind: ProviderRecordKind
    locator_template: str = Field(pattern=r"^https://[^\s]+$")
    item_kind: ProviderItemKind
    required: bool = True
    accepted_mime_types: list[str] = Field(min_length=1, max_length=10)
    maximum_response_bytes: int = Field(ge=1_024, le=500_000_000)
    expected_sha256: str | None = Field(default=None, pattern=r"^[a-f0-9]{64}$")
    compression: Literal["gzip"] | None = None
    operation_id: str | None = Field(default=None, pattern=r"^[a-z][a-z0-9_.-]{2,159}$")
    output_schema: str | None = Field(default=None, min_length=2, max_length=160)
    resource_scope: ProviderResourceScope = ProviderResourceScope.SCIENTIFIC_DATA
    expansion_mode: ProviderExpansionMode = ProviderExpansionMode.STATIC
    required_credential: Literal["epa_comptox_api_key"] | None = None
    access_mode: ProviderAccessMode = ProviderAccessMode.PUBLIC_RELEASE
    identifier_path_kind: Literal["opaque", "pubchem_cid"] = "opaque"
    bulk_download_prohibited: bool = False


class PreapprovalProviderReleaseManifest(ImmutableV2Contract):
    release_id: str = Field(min_length=3, max_length=160)
    provider: str = Field(min_length=2, max_length=120)
    adapter_id: str = Field(min_length=3, max_length=160)
    adapter_version: str = Field(pattern=r"^\d+\.\d+\.\d+$")
    source_system: str = Field(min_length=2, max_length=160)
    source_version: str = Field(min_length=1, max_length=160)
    source_locator: str = Field(pattern=r"^https://[^\s]+$")
    reviewed_hosts: list[str] = Field(min_length=1, max_length=20)
    retrieval_mode: ProviderRetrievalMode
    resources: list[ProviderResourceTemplate] = Field(min_length=1, max_length=50)
    page_size: int | None = Field(default=None, ge=1, le=10_000)
    maximum_pages: int = Field(default=100, ge=1, le=10_000)
    identifier_batch_size: int | None = Field(default=None, ge=1, le=1_000)
    licence_and_provenance: list[str] = Field(min_length=1, max_length=50)
    release_checksum_or_citation: str = Field(min_length=3, max_length=500)
    heavy_data_prohibited: list[str] = Field(default_factory=list, max_length=50)
    default_access_mode: ProviderAccessMode = ProviderAccessMode.PUBLIC_RELEASE
    optional_access_modes: list[ProviderAccessMode] = Field(default_factory=list, max_length=5)
    public_capabilities: list[ProviderCapabilityDeclaration] = Field(
        default_factory=list, max_length=50
    )
    api_contract: ProviderApiContract | None = None
    release_authority: ProviderReleaseAuthority | None = None
    bulk_fallback_resources: list[ProviderBulkFallbackResource] = Field(
        default_factory=list, max_length=20
    )
    completion_rules: list[str] = Field(default_factory=list, max_length=50)
    cache_replay_rules: list[str] = Field(default_factory=list, max_length=20)
    unavailable_relationship_fields: list[str] = Field(default_factory=list, max_length=50)
    manifest_fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")

    @model_validator(mode="after")
    def validate_manifest(self) -> PreapprovalProviderReleaseManifest:
        resource_ids = [item.resource_id for item in self.resources]
        if len(resource_ids) != len(set(resource_ids)):
            raise ValueError("provider release resource IDs must be unique")
        reviewed = set(self.reviewed_hosts)
        for value in [self.source_locator, *(item.locator_template for item in self.resources)]:
            host = (urlparse(value).hostname or "").lower()
            if host not in reviewed:
                raise ValueError(f"provider locator host is not reviewed: {host}")
        if (
            self.retrieval_mode
            in {
                ProviderRetrievalMode.CURSOR_PAGES,
                ProviderRetrievalMode.CURSOR_HYDRATION,
            }
            and not self.page_size
        ):
            raise ValueError("cursor providers require a page size")
        if (
            self.retrieval_mode is ProviderRetrievalMode.CURSOR_HYDRATION
            and not self.identifier_batch_size
        ):
            raise ValueError("hydrating cursor providers require an identifier batch size")
        if (
            self.retrieval_mode is ProviderRetrievalMode.IDENTIFIER_BATCHES
            and not self.identifier_batch_size
        ):
            raise ValueError("identifier providers require an identifier batch size")
        if self.retrieval_mode is ProviderRetrievalMode.TOXCAST_AUTHORITATIVE:
            if (
                self.provider != "toxcast"
                or self.api_contract is None
                or self.release_authority is None
            ):
                raise ValueError(
                    "authoritative ToxCast retrieval requires API and release-authority contracts"
                )
            required_public_operations = {
                "list_release_datasets",
                "list_dataset_files",
                "inspect_public_archive_metadata",
                "public_assay_annotations",
                "public_assay_target_mappings",
                "public_chemical_analytical_qc",
                "public_cytotox_reference",
            }
            optional_authenticated_operations = {
                "all_assays",
                "summary_by_aeid",
                "details_by_aeid",
            }
            operations = [item.operation_id for item in self.resources]
            if not required_public_operations <= set(operations):
                raise ValueError("public ToxCast manifest lacks a required typed operation")
            if not optional_authenticated_operations <= set(operations):
                raise ValueError("optional authenticated ToxCast contract is incomplete")
            if self.default_access_mode is not ProviderAccessMode.PUBLIC_RELEASE:
                raise ValueError("ToxCast must default to its reproducible public release")
            if ProviderAccessMode.AUTHENTICATED_API not in self.optional_access_modes:
                raise ValueError("authenticated ToxCast API must be declared optional")
            public_resources = [
                item
                for item in self.resources
                if item.access_mode is ProviderAccessMode.PUBLIC_RELEASE
            ]
            if any(item.required_credential for item in public_resources):
                raise ValueError("public ToxCast release resources cannot require credentials")
            authenticated_resources = [
                item
                for item in self.resources
                if item.access_mode is ProviderAccessMode.AUTHENTICATED_API
            ]
            if any(
                item.required_credential != "epa_comptox_api_key"
                for item in authenticated_resources
            ):
                raise ValueError("authenticated ToxCast resources require the reviewed binding")
            scientific_operations = [
                item.operation_id
                for item in self.resources
                if item.resource_scope is ProviderResourceScope.SCIENTIFIC_DATA
            ]
            if None in scientific_operations or len(scientific_operations) != len(
                set(scientific_operations)
            ):
                raise ValueError(
                    "every ToxCast scientific role requires a distinct typed operation"
                )
            expected_schemas = {
                "all_assays": (ProviderRecordKind.ASSAY, "AssayAnnotation[]"),
                "summary_by_aeid": (
                    ProviderRecordKind.ACTIVITY_SUMMARY,
                    "AssayAgg[]",
                ),
                "details_by_aeid": (
                    ProviderRecordKind.ACTIVITY,
                    "BioactivityDataAll[]",
                ),
            }
            for item in self.resources:
                if (
                    item.resource_scope is ProviderResourceScope.RELEASE_METADATA
                    and item.record_kind is not ProviderRecordKind.RELEASE
                ):
                    raise ValueError("release metadata cannot declare a scientific row kind")
                if (
                    item.resource_scope is ProviderResourceScope.SCIENTIFIC_DATA
                    and item.record_kind is ProviderRecordKind.RELEASE
                ):
                    raise ValueError("scientific operations cannot emit release descriptors")
                if (
                    item.operation_id in expected_schemas
                    and (
                        item.record_kind,
                        item.output_schema,
                    )
                    != expected_schemas[item.operation_id]
                ):
                    raise ValueError(
                        "ToxCast operation output schema does not match its declared role"
                    )
            if any(item.automatic_download for item in self.bulk_fallback_resources):
                raise ValueError("bulk fallback must never be downloaded automatically")
        if self.manifest_fingerprint == "0" * 64:
            return self
        payload = self.model_dump(mode="json", exclude={"manifest_fingerprint"})
        if deterministic_fingerprint(payload) != self.manifest_fingerprint:
            raise ValueError("provider release manifest fingerprint does not match")
        return self

    @classmethod
    def create(cls, **values: Any) -> PreapprovalProviderReleaseManifest:
        draft = cls.model_validate({**values, "manifest_fingerprint": "0" * 64})
        normalized = draft.model_dump(mode="json", exclude={"manifest_fingerprint"})
        normalized["manifest_fingerprint"] = deterministic_fingerprint(normalized)
        return cls.model_validate(normalized)


class PreapprovalProviderReleaseRegistry(ImmutableV2Contract):
    registry_version: Literal["1.0.0"] = PREAPPROVAL_PROVIDER_REGISTRY_VERSION
    releases: list[PreapprovalProviderReleaseManifest] = Field(min_length=1, max_length=100)

    @model_validator(mode="after")
    def unique_releases(self) -> PreapprovalProviderReleaseRegistry:
        keys = [(item.provider, item.release_id) for item in self.releases]
        if len(keys) != len(set(keys)):
            raise ValueError("provider release manifests must be unique")
        return self

    def release(self, provider: str, release_id: str) -> PreapprovalProviderReleaseManifest:
        for item in self.releases:
            if item.provider == provider and item.release_id == release_id:
                return item
        raise KeyError((provider, release_id))


def load_preapproval_provider_registry(repo_root: Path) -> PreapprovalProviderReleaseRegistry:
    path = Path(repo_root) / "registry" / "data" / "preapproval_provider_releases.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["releases"] = [
        PreapprovalProviderReleaseManifest.create(
            **{key: value for key, value in item.items() if key != "manifest_fingerprint"}
        )
        for item in payload["releases"]
    ]
    return PreapprovalProviderReleaseRegistry.model_validate(payload)


class ProviderMetadataQuery(ImmutableV2Contract):
    biological_target: str | None = Field(default=None, max_length=300)
    target_identifiers: list[str] = Field(default_factory=list, max_length=50)
    target_synonyms: list[str] = Field(default_factory=list, max_length=100)
    modality: str | None = Field(default=None, max_length=120)
    publication_identifiers: list[str] = Field(default_factory=list, max_length=100)
    source_hints: list[str] = Field(default_factory=list, max_length=0)


class ProviderMetadataExecutionInput(ImmutableV2Contract):
    ledger_record: DiscoveryExecutionRecord
    release_id: str = Field(min_length=3, max_length=160)
    query: ProviderMetadataQuery
    upstream_identifier_artifact: ArtifactReference | None = None
    upstream_identifier_count: int = Field(default=0, ge=0)
    validation_scope: Literal["production_full", "bounded_smoke"] = "production_full"
    access_mode: ProviderAccessMode = ProviderAccessMode.PUBLIC_RELEASE
    maximum_summary_aeids: int | None = Field(default=None, ge=1, le=20)
    maximum_detail_aeids: int | None = Field(default=None, ge=1, le=5)
    maximum_hydration_candidates: int | None = Field(default=None, ge=1, le=500)

    @model_validator(mode="after")
    def identifier_artifact_is_complete(self) -> ProviderMetadataExecutionInput:
        if bool(self.upstream_identifier_artifact) != bool(self.upstream_identifier_count):
            raise ValueError(
                "upstream identifier artifact and exact identifier count must be supplied together"
            )
        if self.access_mode is ProviderAccessMode.PUBLIC_RELEASE and (
            self.maximum_summary_aeids is not None or self.maximum_detail_aeids is not None
        ):
            raise ValueError("public-release retrieval does not use CTX AEID request bounds")
        if self.validation_scope == "production_full" and (
            self.maximum_summary_aeids is not None or self.maximum_detail_aeids is not None
        ):
            raise ValueError("production-full retrieval cannot truncate matched AEIDs")
        if (
            self.access_mode is ProviderAccessMode.AUTHENTICATED_API
            and self.validation_scope == "bounded_smoke"
            and (self.maximum_summary_aeids is None or self.maximum_detail_aeids is None)
        ):
            raise ValueError("bounded ToxCast smoke requires explicit summary and detail bounds")
        return self


def compile_provider_metadata_input(
    task: ProviderDiscoveryTask,
    ledger_record: DiscoveryExecutionRecord,
    *,
    upstream_identifier_artifact: ArtifactReference | None = None,
    upstream_identifier_count: int = 0,
    maximum_hydration_candidates: int | None = None,
) -> ProviderMetadataExecutionInput:
    """Compile a semantics-v2 task into the exact typed provider input."""

    if (
        task.task_id != ledger_record.task_id
        or task.provider != ledger_record.provider
        or task.evidence_role is not ledger_record.evidence_role
        or task.modality != ledger_record.modality
    ):
        raise ValueError("discovery task and ledger record are not identical")
    release_id = _text(task.query_parameters.get("release_id"))
    if release_id is None:
        raise ValueError("operational provider task lacks its reviewed release ID")
    return ProviderMetadataExecutionInput(
        ledger_record=ledger_record,
        release_id=release_id,
        query=ProviderMetadataQuery(
            biological_target=_text(task.query_parameters.get("biological_target")),
            target_identifiers=list(task.query_parameters.get("target_identifiers") or []),
            target_synonyms=list(task.query_parameters.get("target_synonyms") or []),
            modality=task.modality,
            publication_identifiers=list(
                task.query_parameters.get("publication_identifiers") or []
            ),
            source_hints=list(task.query_parameters.get("source_hints") or []),
        ),
        upstream_identifier_artifact=upstream_identifier_artifact,
        upstream_identifier_count=upstream_identifier_count,
        maximum_hydration_candidates=maximum_hydration_candidates,
    )


class ProviderNormalizedCounts(ImmutableV2Contract):
    raw_records: int = Field(ge=0)
    source_candidates: int = Field(ge=0)
    hydrated_sources: int = Field(ge=0)
    compound_mappings: int = Field(ge=0)
    activity_records: int = Field(ge=0)
    activity_summaries: int = Field(default=0, ge=0)
    transcriptomic_records: int = Field(ge=0)
    assay_relationships: int = Field(ge=0)
    assay_traversal_records: int = Field(ge=0)
    assay_annotations: int = Field(default=0, ge=0)
    release_inventory_records: int = Field(default=0, ge=0)
    publication_dataset_links: int = Field(ge=0)
    exact_identity_mappings: int = Field(ge=0)
    ambiguous_identity_mappings: int = Field(ge=0)
    missing_identity_mappings: int = Field(ge=0)


class ProviderDatasetCompletionProof(ImmutableV2Contract):
    provider: str
    release_id: str
    task_id: str
    declared_resource_count: int = Field(ge=0)
    completed_resource_count: int = Field(ge=0)
    requested_identifier_count: int = Field(ge=0)
    processed_identifier_count: int = Field(ge=0)
    cursor_exhausted: bool
    file_manifest_reconciled: bool
    safety_truncated: bool
    all_upstream_compounds_processed: bool
    completed: bool
    reason: str = Field(min_length=3, max_length=1000)
    validation_scope: Literal["production_full", "bounded_smoke"] = "production_full"
    matched_aeid_count: int = Field(default=0, ge=0)
    summary_aeid_count: int = Field(default=0, ge=0)
    detail_aeid_count: int = Field(default=0, ge=0)
    discovered_candidate_count: int = Field(default=0, ge=0)
    hydrated_candidate_count: int = Field(default=0, ge=0)
    shortlisted_candidate_ids: list[str] = Field(default_factory=list, max_length=500)
    excluded_candidate_ids: list[str] = Field(default_factory=list, max_length=500)
    excluded_candidate_count: int = Field(default=0, ge=0)
    candidate_exclusion_reason: str | None = Field(default=None, max_length=500)
    shortlist_decision_artifact: ArtifactReference | None = None
    activity_availability_attempt_count: int = Field(default=0, ge=0)
    relationship_attempt_count: int = Field(default=0, ge=0)
    resolved_geo_accession_count: int = Field(default=0, ge=0)
    geo_soft_attempt_count: int = Field(default=0, ge=0)
    structural_validation_codes: list[str] = Field(default_factory=list, max_length=20)


class ProviderPageYieldObservation(ImmutableV2Contract):
    resource_id: str
    operation: str
    transport_request_count: int = Field(ge=0)
    raw_rows: int = Field(ge=0)
    normalized_rows: int = Field(ge=0)
    provider_unique_records: int = Field(ge=0)
    new_stable_compound_identifiers: int = Field(ge=0)
    new_compact_candidates: int = Field(ge=0)
    duplicates: int = Field(ge=0)
    joinability_contribution: int = Field(ge=0)
    continuation_token: str | None = None


class DiskBackedSourceCandidateSet(ImmutableV2Contract):
    artifact: ArtifactReference
    table_name: Literal["source_candidates"] = "source_candidates"
    exact_count: int = Field(ge=0)
    preview: list[SourceCandidate] = Field(default_factory=list, max_length=PREVIEW_LIMIT)
    preview_is_complete: bool


class DiskBackedHydratedSourceSet(ImmutableV2Contract):
    artifact: ArtifactReference
    table_name: Literal["hydrated_sources"] = "hydrated_sources"
    exact_count: int = Field(ge=0)
    preview: list[HydratedSource] = Field(default_factory=list, max_length=PREVIEW_LIMIT)
    preview_is_complete: bool


class ProviderNormalizedDatasetManifest(ImmutableV2Contract):
    provider: str
    release_id: str
    source_system: str
    source_version: str
    adapter_id: str
    adapter_version: str
    release_manifest_fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")
    normalized_schema_version: Literal["1.0.0"] = PREAPPROVAL_NORMALIZED_SCHEMA_VERSION
    sqlite_artifact: ArtifactReference
    raw_artifacts: list[ArtifactReference] = Field(default_factory=list, max_length=10_000)
    counts: ProviderNormalizedCounts
    modalities: list[str] = Field(default_factory=list, max_length=100)
    provider_provenance_partitions: dict[str, int] = Field(default_factory=dict)
    relationship_findings: list[str] = Field(default_factory=list, max_length=100)
    operation_hashes: dict[str, list[str]] = Field(default_factory=dict)
    index_names: list[str] = Field(default_factory=list, max_length=100)
    completion_proof: ProviderDatasetCompletionProof
    bundle_fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")
    access_mode: ProviderAccessMode = ProviderAccessMode.PUBLIC_RELEASE
    capability_findings: list[ProviderCapabilityDeclaration] = Field(
        default_factory=list, max_length=50
    )
    expression_values_retrieved: Literal[False] = False
    final_source_selection_performed: Literal[False] = False
    strategy_proposal_created: Literal[False] = False
    assembly_recipe_created: Literal[False] = False


class ToxCastPublicActivityCoverage(ImmutableV2Contract):
    """Bounded workflow view over the reusable immutable provider cache."""

    normalized_bundle: ToxCastNormalizationResult
    coverage: ToxCastCoverageResult
    matched_aeids: list[str] = Field(max_length=20_000)
    modality_by_aeid: dict[str, str] = Field(default_factory=dict, max_length=20_000)
    archive_network_request_count: Literal[0] = 0
    extraction_ran: Literal[False] = False
    normalization_ran: Literal[False] = False
    final_source_selection_performed: Literal[False] = False


class ProviderIdentifierReconciliation(ImmutableV2Contract):
    received_identifier_count: int = Field(ge=0)
    normalized_unique_identifier_count: int = Field(ge=0)
    duplicate_identifier_count: int = Field(ge=0)
    invalid_identifier_count: int = Field(ge=0)
    invalid_identifiers_preview: list[str] = Field(default_factory=list, max_length=20)
    provider_batch_size: int = Field(ge=1)
    provider_batch_count: int = Field(ge=0)
    completed_identifier_count: int = Field(default=0, ge=0)
    missing_identifier_count: int = Field(default=0, ge=0)
    mapping_artifact: ArtifactReference


class ProviderMetadataExecutionOutput(ImmutableV2Contract):
    provider: str
    release_id: str
    ledger_record: DiscoveryExecutionRecord
    execution_manifest_artifact: ArtifactReference
    exact_execution_count: int = Field(ge=0)
    page_or_file_execution_preview: list[ProviderExecutionOutcome] = Field(
        default_factory=list, max_length=PREVIEW_LIMIT
    )
    execution_preview_is_complete: bool
    completion_proof: ProviderDatasetCompletionProof
    dataset_manifest: ProviderNormalizedDatasetManifest | None = None
    dataset_manifest_artifact: ArtifactReference | None = None
    candidates: DiskBackedSourceCandidateSet | None = None
    hydrated_sources: DiskBackedHydratedSourceSet | None = None
    normalization_ran: bool = False
    cache_only_replay: bool
    access_mode: ProviderAccessMode = ProviderAccessMode.PUBLIC_RELEASE
    optional_authenticated_api_available: bool = False
    capability_findings: list[ProviderCapabilityDeclaration] = Field(
        default_factory=list, max_length=50
    )
    toxcast_public_activity: ToxCastPublicActivityCoverage | None = None
    toxcast_public_activity_artifact: ArtifactReference | None = None
    pagination_yield: list[ProviderPageYieldObservation] = Field(
        default_factory=list, max_length=10_000
    )
    identifier_reconciliation: ProviderIdentifierReconciliation | None = None


def _artifact_reference(descriptor: Any, artifact_type: str | None = None) -> ArtifactReference:
    return ArtifactReference(
        artifact_id=descriptor.id,
        sha256=descriptor.sha256,
        artifact_type=artifact_type or descriptor.artifact_type,
    )


def _safe_token(value: str) -> str:
    return re.sub(r"[^a-z0-9_.-]", "-", value.lower()).strip("-")[:120]


def _render_locator(
    template: str,
    *,
    query: ProviderMetadataQuery,
    page: int,
    page_size: int,
    cursor: str | None,
    identifiers: list[str],
    identifier_path_kind: Literal["opaque", "pubchem_cid"] = "opaque",
    aeid: int | None = None,
) -> str:
    search_terms = list(
        dict.fromkeys(
            [
                *(query.target_identifiers or []),
                *(query.target_synonyms or []),
                *([query.biological_target] if query.biological_target else []),
                *([query.modality] if query.modality else []),
            ]
        )
    )
    rendered_identifiers = identifiers
    if identifier_path_kind == "pubchem_cid":
        rendered_identifiers = []
        for identifier in identifiers:
            match = re.fullmatch(r"CID:([1-9][0-9]*)", identifier)
            if match is None:
                raise ValueError("PubChem CID path requires canonical CID:<integer> identifiers")
            rendered_identifiers.append(match.group(1))
    values = {
        "encoded_query": quote(" AND ".join(search_terms), safe=""),
        "page": str(page),
        "page_start": str((page - 1) * page_size),
        "page_size": str(page_size),
        "cursor": quote(cursor or "", safe=""),
        "identifier_batch": quote(",".join(rendered_identifiers), safe=","),
        "publication_batch": quote(",".join(query.publication_identifiers), safe=","),
        "aeid": str(aeid or ""),
    }
    try:
        return template.format(**values)
    except KeyError as exc:
        raise ValueError(f"unsupported reviewed locator placeholder: {exc.args[0]}") from exc


def _json_records(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return [
            dict(item) if isinstance(item, dict) else {"source_identifier": str(item)}
            for item in payload
        ]
    if not isinstance(payload, dict):
        return []
    records = payload.get("records")
    if isinstance(records, list):
        return [dict(item) for item in records]
    properties = payload.get("PropertyTable", {}).get("Properties")
    if isinstance(properties, list):
        return [dict(item) for item in properties]
    information = payload.get("InformationList", {}).get("Information")
    if isinstance(information, list):
        return [dict(item) for item in information]
    summaries = payload.get("DocumentSummarySet", {}).get("DocumentSummary")
    if isinstance(summaries, list):
        return [dict(item) for item in summaries]
    result = payload.get("result")
    if isinstance(result, dict) and isinstance(result.get("uids"), list):
        return [
            {
                **dict(result.get(str(identifier)) or {}),
                "source_identifier": str(identifier),
            }
            for identifier in result["uids"]
            if isinstance(result.get(str(identifier)), dict)
        ]
    identifier_list = payload.get("IdentifierList")
    if isinstance(identifier_list, dict):
        return [
            {"cid": str(identifier), "source_identifier": f"CID:{identifier}"}
            for identifier in identifier_list.get("CID", [])
            if isinstance(identifier, int) or str(identifier).isdigit()
        ]
    link_sets = payload.get("linksets")
    if isinstance(link_sets, list):
        linked_records: list[dict[str, Any]] = []
        for link_set in link_sets:
            if not isinstance(link_set, dict):
                continue
            parents = [str(item) for item in link_set.get("ids", [])]
            for database in link_set.get("linksetdbs", []):
                if not isinstance(database, dict):
                    continue
                relationship = str(database.get("linkname") or "source_declared_link")
                for parent in parents or [""]:
                    linked_records.extend(
                        {
                            "source_identifier": f"AID:{parent}:{relationship}:AID:{child}",
                            "parent_assay_identifier": f"AID:{parent}",
                            "child_assay_identifier": f"AID:{child}",
                            "relationship_type": relationship,
                            "relationship_provenance": {
                                "provider": "NCBI PubChem ELink",
                                "linkname": relationship,
                                "confidence": "source_backed",
                            },
                        }
                        for child in database.get("links", [])
                    )
        return linked_records
    identifiers = (
        payload.get("esearchresult", {}).get("idlist")
        or payload.get("result", {}).get("uids")
        or []
    )
    return [{"source_identifier": str(identifier)} for identifier in identifiers]


def _elink_xml_records(text: str) -> list[dict[str, Any]]:
    root = ET.fromstring(text)
    records: list[dict[str, Any]] = []
    for link_set in root.findall(".//LinkSet"):
        parents = [
            str(node.text).strip()
            for node in link_set.findall("./IdList/Id")
            if node.text and str(node.text).strip().isdigit()
        ]
        for database in link_set.findall("./LinkSetDb"):
            relationship = (database.findtext("LinkName") or "source_declared_link").strip()
            children = [
                str(node.text).strip()
                for node in database.findall("./Link/Id")
                if node.text and str(node.text).strip().isdigit()
            ]
            for parent in parents or [""]:
                records.extend(
                    {
                        "source_identifier": (f"AID:{parent}:{relationship}:AID:{child}"),
                        "parent_assay_identifier": f"AID:{parent}",
                        "child_assay_identifier": f"AID:{child}",
                        "relationship_type": relationship,
                        "relationship_provenance": {
                            "provider": "NCBI PubChem ELink",
                            "linkname": relationship,
                            "confidence": "source_backed",
                        },
                    }
                    for child in children
                )
    return records


def _assay_description_relationship_rows(payload: Any) -> list[dict[str, Any]]:
    if not isinstance(payload, dict):
        return []
    records: list[dict[str, Any]] = []
    for container in payload.get("PC_AssayContainer", []):
        if not isinstance(container, dict):
            continue
        description = container.get("assay", {}).get("descr", {})
        if not isinstance(description, dict):
            continue
        parent = str(description.get("aid", {}).get("id") or "")
        if not parent.isdigit():
            continue
        for cross_reference in description.get("xref", []):
            if not isinstance(cross_reference, dict):
                continue
            child = str(cross_reference.get("xref", {}).get("aid") or "")
            if not child.isdigit():
                continue
            comment = _text(cross_reference.get("comment"))
            folded_comment = (comment or "").casefold()
            relationship = (
                "source_declared_counterscreen"
                if any(
                    marker in folded_comment
                    for marker in (
                        "counter screen",
                        "counterscreen",
                        "auto fluorescence",
                        "cytotoxic",
                        "viability",
                    )
                )
                else "source_declared_assay_cross_reference"
            )
            records.append(
                {
                    "source_identifier": f"AID:{parent}:xref:AID:{child}",
                    "parent_assay_identifier": f"AID:{parent}",
                    "child_assay_identifier": f"AID:{child}",
                    "relationship_type": relationship,
                    "relationship_provenance": {
                        "provider": "PubChem PUG REST assay description",
                        "source_comment": comment,
                        "confidence": "source_backed",
                    },
                }
            )
    return records


def _canonicalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(row)
    aliases = {
        "CID": "cid",
        "SID": "sid",
        "Title": "canonical_name",
        "InChIKey": "inchikey",
        "CanonicalSMILES": "structure_identifier",
        "IsomericSMILES": "isomeric_structure_identifier",
    }
    for source, target in aliases.items():
        if source in normalized and target not in normalized:
            normalized[target] = normalized[source]
    if normalized.get("cid") is not None:
        normalized.setdefault("source_identifier", f"CID:{normalized['cid']}")
        normalized.setdefault("compound_identifier", f"CID:{normalized['cid']}")
        normalized.setdefault("mapping_status", "exact")
    return normalized


def _row_value(row: dict[str, Any], *names: str) -> Any:
    """Return the first source value matching a reviewed alias, case-insensitively."""

    folded = {str(key).casefold(): value for key, value in row.items()}
    for name in names:
        value = folded.get(name.casefold())
        if value not in (None, ""):
            return value
    return None


def _modality_from_source(query: ProviderMetadataQuery, *values: Any) -> str | None:
    text = " ".join(str(value or "") for value in values).casefold()
    for modality, markers in (
        ("antagonism", ("antagon", "inhibitor mode")),
        ("agonism", ("agon", "activator mode")),
        ("binding", ("binding", "bind", "displacement")),
    ):
        if any(marker in text for marker in markers):
            return modality
    return query.modality


def _assay_summary_row(
    row: dict[str, Any],
    *,
    provider: str,
    query: ProviderMetadataQuery,
) -> dict[str, Any]:
    aid = str(_row_value(row, "aid", "uid", "source_identifier") or "").removeprefix("AID:")
    if not aid.isdigit():
        raise ValueError("PubChem assay summary lacks a numeric AID")
    title = _text(_row_value(row, "title", "name", "assayname"))
    description = _text(_row_value(row, "description", "assaydescription", "summary"))
    target = _text(_row_value(row, "targetname", "target", "protein_target"))
    protein_targets = row.get("proteintargetlist")
    if target is None and isinstance(protein_targets, list):
        target_names = sorted(
            {
                str(item.get("name")).strip()
                for item in protein_targets
                if isinstance(item, dict) and str(item.get("name") or "").strip()
            }
        )
        target = "; ".join(target_names) or None
    assay_type = _text(
        _row_value(
            row,
            "assaytype",
            "assay_type",
            "activityoutcomemethod",
            "detectionmethod",
        )
    )
    activity_outcome = _text(
        _row_value(row, "activityoutcome", "activity_outcome", "activityoutcomemethod")
    )
    requested_targets = [
        value.casefold()
        for value in (
            query.biological_target,
            *query.target_identifiers,
            *query.target_synonyms,
        )
        if value and len(value.strip()) >= 3
    ]
    target_relevance_status = "unresolved"
    exclusion_reason = None
    if target is not None and requested_targets:
        observed_target = target.casefold()
        target_relevance_status = (
            "matched"
            if any(
                requested in observed_target or observed_target in requested
                for requested in requested_targets
            )
            else "unrelated"
        )
        if target_relevance_status == "unrelated":
            exclusion_reason = (
                "The source-declared assay target does not match the reviewed discovery target."
            )
    missing = [
        label
        for label, value in (
            ("assay title", title),
            ("target/mechanism", target),
            ("assay format", assay_type),
            ("activity endpoint", activity_outcome),
        )
        if value is None
    ]
    return {
        "record_kind": ProviderRecordKind.ASSAY.value,
        "source_identifier": f"AID:{aid}",
        "assay_identifier": f"AID:{aid}",
        "assay_name": title,
        "endpoint_name": activity_outcome,
        "assay_source_name": "Tox21" if provider == "tox21" else "PubChem BioAssay",
        "source_partition": provider,
        "target": target,
        "intended_target_family": target,
        "target_relevance_status": target_relevance_status,
        "exclusion_reason": exclusion_reason,
        "assay_design": description,
        "assay_format": assay_type,
        "organism": _text(_row_value(row, "organism", "taxname")),
        "pubchem_aid": aid,
        "modality": _modality_from_source(query, title, description, assay_type),
        "activity_field": activity_outcome,
        "compound_identifier_field": "PubChem CID/SID from public result table",
        "missing_fields": missing,
        "completeness_status": (
            HydrationCompleteness.PARTIAL.value if missing else HydrationCompleteness.COMPLETE.value
        ),
        "source_fields": {
            **row,
            "target_relevance_status": target_relevance_status,
            "exclusion_reason": exclusion_reason,
        },
        "provenance": [
            "NCBI PubChem BioAssay ESummary",
            *(["Tox21[SourceName] source-constrained catalogue"] if provider == "tox21" else []),
        ],
    }


def _assay_activity_row(
    row: dict[str, Any],
    *,
    aid: str,
    query: ProviderMetadataQuery,
    provider: str,
) -> dict[str, Any]:
    cid = _text(_row_value(row, "PUBCHEM_CID", "cid"))
    sid = _text(_row_value(row, "PUBCHEM_SID", "sid"))
    compound_identifier = f"CID:{cid}" if cid else f"SID:{sid}" if sid else None
    outcome = _text(
        _row_value(row, "PUBCHEM_ACTIVITY_OUTCOME", "activity outcome", "activityoutcome")
    )
    value_fields = {
        str(key): value
        for key, value in row.items()
        if value not in (None, "")
        and str(key).casefold()
        not in {
            "pubchem_result_tag",
            "pubchem_sid",
            "pubchem_cid",
            "pubchem_activity_outcome",
        }
    }
    first_value = next(iter(value_fields.values()), None)
    first_unit = next(iter(value_fields.keys()), None)
    return {
        "record_kind": ProviderRecordKind.ACTIVITY.value,
        "source_identifier": (
            f"AID:{aid}:{compound_identifier or deterministic_fingerprint(row)[:16]}"
        ),
        "assay_identifier": f"AID:{aid}",
        "compound_identifier": compound_identifier,
        "cid": cid,
        "sid": sid,
        "target": query.biological_target,
        "modality": _modality_from_source(query),
        "activity_call": outcome,
        "activity_value": _text(first_value),
        "activity_unit": _text(first_unit),
        "original_activity_fields": row,
        "caution_fields": {
            str(key): value
            for key, value in row.items()
            if any(term in str(key).casefold() for term in ("flag", "caution", "qualifier"))
        },
        "missing_fields": ([] if compound_identifier else ["stable compound identifier"]),
        "provenance": [
            "Public PubChem BioAssay result CSV",
            f"provider={provider}",
            f"AID={aid}",
        ],
    }


def _reviewed_geo_download_locator(*values: Any) -> str | None:
    for value in values:
        candidates = value if isinstance(value, list) else [value]
        for candidate in candidates:
            locator = _text(candidate)
            if locator is None:
                continue
            parsed = urlparse(locator)
            if parsed.scheme in {"https", "ftp"} and (parsed.hostname or "").lower() in {
                "ftp.ncbi.nlm.nih.gov",
                "www.ncbi.nlm.nih.gov",
            }:
                return locator
    return None


def _geo_summary_row(row: dict[str, Any], query: ProviderMetadataQuery) -> dict[str, Any]:
    uid = str(_row_value(row, "uid", "source_identifier") or "").removeprefix("GDS_UID:")
    accession = str(_row_value(row, "accession", "gse") or "").upper()
    title = _text(_row_value(row, "title"))
    summary = _text(_row_value(row, "summary")) or ""
    organism = _text(_row_value(row, "taxon", "organism"))
    biological_model = _text(_row_value(row, "gdstype", "entrytype"))
    ftp_link = _reviewed_geo_download_locator(_row_value(row, "ftplink", "ftpLink"))
    supplementary = _row_value(row, "suppfile")
    supplementary_locator = _reviewed_geo_download_locator(supplementary)
    combined = " ".join(item for item in (title, summary) if item)
    dose_match = re.search(r"\b\d+(?:\.\d+)?\s*(?:nM|uM|µM|mM|mg/?kg)\b", combined, re.I)
    time_match = re.search(r"\b\d+(?:\.\d+)?\s*(?:h|hr|hrs|hour|hours|day|days)\b", combined, re.I)
    perturbation = bool(
        re.search(r"\b(?:compound|chemical|drug|treated?|exposed?|perturb)\w*\b", combined, re.I)
    )
    transcriptomic = bool(
        re.search(
            r"\b(?:expression|transcriptom|microarray|rna[- ]?seq|profil)\w*\b", combined, re.I
        )
    )
    matrix_available = bool(ftp_link or supplementary_locator)
    if not transcriptomic:
        outcome = "not_transcriptomic"
    elif not perturbation:
        outcome = "unrelated_name_match"
    elif not matrix_available:
        outcome = "processed_matrix_missing"
    elif not organism or not biological_model:
        outcome = "biological_context_incomplete"
    else:
        outcome = "candidate_requires_manual_verification"
    missing = [
        label
        for label, value in (
            (
                "GSE accession",
                accession if re.fullmatch(r"GSE[1-9][0-9]{1,8}", accession) else None,
            ),
            ("organism", organism),
            ("cell or tissue model", biological_model),
            ("processed matrix locator", ftp_link or supplementary_locator),
        )
        if value is None or value == ""
    ]
    return {
        "record_kind": ProviderRecordKind.TRANSCRIPTOMIC.value,
        "source_identifier": accession or f"GDS_UID:{uid}",
        "accession": accession or None,
        "organism": organism,
        "biological_model": biological_model,
        "dose": dose_match.group(0) if dose_match else None,
        "exposure_time": time_match.group(0) if time_match else None,
        "processing_level": "processed public GEO resource" if matrix_available else None,
        "matrix_locator": ftp_link or supplementary_locator,
        "perturbation_verified": False,
        "outcome": outcome,
        "title": title,
        "summary": summary,
        "query_context": query.biological_target,
        "missing_fields": missing,
        "completeness_status": HydrationCompleteness.PARTIAL.value,
        "source_fields": row,
        "provenance": ["NCBI GEO DataSets ESummary", f"GDS_UID={uid}"],
    }


def _parse_geo_soft(text: str) -> dict[str, Any]:
    fields: dict[str, list[str]] = {}
    for line in text.splitlines():
        if not line.startswith("!") or " = " not in line:
            continue
        key, value = line[1:].split(" = ", 1)
        fields.setdefault(key.strip(), []).append(value.strip())
    accession = next(iter(fields.get("Series_geo_accession", [])), "")
    if not re.fullmatch(r"GSE[1-9][0-9]{1,8}", accession):
        raise ValueError("GEO SOFT response lacks a valid Series accession")
    title = next(iter(fields.get("Series_title", [])), "")
    summary = " ".join(fields.get("Series_summary", []))
    organisms = fields.get("Sample_organism_ch1", [])
    characteristics = fields.get("Sample_characteristics_ch1", [])
    treatments = fields.get("Sample_treatment_protocol_ch1", [])
    supplementary = fields.get("Series_supplementary_file", [])
    combined = " ".join([title, summary, *characteristics, *treatments])
    dose_match = re.search(r"\b\d+(?:\.\d+)?\s*(?:nM|uM|µM|mM|mg/?kg)\b", combined, re.I)
    time_match = re.search(r"\b\d+(?:\.\d+)?\s*(?:h|hr|hrs|hour|hours|day|days)\b", combined, re.I)
    perturbation = bool(
        re.search(r"\b(?:compound|chemical|drug|treated?|exposed?|perturb)\w*\b", combined, re.I)
    )
    transcriptomic = bool(
        re.search(
            r"\b(?:expression|transcriptom|microarray|rna[- ]?seq|profil)\w*\b", combined, re.I
        )
    )
    matrix_locator = _reviewed_geo_download_locator(supplementary)
    if not transcriptomic:
        outcome = "not_transcriptomic"
    elif not perturbation:
        outcome = "unrelated_name_match"
    elif not matrix_locator:
        outcome = "processed_matrix_missing"
    elif not organisms or not characteristics:
        outcome = "biological_context_incomplete"
    else:
        outcome = "verified_compound_perturbation"
    return {
        "record_kind": ProviderRecordKind.TRANSCRIPTOMIC.value,
        "source_identifier": accession,
        "accession": accession,
        "organism": next(iter(organisms), None),
        "biological_model": "; ".join(dict.fromkeys(characteristics)) or None,
        "dose": dose_match.group(0) if dose_match else None,
        "exposure_time": time_match.group(0) if time_match else None,
        "processing_level": "processed supplementary resource" if matrix_locator else None,
        "matrix_locator": matrix_locator,
        "perturbation_verified": perturbation and transcriptomic,
        "outcome": outcome,
        "title": title,
        "summary": summary,
        "sample_design": {
            "sample_count": len(fields.get("Sample_geo_accession", [])),
            "treatment_protocols": list(dict.fromkeys(treatments)),
            "characteristics": list(dict.fromkeys(characteristics)),
        },
        "supplementary_files": supplementary,
        "missing_fields": [
            label
            for label, value in (
                ("organism", organisms),
                ("cell or tissue model", characteristics),
                ("processed matrix locator", matrix_locator),
            )
            if not value
        ],
        "completeness_status": (
            HydrationCompleteness.COMPLETE.value
            if outcome == "verified_compound_perturbation"
            else HydrationCompleteness.PARTIAL.value
        ),
        "source_fields": fields,
        "provenance": ["NCBI GEO Series SOFT", f"accession={accession}"],
    }


def _compact_parser(resource: ProviderResourceRequest, content: bytes) -> dict[str, Any]:
    """Return cache-safe counts/cursors only; scientific rows remain in raw artifacts."""

    if content.startswith(b"PK\x03\x04"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet_counts: dict[str, int] = {}
        try:
            for worksheet in workbook.worksheets:
                # EPA's reviewed v4.3 workbooks retain an incorrect A1 dimension.
                # Resetting it is necessary to stream the actual immutable rows.
                worksheet.reset_dimensions()
                rows = worksheet.iter_rows(values_only=True)
                header = next(rows, None)
                sheet_counts[worksheet.title] = sum(1 for _ in rows) if header else 0
        finally:
            workbook.close()
        return {
            "record_count": sum(sheet_counts.values()),
            "sheet_record_counts": sheet_counts,
            "next_cursor": None,
            "terminal": True,
        }
    raw = gzip.decompress(content) if content.startswith(b"\x1f\x8b") else content
    text = raw.decode("utf-8-sig")
    if resource.operation_id == "geo_series_soft":
        _parse_geo_soft(text)
        return {"record_count": 1, "next_cursor": None, "terminal": True}
    stripped = text.lstrip()
    if stripped.startswith("<"):
        records = _elink_xml_records(text)
        return {
            "record_count": len(records),
            "next_cursor": None,
            "terminal": True,
        }
    if stripped.startswith("{") or stripped.startswith("["):
        payload = json.loads(text)
        records = (
            _assay_description_relationship_rows(payload)
            if resource.operation_id == "assay_relationships"
            else _json_records(payload)
        )
        next_cursor = payload.get("next_cursor") if isinstance(payload, dict) else None
        terminal = bool(
            payload.get("terminal", not next_cursor) if isinstance(payload, dict) else True
        )
        total_count = payload.get("total_count") if isinstance(payload, dict) else None
        if isinstance(payload, dict) and isinstance(payload.get("esearchresult"), dict):
            search = payload["esearchresult"]
            total_count = int(search.get("count") or len(records))
            retstart = int(search.get("retstart") or 0)
            next_start = retstart + len(records)
            terminal = next_start >= total_count or not records
            next_cursor = None if terminal else str(next_start)
        return {
            "record_count": len(records),
            "next_cursor": next_cursor,
            "terminal": terminal,
            "total_count": total_count,
        }
    dialect = csv.excel_tab if "\t" in text.partition("\n")[0] else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return {"record_count": sum(1 for _ in reader), "next_cursor": None, "terminal": True}


def _authoritative_toxcast_parser(
    resource: ProviderResourceRequest, content: bytes
) -> dict[str, Any]:
    compact = _compact_parser(resource, content)
    operation = resource.operation_id
    if operation in {
        "public_assay_annotations",
        "public_assay_target_mappings",
        "public_chemical_analytical_qc",
        "public_cytotox_reference",
    }:
        required_sheets = {
            "public_assay_annotations": "annotations_combined",
            "public_assay_target_mappings": "Sheet 1",
            "public_chemical_analytical_qc": "Sheet 1",
            "public_cytotox_reference": "Sheet 1",
        }
        sheet_counts = compact.get("sheet_record_counts") or {}
        required_sheet = required_sheets[operation]
        if int(sheet_counts.get(required_sheet, 0)) <= 0:
            raise ValueError("public ToxCast workbook lacks its reviewed populated sheet")
        return {**compact, "validated_output_schema": operation}
    payload = json.loads(content.decode("utf-8-sig"))
    if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
        raise ValueError("authoritative ToxCast operation did not return an object array")
    if operation == "list_release_datasets" and not all(
        (item.get("id") or item.get("_id")) and item.get("name") for item in payload
    ):
        raise ValueError("Clowder release listing lacks dataset identifiers or names")
    if operation == "list_dataset_files" and not all(
        (item.get("id") or item.get("_id")) and (item.get("filename") or item.get("name"))
        for item in payload
    ):
        raise ValueError("Clowder file listing lacks file identifiers or names")
    if operation == "inspect_public_archive_metadata":
        inventories = [
            item.get("content", {}).get("inventory")
            for item in payload
            if isinstance(item.get("content"), dict)
        ]
        inventory = next((item for item in inventories if isinstance(item, list)), None)
        if not inventory or not all(isinstance(item, str) for item in inventory):
            raise ValueError("Clowder archive metadata lacks a bounded member inventory")
        return {
            **compact,
            "record_count": len(inventory),
            "archive_members": inventory,
            "validated_output_schema": operation,
        }
    if operation == "all_assays" and not all(
        str(item.get("aeid") or "").isdigit()
        and (item.get("assayComponentEndpointName") or item.get("assayName"))
        for item in payload
    ):
        raise ValueError("CTX all-assays response does not match AssayAnnotation[]")
    if operation == "summary_by_aeid" and not all(
        str(item.get("aeid") or "").isdigit()
        and all(key in item for key in ("activeMc", "totalMc", "activeSc", "totalSc"))
        for item in payload
    ):
        raise ValueError("CTX summary response does not match AssayAgg[]")
    if operation == "details_by_aeid" and not all(
        str(item.get("aeid") or "").isdigit()
        and any(key in item for key in ("dtxsid", "spid", "chid"))
        and "hitc" in item
        for item in payload
    ):
        raise ValueError("CTX details response does not match BioactivityDataAll[]")
    return {**compact, "validated_output_schema": operation}


def _toxcast_source_partition(row: dict[str, Any]) -> str:
    source_name = _text(row.get("assaySourceName"))
    source_id = _text(row.get("asid"))
    source_text = " ".join(item for item in (source_name, source_id) if item).casefold()
    if "tox21" in source_text:
        return "tox21"
    if source_text:
        return "toxcast_source_backed"
    return "unclassified_missing_source"


def _toxcast_modality(row: dict[str, Any]) -> str | None:
    text = " ".join(
        str(row.get(field) or "")
        for field in (
            "assayComponentEndpointName",
            "assayComponentEndpointDesc",
            "assayFunctionType",
            "assayDesignType",
            "assayDesignTypeSub",
            "assayName",
            "assayDesc",
        )
    ).casefold()
    for modality, terms in (
        ("antagonism", ("antagonist", "antagonism")),
        ("agonism", ("agonist", "agonism")),
        ("binding", ("binding", "bind", "ligand displacement")),
    ):
        if any(term in text for term in terms):
            return modality
    return None


def _toxcast_gene_text(value: Any) -> str | None:
    values = value if isinstance(value, list) else [value]
    labels: list[str] = []
    for item in values:
        if not isinstance(item, dict):
            continue
        for field in (
            "officialSymbol",
            "geneSymbol",
            "officialFullName",
            "geneName",
            "uniprotAccessionNumber",
            "entrezGeneId",
        ):
            if label := _text(item.get(field)):
                labels.append(label)
    return "; ".join(dict.fromkeys(labels)) or None


def _toxcast_assay_row(row: dict[str, Any]) -> dict[str, Any]:
    aeid = _text(row.get("aeid"))
    if aeid is None or not aeid.isdigit():
        raise ValueError("CTX Bioactivity assay annotation lacks a numeric AEID")
    target = _toxcast_gene_text(row.get("gene")) or _text(row.get("assayComponentTargetDesc"))
    source = _text(row.get("assaySourceName"))
    return {
        "record_kind": ProviderRecordKind.ASSAY.value,
        "source_identifier": f"AEID:{aeid}",
        "assay_identifier": f"AEID:{aeid}",
        "assay_name": _text(row.get("assayName")),
        "component_name": _text(row.get("assayComponentName")),
        "endpoint_name": _text(row.get("assayComponentEndpointName")),
        "assay_source_name": source,
        "source_partition": _toxcast_source_partition(row),
        "target": target,
        "intended_target_type": _text(row.get("intendedTargetType")),
        "intended_target_family": _text(row.get("intendedTargetFamily")),
        "biological_process": _text(row.get("biologicalProcessTarget")),
        "assay_design": _text(row.get("assayDesignType")),
        "assay_format": _text(row.get("assayFormatType")),
        "signal_direction": _text(row.get("signalDirection")),
        "organism": _text(row.get("organism")) or _text(row.get("taxonName")),
        "tissue": _text(row.get("tissue")),
        "cell_model": _text(row.get("cellShortName")),
        "timepoint_hours": _text(row.get("timepointHr")),
        "viability_annotation": _text(row.get("cellViabilityAssay")),
        "pubchem_aid": _text(row.get("aid")),
        "modality": _toxcast_modality(row),
        "citations": row.get("citations") or [],
        "source_fields": row,
        "provenance": [
            "EPA CTX Bioactivity API assay-all projection",
            *([f"assaySourceName={source}"] if source else []),
        ],
    }


def _toxcast_relationship_rows(row: dict[str, Any]) -> list[dict[str, Any]]:
    aeid = _text(row.get("aeid"))
    if aeid is None:
        return []
    endpoint = f"AEID:{aeid}"
    relationships: list[dict[str, Any]] = []

    def add(subject: str, relationship: str, object_identifier: str, fields: list[str]) -> None:
        relationships.append(
            {
                "record_kind": ProviderRecordKind.RELATIONSHIP.value,
                "source_identifier": (f"{subject}:{relationship}:{object_identifier}"),
                "parent_assay_identifier": subject,
                "child_assay_identifier": object_identifier,
                "relationship_type": relationship,
                "relationship_provenance": {
                    "provider": "toxcast",
                    "exact_source_fields": fields,
                    "confidence": "source_backed",
                },
                "provenance": ["EPA CTX Bioactivity API assay annotation"],
            }
        )

    acid = _text(row.get("acid"))
    asid = _text(row.get("asid"))
    aid = _text(row.get("aid"))
    if asid and acid:
        add(f"ASID:{asid}", "assay_has_component", f"ACID:{acid}", ["asid", "acid"])
    if acid:
        add(f"ACID:{acid}", "component_has_endpoint", endpoint, ["acid", "aeid"])
    if aid:
        add(endpoint, "linked_to_pubchem_assay", f"AID:{aid}", ["aeid", "aid"])
    if row.get("cellViabilityAssay") is not None:
        add(
            endpoint,
            "has_viability_annotation",
            f"VIABILITY:{row.get('cellViabilityAssay')}",
            ["aeid", "cellViabilityAssay"],
        )
    return relationships


def _toxcast_summary_row(row: dict[str, Any]) -> dict[str, Any]:
    aeid = _text(row.get("aeid"))
    if aeid is None or not aeid.isdigit():
        raise ValueError("CTX Bioactivity summary lacks a numeric AEID")
    return {
        "record_kind": ProviderRecordKind.ACTIVITY_SUMMARY.value,
        "source_identifier": f"AEID:{aeid}",
        "assay_identifier": f"AEID:{aeid}",
        "aeid": aeid,
        "active_mc": row.get("activeMc"),
        "total_mc": row.get("totalMc"),
        "active_sc": row.get("activeSc"),
        "total_sc": row.get("totalSc"),
        "source_fields": row,
        "provenance": ["EPA CTX Bioactivity API summary-by-AEID"],
    }


def _toxcast_detail_row(
    row: dict[str, Any], assay_by_aeid: dict[str, dict[str, Any]]
) -> dict[str, Any]:
    aeid = _text(row.get("aeid"))
    if aeid is None or not aeid.isdigit():
        raise ValueError("CTX Bioactivity detail lacks a numeric AEID")
    dtxsid = _text(row.get("dtxsid"))
    spid = _text(row.get("spid"))
    chid = _text(row.get("chid"))
    compound_identifier = (
        f"DTXSID:{dtxsid}"
        if dtxsid
        else f"SPID:{spid}"
        if spid
        else f"CHID:{chid}"
        if chid
        else None
    )
    assay = assay_by_aeid.get(aeid, {})
    return {
        "record_kind": ProviderRecordKind.ACTIVITY.value,
        "source_identifier": (
            f"AEID:{aeid}:{compound_identifier or deterministic_fingerprint(row)[:16]}"
        ),
        "assay_identifier": f"AEID:{aeid}",
        "compound_identifier": compound_identifier,
        "dtxsid": dtxsid,
        "spid": spid,
        "chid": chid,
        "canonical_name": _text(row.get("chnm")),
        "casrn": _text(row.get("casn")),
        "target": _text(assay.get("target")),
        "modality": _text(assay.get("modality")),
        "activity_call": _text(row.get("hitc")),
        "activity_value": _text(row.get("actp")),
        "activity_unit": _text(row.get("testedConcUnit")),
        "original_activity_fields": row,
        "caution_fields": {
            key: row.get(key) for key in ("fitc", "coff", "nrep", "nconc", "npts") if key in row
        },
        "source_fields": row,
        "provenance": ["EPA CTX Bioactivity API complete details-by-AEID"],
    }


def _toxcast_public_targets_by_aeid(
    rows: Iterator[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    targets: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        aeid = _text(row.get("aeid"))
        if aeid is None or not aeid.isdigit():
            continue
        target = {
            "target_id": _text(row.get("target_id")),
            "target_type": _text(row.get("target_type")),
            "official_full_name": _text(row.get("official_full_name")),
            "official_symbol": _text(row.get("official_symbol")),
            "ncbi_taxon_id": _text(row.get("ncbi_taxon_id")),
        }
        targets.setdefault(aeid, []).append(target)
    return targets


def _toxcast_public_assay_row(
    row: dict[str, Any], targets_by_aeid: dict[str, list[dict[str, Any]]]
) -> dict[str, Any]:
    aeid = _text(row.get("aeid"))
    if aeid is None or not aeid.isdigit():
        raise ValueError("public invitroDB assay annotation lacks a numeric AEID")
    gene = [
        {
            "officialSymbol": item.get("official_symbol"),
            "officialFullName": item.get("official_full_name"),
            "entrezGeneId": item.get("target_id"),
        }
        for item in targets_by_aeid.get(aeid, [])
        if item.get("target_type") == "entrez_gene_id"
        and any(
            item.get(field) not in {None, "#N/A"}
            for field in ("official_symbol", "official_full_name")
        )
    ]
    projected = {
        "aid": row.get("aid"),
        "asid": row.get("asid"),
        "acid": row.get("acid"),
        "aeid": aeid,
        "assayName": row.get("assay_name"),
        "assayDesc": row.get("assay_desc"),
        "timepointHr": row.get("timepoint_hr"),
        "organism": row.get("organism"),
        "tissue": row.get("tissue"),
        "cellShortName": row.get("cell_short_name"),
        "assayFormatType": row.get("assay_format_type"),
        "assaySourceName": row.get("assay_source_name"),
        "assayComponentName": row.get("assay_component_name"),
        "assayComponentDesc": row.get("assay_component_desc"),
        "assayComponentTargetDesc": row.get("assay_component_target_desc"),
        "assayDesignType": row.get("assay_design_type"),
        "assayDesignTypeSub": row.get("assay_design_type_sub"),
        "biologicalProcessTarget": row.get("biological_process_target"),
        "assayComponentEndpointName": row.get("assay_component_endpoint_name"),
        "assayComponentEndpointDesc": row.get("assay_component_endpoint_desc"),
        "assayFunctionType": row.get("assay_function_type"),
        "signalDirection": row.get("signal_direction"),
        "intendedTargetType": row.get("intended_target_type"),
        "intendedTargetFamily": row.get("intended_target_family"),
        "cellViabilityAssay": row.get("cell_viability_assay"),
        "gene": gene,
    }
    assay = _toxcast_assay_row(projected)
    return {
        **assay,
        "data_usability": row.get("data_usability"),
        "source_fields": row,
        "provenance": [
            "EPA invitroDB v4.3 public assay_annotations workbook",
            f"AEID={aeid}",
        ],
    }


def _toxcast_public_target_relationship_rows(row: dict[str, Any]) -> list[dict[str, Any]]:
    aeid = _text(row.get("aeid"))
    target_type = _text(row.get("target_type"))
    target_id = _text(row.get("target_id"))
    if aeid is None or not aeid.isdigit() or not target_type or not target_id:
        return []
    target_label = (
        _text(row.get("official_symbol")) or _text(row.get("official_full_name")) or target_id
    )
    return [
        {
            "record_kind": ProviderRecordKind.RELATIONSHIP.value,
            "source_identifier": f"AEID:{aeid}:maps-to:{target_type}:{target_id}",
            "parent_assay_identifier": f"AEID:{aeid}",
            "child_assay_identifier": f"TARGET:{target_type}:{target_id}",
            "relationship_type": "assay_endpoint_maps_to_reviewed_target",
            "relationship_provenance": {
                "provider": "toxcast",
                "exact_source_fields": sorted(row),
                "target_label": target_label,
                "confidence": "source_backed",
            },
            "provenance": ["EPA invitroDB v4.3 public assay_target_mappings workbook"],
        }
    ]


def _toxcast_public_compound_row(row: dict[str, Any], *, operation_id: str) -> dict[str, Any]:
    dtxsid = _text(row.get("dsstox_substance_id"))
    spid = _text(row.get("spid"))
    chid = _text(row.get("chid"))
    compound_identifier = (
        f"DTXSID:{dtxsid}"
        if dtxsid
        else f"SPID:{spid}"
        if spid
        else f"CHID:{chid}"
        if chid
        else None
    )
    source_key = (
        _text(row.get("analytical_qc_id"))
        or compound_identifier
        or deterministic_fingerprint(row)[:16]
    )
    source_name = (
        "chemical_analytical_qc"
        if operation_id == "public_chemical_analytical_qc"
        else "cytotox_reference"
    )
    return {
        "record_kind": ProviderRecordKind.COMPOUND.value,
        "source_identifier": f"{source_name}:{source_key}",
        "compound_identifier": compound_identifier,
        "dtxsid": dtxsid,
        "spid": spid,
        "chid": chid,
        "canonical_name": _text(row.get("chnm")),
        "casrn": _text(row.get("casn")),
        "mapping_status": "exact" if compound_identifier else "missing",
        "structure_status": "not_exposed_by_public_reference_table",
        "quality_fields": row,
        "source_fields": row,
        "provenance": [f"EPA invitroDB v4.3 public {source_name} workbook"],
    }


def _toxcast_rows_for_operation(
    operation_id: str | None,
    row: dict[str, Any],
    assay_by_aeid: dict[str, dict[str, Any]],
    public_targets_by_aeid: dict[str, list[dict[str, Any]]] | None = None,
) -> list[dict[str, Any]]:
    if operation_id == "all_assays":
        assay = _toxcast_assay_row(row)
        return [assay, *_toxcast_relationship_rows(row)]
    if operation_id == "summary_by_aeid":
        return [_toxcast_summary_row(row)]
    if operation_id == "details_by_aeid":
        return [_toxcast_detail_row(row, assay_by_aeid)]
    if operation_id == "public_assay_annotations":
        assay = _toxcast_public_assay_row(row, public_targets_by_aeid or {})
        projected = dict(assay.get("source_fields") or {})
        projected.update(
            {
                "aeid": row.get("aeid"),
                "asid": row.get("asid"),
                "acid": row.get("acid"),
                "aid": row.get("aid"),
                "cellViabilityAssay": row.get("cell_viability_assay"),
            }
        )
        return [assay, *_toxcast_relationship_rows(projected)]
    if operation_id == "public_assay_target_mappings":
        return _toxcast_public_target_relationship_rows(row)
    if operation_id in {"public_chemical_analytical_qc", "public_cytotox_reference"}:
        return [_toxcast_public_compound_row(row, operation_id=operation_id)]
    return [dict(row)]


def _iter_source_records(path: Path, operation_id: str | None = None) -> Iterator[dict[str, Any]]:
    with path.open("rb") as probe:
        magic = probe.read(4)
    if magic == b"PK\x03\x04":
        with path.open("rb") as workbook_handle:
            workbook = load_workbook(workbook_handle, read_only=True, data_only=True)
            selected_sheet = (
                {
                    "public_assay_annotations": "annotations_combined",
                    "public_assay_target_mappings": "Sheet 1",
                    "public_chemical_analytical_qc": "Sheet 1",
                    "public_cytotox_reference": "Sheet 1",
                }.get(operation_id)
                if operation_id is not None
                else None
            )
            try:
                worksheets = (
                    [workbook[selected_sheet]] if selected_sheet else list(workbook.worksheets)
                )
                for worksheet in worksheets:
                    worksheet.reset_dimensions()
                    rows = worksheet.iter_rows(values_only=True)
                    header_row = next(rows, None)
                    if header_row is None:
                        continue
                    headers = [
                        str(value).strip() if value is not None else "" for value in header_row
                    ]
                    if not headers or any(not value for value in headers):
                        raise ValueError("public ToxCast workbook contains a blank header")
                    for values in rows:
                        yield _canonicalize_row(
                            {header: value for header, value in zip(headers, values, strict=False)}
                        )
            finally:
                workbook.close()
        return
    with path.open("rb") as probe:
        compressed = probe.read(2) == b"\x1f\x8b"
    opener = gzip.open if compressed else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        if operation_id == "geo_series_soft":
            yield _parse_geo_soft(handle.read())
            return
        first = handle.read(1)
        handle.seek(0)
        if first == "<":
            yield from (_canonicalize_row(item) for item in _elink_xml_records(handle.read()))
            return
        if first in {"{", "["}:
            payload = json.load(handle)
            records = (
                _assay_description_relationship_rows(payload)
                if operation_id == "assay_relationships"
                else _json_records(payload)
            )
            yield from (_canonicalize_row(item) for item in records)
            return
        header = handle.readline()
        handle.seek(0)
        dialect = csv.excel_tab if "\t" in header else csv.excel
        yield from (
            _canonicalize_row(dict(item)) for item in csv.DictReader(handle, dialect=dialect)
        )


def _provider_rows_for_operation(
    *,
    provider: str,
    operation_id: str | None,
    resource_id: str,
    kind: ProviderRecordKind,
    source_row: dict[str, Any],
    query: ProviderMetadataQuery,
) -> list[dict[str, Any]]:
    """Project reviewed provider responses without changing their scientific values."""

    if operation_id == "assay_summary_batch":
        return [_assay_summary_row(source_row, provider=provider, query=query)]
    if operation_id == "assay_compound_identifiers":
        aid_match = re.search(r"aid-([0-9]+)", resource_id)
        if aid_match is None:
            raise ValueError("compound-identifier response lacks its bound assay identifier")
        cid = _text(source_row.get("cid"))
        if cid is None or not cid.isdigit():
            return []
        return [
            {
                "record_kind": ProviderRecordKind.COMPOUND.value,
                "source_identifier": f"CID:{cid}",
                "compound_identifier": f"CID:{cid}",
                "cid": cid,
                "mapping_status": "source_backed",
                "structure_status": "not_requested",
                "provenance": [
                    "PubChem BioAssay active-CID index",
                    f"AID={aid_match.group(1)}",
                ],
                "source_fields": source_row,
            }
        ]
    if operation_id == "assay_activity_rows":
        aid_match = re.search(r"aid-([0-9]+)", resource_id)
        if aid_match is None:
            raise ValueError("activity result response lacks its bound assay identifier")
        return [
            _assay_activity_row(
                source_row,
                aid=aid_match.group(1),
                query=query,
                provider=provider,
            )
        ]
    if operation_id == "assay_relationships":
        return [
            {
                "record_kind": ProviderRecordKind.RELATIONSHIP.value,
                **source_row,
                "provenance": ["NCBI PubChem assay-to-assay ELink"],
            }
        ]
    if operation_id == "geo_summary_batch":
        return [_geo_summary_row(source_row, query)]
    if operation_id == "geo_series_soft":
        return [source_row]
    if operation_id == "search_candidates":
        identifier = _text(source_row.get("source_identifier")) or "unresolved"
        prefix = "GDS_UID" if provider == "ncbi-geo" else "AID"
        return [
            {
                "record_kind": kind.value,
                "source_identifier": f"{prefix}:{identifier}",
                "assay_identifier": (
                    f"AID:{identifier}" if kind is ProviderRecordKind.ASSAY else None
                ),
                "modality": query.modality,
                "missing_fields": (
                    [
                        "assay title",
                        "target/mechanism",
                        "activity result availability",
                    ]
                    if kind is ProviderRecordKind.ASSAY
                    else ["GSE accession", "study design", "processed matrix availability"]
                ),
                "completeness_status": HydrationCompleteness.PARTIAL.value,
                "provenance": ["Reviewed provider candidate search"],
                "source_fields": source_row,
            }
        ]
    return [source_row]


def _normalized_search_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def _toxcast_matched_aeids(
    rows: list[dict[str, Any]], query: ProviderMetadataQuery
) -> tuple[list[int], dict[str, dict[str, Any]]]:
    target_terms = [
        item
        for item in (
            *query.target_identifiers,
            *query.target_synonyms,
            *([query.biological_target] if query.biological_target else []),
        )
        if _normalized_search_text(item)
    ]
    normalized_terms = [_normalized_search_text(item) for item in target_terms]
    matched: list[int] = []
    assay_by_aeid: dict[str, dict[str, Any]] = {}
    for raw in rows:
        assay = _toxcast_assay_row(raw)
        aeid = assay["assay_identifier"].removeprefix("AEID:")
        assay_by_aeid[aeid] = assay
        haystack = _normalized_search_text(
            " ".join(
                str(raw.get(field) or "")
                for field in (
                    "assayComponentEndpointName",
                    "assayComponentEndpointDesc",
                    "assayComponentTargetDesc",
                    "intendedTargetType",
                    "intendedTargetFamily",
                    "intendedTargetFamilySub",
                    "biologicalProcessTarget",
                    "assayName",
                    "assayDesc",
                    "assayComponentName",
                    "assayComponentDesc",
                )
            )
            + " "
            + (_toxcast_gene_text(raw.get("gene")) or "")
        )
        target_match = not normalized_terms or any(
            term in haystack or all(word in haystack.split() for word in term.split())
            for term in normalized_terms
        )
        modality_match = query.modality is None or assay.get("modality") == query.modality
        belongs_to_toxcast_view = assay.get("source_partition") != "tox21"
        if target_match and modality_match and belongs_to_toxcast_view:
            matched.append(int(aeid))
    return sorted(set(matched)), assay_by_aeid


def _text(value: Any) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def _truth(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "verified", "available"}


def _json(value: Any) -> str:
    def encode(source_value: Any) -> str:
        if isinstance(source_value, date | datetime):
            return source_value.isoformat()
        raise TypeError(f"Unsupported provider source value: {type(source_value).__name__}")

    return json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
        default=encode,
    )


def _source_identifier(row: dict[str, Any], kind: ProviderRecordKind) -> str:
    fields = {
        ProviderRecordKind.ASSAY: ("source_identifier", "assay_identifier", "aid", "assay_id"),
        ProviderRecordKind.ACTIVITY: (
            "source_identifier",
            "activity_record_id",
            "record_id",
            "assay_identifier",
        ),
        ProviderRecordKind.ACTIVITY_SUMMARY: (
            "source_identifier",
            "assay_identifier",
            "aeid",
        ),
        ProviderRecordKind.COMPOUND: (
            "source_identifier",
            "cid",
            "inchikey",
            "sid",
            "compound_identifier",
        ),
        ProviderRecordKind.RELATIONSHIP: (
            "source_identifier",
            "relationship_id",
            "parent_assay_identifier",
        ),
        ProviderRecordKind.TRANSCRIPTOMIC: (
            "source_identifier",
            "accession",
            "profile_identifier",
            "record_id",
        ),
        ProviderRecordKind.PUBLICATION: (
            "source_identifier",
            "publication_identifier",
            "pmid",
            "doi",
        ),
        ProviderRecordKind.DATASET_LINK: (
            "source_identifier",
            "dataset_identifier",
            "accession",
        ),
        ProviderRecordKind.RELEASE: ("source_identifier", "release_id", "version"),
    }
    for field in fields[kind]:
        if value := _text(row.get(field)):
            return value
    return f"record-{deterministic_fingerprint(row)[:24]}"


def _create_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        PRAGMA journal_mode=OFF;
        PRAGMA synchronous=OFF;
        CREATE TABLE raw_records (
            record_id TEXT PRIMARY KEY, provider TEXT NOT NULL, release_id TEXT NOT NULL,
            resource_id TEXT NOT NULL, record_kind TEXT NOT NULL,
            source_identifier TEXT NOT NULL, payload_json TEXT NOT NULL,
            source_artifact_id TEXT NOT NULL, source_sha256 TEXT NOT NULL
        );
        CREATE TABLE source_candidates (
            candidate_id TEXT PRIMARY KEY, source_identifier TEXT NOT NULL,
            evidence_role TEXT NOT NULL, modality TEXT, task_id TEXT NOT NULL,
            candidate_status TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE hydrated_sources (
            hydrated_source_id TEXT PRIMARY KEY, candidate_id TEXT NOT NULL,
            source_identifier TEXT NOT NULL, evidence_role TEXT NOT NULL, modality TEXT,
            verified_metadata_json TEXT NOT NULL, compound_index_available INTEGER NOT NULL,
            label_fields_available INTEGER NOT NULL, organism TEXT,
            experimental_context_json TEXT NOT NULL, completeness_status TEXT NOT NULL,
            missing_fields_json TEXT NOT NULL, exclusion_reason TEXT, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE compound_index (
            mapping_id TEXT PRIMARY KEY, source_identifier TEXT NOT NULL,
            compound_identifier TEXT, cid TEXT, sid TEXT, inchikey TEXT,
            canonical_name TEXT, exact_synonyms_json TEXT NOT NULL,
            mapping_status TEXT NOT NULL, structure_status TEXT,
            parent_identifier TEXT, provenance_json TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE activity_records (
            activity_record_id TEXT PRIMARY KEY, source_identifier TEXT NOT NULL,
            assay_identifier TEXT, compound_identifier TEXT, cid TEXT, target TEXT,
            modality TEXT, activity_call TEXT, activity_value TEXT, activity_unit TEXT,
            original_activity_fields_json TEXT NOT NULL, caution_fields_json TEXT NOT NULL,
            provenance_json TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE activity_summaries (
            summary_id TEXT PRIMARY KEY, aeid TEXT NOT NULL,
            active_mc INTEGER, total_mc INTEGER, active_sc INTEGER, total_sc INTEGER,
            provenance_json TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE assay_annotations (
            annotation_id TEXT PRIMARY KEY, aeid TEXT NOT NULL,
            assay_name TEXT, component_name TEXT, endpoint_name TEXT,
            assay_source_name TEXT, source_partition TEXT NOT NULL,
            intended_target_type TEXT, intended_target_family TEXT,
            biological_process TEXT, assay_design TEXT, assay_format TEXT,
            signal_direction TEXT, organism TEXT, tissue TEXT, cell_model TEXT,
            timepoint_hours TEXT, viability_annotation TEXT, pubchem_aid TEXT,
            modality TEXT, source_fields_json TEXT NOT NULL,
            provenance_json TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE transcriptomic_records (
            transcriptomic_record_id TEXT PRIMARY KEY, source_identifier TEXT NOT NULL,
            accession TEXT, compound_identifier TEXT, profile_identifier TEXT,
            organism TEXT, biological_model TEXT, dose TEXT, exposure_time TEXT,
            processing_level TEXT, matrix_locator TEXT, perturbation_verified INTEGER NOT NULL,
            outcome TEXT NOT NULL, provenance_json TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE assay_relationships (
            relationship_id TEXT PRIMARY KEY, source_identifier TEXT NOT NULL,
            parent_assay_identifier TEXT, child_assay_identifier TEXT,
            relationship_type TEXT, relationship_provenance_json TEXT NOT NULL,
            raw_record_id TEXT NOT NULL
        );
        CREATE TABLE assay_traversal_ledger (
            traversal_id TEXT PRIMARY KEY, assay_identifier TEXT NOT NULL,
            discovered_from TEXT, relationship_type TEXT NOT NULL,
            visited INTEGER NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE publication_dataset_links (
            link_id TEXT PRIMARY KEY, source_identifier TEXT NOT NULL,
            publication_identifier TEXT, metadata_field_or_relation TEXT,
            dataset_identifier TEXT, target_provider TEXT, verification_status TEXT,
            provenance_json TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE TABLE release_inventory (
            inventory_id TEXT PRIMARY KEY, descriptor_type TEXT NOT NULL,
            provider_object_id TEXT, name TEXT, content_type TEXT, response_bytes INTEGER,
            source_checksum TEXT, source_fields_json TEXT NOT NULL,
            provenance_json TEXT NOT NULL, raw_record_id TEXT NOT NULL
        );
        CREATE INDEX idx_raw_kind_source ON raw_records(record_kind, source_identifier);
        CREATE INDEX idx_candidates_source ON source_candidates(source_identifier, modality);
        CREATE INDEX idx_hydrated_source ON hydrated_sources(source_identifier, modality);
        CREATE INDEX idx_compound_cid ON compound_index(cid);
        CREATE INDEX idx_compound_inchikey ON compound_index(inchikey);
        CREATE INDEX idx_compound_sid ON compound_index(sid);
        CREATE INDEX idx_activity_assay_compound
            ON activity_records(assay_identifier, compound_identifier);
        CREATE INDEX idx_activity_modality ON activity_records(modality, target);
        CREATE INDEX idx_activity_summary_aeid ON activity_summaries(aeid);
        CREATE INDEX idx_assay_annotation_aeid ON assay_annotations(aeid);
        CREATE INDEX idx_assay_source_partition ON assay_annotations(source_partition);
        CREATE INDEX idx_transcriptomic_compound ON transcriptomic_records(compound_identifier);
        CREATE INDEX idx_transcriptomic_accession ON transcriptomic_records(accession);
        CREATE INDEX idx_relationship_parent ON assay_relationships(parent_assay_identifier);
        CREATE INDEX idx_traversal_assay ON assay_traversal_ledger(assay_identifier);
        CREATE INDEX idx_dataset_identifier ON publication_dataset_links(dataset_identifier);
        CREATE INDEX idx_release_inventory_object ON release_inventory(provider_object_id);
        """
    )


def _candidate_role(kind: ProviderRecordKind) -> EvidenceRole | None:
    if kind in {ProviderRecordKind.ASSAY, ProviderRecordKind.ACTIVITY}:
        return EvidenceRole.ACTIVITY
    if kind is ProviderRecordKind.TRANSCRIPTOMIC:
        return EvidenceRole.TRANSCRIPTOMIC
    if kind is ProviderRecordKind.COMPOUND:
        return EvidenceRole.IDENTITY
    if kind in {ProviderRecordKind.PUBLICATION, ProviderRecordKind.DATASET_LINK}:
        return EvidenceRole.SUPPORTING_METADATA
    return None


def _insert_candidate_and_hydration(
    connection: sqlite3.Connection,
    *,
    provider: str,
    source_identifier: str,
    kind: ProviderRecordKind,
    row: dict[str, Any],
    task_id: str,
    raw_record_id: str,
) -> None:
    role = _candidate_role(kind)
    if role is None:
        return
    modality = _text(row.get("modality"))
    candidate_id = (
        f"candidate-{deterministic_fingerprint([provider, role, source_identifier, modality])[:32]}"
    )
    connection.execute(
        "INSERT OR IGNORE INTO source_candidates VALUES (?,?,?,?,?,?,?)",
        (
            candidate_id,
            source_identifier,
            role.value,
            modality,
            task_id,
            CandidateStatus.METADATA_CANDIDATE.value,
            raw_record_id,
        ),
    )
    compound_available = any(
        _text(row.get(field))
        for field in ("compound_identifier", "cid", "inchikey", "compound_identifier_field")
    )
    label_available = any(
        row.get(field) is not None
        for field in ("activity_call", "activity_value", "activity_field")
    )
    missing = row.get("missing_fields") or []
    if isinstance(missing, str):
        missing = [missing]
    completeness = _text(row.get("completeness_status")) or (
        HydrationCompleteness.PARTIAL.value if missing else HydrationCompleteness.COMPLETE.value
    )
    hydrated_id = f"hydrated-{deterministic_fingerprint([candidate_id, raw_record_id])[:32]}"
    connection.execute(
        "INSERT OR IGNORE INTO hydrated_sources VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            hydrated_id,
            candidate_id,
            source_identifier,
            role.value,
            modality,
            _json(row),
            int(compound_available),
            int(label_available),
            _text(row.get("organism")),
            _json(row.get("experimental_context") or {}),
            completeness,
            _json(missing),
            _text(row.get("exclusion_reason")),
            raw_record_id,
        ),
    )


def _insert_kind_record(
    connection: sqlite3.Connection,
    *,
    kind: ProviderRecordKind,
    source_identifier: str,
    row: dict[str, Any],
    raw_record_id: str,
) -> None:
    provenance = row.get("provenance") or []
    if kind is ProviderRecordKind.ASSAY:
        assay_identifier = _text(row.get("assay_identifier")) or source_identifier
        traversal_id = (
            f"traversal-{deterministic_fingerprint([assay_identifier, raw_record_id])[:32]}"
        )
        connection.execute(
            "INSERT OR IGNORE INTO assay_traversal_ledger VALUES (?,?,?,?,?,?)",
            (
                traversal_id,
                assay_identifier,
                None,
                "root_candidate",
                1,
                raw_record_id,
            ),
        )
        annotation_id = (
            f"annotation-{deterministic_fingerprint([assay_identifier, raw_record_id])[:32]}"
        )
        connection.execute(
            """
            INSERT OR IGNORE INTO assay_annotations VALUES (
                ?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?
            )
            """,
            (
                annotation_id,
                assay_identifier.removeprefix("AEID:"),
                _text(row.get("assay_name")),
                _text(row.get("component_name")),
                _text(row.get("endpoint_name")),
                _text(row.get("assay_source_name")),
                _text(row.get("source_partition")) or "unclassified_missing_source",
                _text(row.get("intended_target_type")),
                _text(row.get("intended_target_family")),
                _text(row.get("biological_process")),
                _text(row.get("assay_design")),
                _text(row.get("assay_format")),
                _text(row.get("signal_direction")),
                _text(row.get("organism")),
                _text(row.get("tissue")),
                _text(row.get("cell_model")),
                _text(row.get("timepoint_hours")),
                _text(row.get("viability_annotation")),
                _text(row.get("pubchem_aid")),
                _text(row.get("modality")),
                _json(row.get("source_fields") or row),
                _json(provenance),
                raw_record_id,
            ),
        )
    elif kind is ProviderRecordKind.COMPOUND:
        mapping_status = _text(row.get("mapping_status")) or "missing"
        mapping_id = f"mapping-{deterministic_fingerprint([source_identifier, row])[:32]}"
        connection.execute(
            "INSERT OR IGNORE INTO compound_index VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                mapping_id,
                source_identifier,
                _text(row.get("compound_identifier")),
                _text(row.get("cid")),
                _text(row.get("sid")),
                _text(row.get("inchikey")),
                _text(row.get("canonical_name")),
                _json(row.get("exact_synonyms") or []),
                mapping_status,
                _text(row.get("structure_status")),
                _text(row.get("parent_identifier")),
                _json(provenance),
                raw_record_id,
            ),
        )
    elif kind is ProviderRecordKind.ACTIVITY:
        activity_id = f"activity-{deterministic_fingerprint([source_identifier, row])[:32]}"
        original = row.get("original_activity_fields") or {
            key: row.get(key)
            for key in ("activity_call", "activity_value", "activity_unit")
            if key in row
        }
        connection.execute(
            "INSERT OR IGNORE INTO activity_records VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                activity_id,
                source_identifier,
                _text(row.get("assay_identifier")),
                _text(row.get("compound_identifier")),
                _text(row.get("cid")),
                _text(row.get("target")),
                _text(row.get("modality")),
                _text(row.get("activity_call")),
                _text(row.get("activity_value")),
                _text(row.get("activity_unit")),
                _json(original),
                _json(row.get("caution_fields") or {}),
                _json(provenance),
                raw_record_id,
            ),
        )
        compound_identifier = _text(row.get("compound_identifier"))
        if compound_identifier is not None:
            mapping_id = (
                f"mapping-{deterministic_fingerprint([compound_identifier, raw_record_id])[:32]}"
            )
            connection.execute(
                "INSERT OR IGNORE INTO compound_index VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    mapping_id,
                    compound_identifier,
                    compound_identifier,
                    None,
                    _text(row.get("spid")),
                    None,
                    _text(row.get("canonical_name")),
                    "[]",
                    "source_backed",
                    None,
                    None,
                    _json(provenance),
                    raw_record_id,
                ),
            )
    elif kind is ProviderRecordKind.ACTIVITY_SUMMARY:
        aeid = _text(row.get("aeid")) or source_identifier.removeprefix("AEID:")
        summary_id = f"summary-{deterministic_fingerprint([aeid, row])[:32]}"
        connection.execute(
            "INSERT OR IGNORE INTO activity_summaries VALUES (?,?,?,?,?,?,?,?)",
            (
                summary_id,
                aeid,
                row.get("active_mc"),
                row.get("total_mc"),
                row.get("active_sc"),
                row.get("total_sc"),
                _json(provenance),
                raw_record_id,
            ),
        )
    elif kind is ProviderRecordKind.TRANSCRIPTOMIC:
        outcome = _text(row.get("outcome")) or "candidate_requires_manual_verification"
        transcriptomic_id = (
            f"transcriptomic-{deterministic_fingerprint([source_identifier, row])[:32]}"
        )
        connection.execute(
            "INSERT OR IGNORE INTO transcriptomic_records VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                transcriptomic_id,
                source_identifier,
                _text(row.get("accession")),
                _text(row.get("compound_identifier")),
                _text(row.get("profile_identifier")),
                _text(row.get("organism")),
                _text(row.get("biological_model")),
                _text(row.get("dose")),
                _text(row.get("exposure_time")),
                _text(row.get("processing_level")),
                _text(row.get("matrix_locator")),
                int(_truth(row.get("perturbation_verified"))),
                outcome,
                _json(provenance),
                raw_record_id,
            ),
        )
    elif kind is ProviderRecordKind.RELATIONSHIP:
        relationship_id = f"relationship-{deterministic_fingerprint([source_identifier, row])[:32]}"
        connection.execute(
            "INSERT OR IGNORE INTO assay_relationships VALUES (?,?,?,?,?,?,?)",
            (
                relationship_id,
                source_identifier,
                _text(row.get("parent_assay_identifier")),
                _text(row.get("child_assay_identifier")),
                _text(row.get("relationship_type")),
                _json(row.get("relationship_provenance") or provenance),
                raw_record_id,
            ),
        )
        child = _text(row.get("child_assay_identifier"))
        if child is not None:
            traversal_id = f"traversal-{deterministic_fingerprint([child, relationship_id])[:32]}"
            connection.execute(
                "INSERT OR IGNORE INTO assay_traversal_ledger VALUES (?,?,?,?,?,?)",
                (
                    traversal_id,
                    child,
                    _text(row.get("parent_assay_identifier")),
                    _text(row.get("relationship_type")) or "source_declared_relationship",
                    1,
                    raw_record_id,
                ),
            )
    elif kind is ProviderRecordKind.DATASET_LINK:
        link_id = f"link-{deterministic_fingerprint([source_identifier, row])[:32]}"
        connection.execute(
            "INSERT OR IGNORE INTO publication_dataset_links VALUES (?,?,?,?,?,?,?,?,?)",
            (
                link_id,
                source_identifier,
                _text(row.get("publication_identifier")),
                _text(row.get("metadata_field_or_relation")),
                _text(row.get("dataset_identifier")),
                _text(row.get("target_provider")),
                _text(row.get("verification_status")),
                _json(provenance),
                raw_record_id,
            ),
        )
    elif kind is ProviderRecordKind.RELEASE:
        object_id = _text(row.get("id")) or _text(row.get("_id"))
        name = _text(row.get("filename")) or _text(row.get("name"))
        descriptor_type = "file" if row.get("filename") is not None else "dataset"
        inventory_id = (
            f"inventory-{deterministic_fingerprint([descriptor_type, object_id, row])[:32]}"
        )
        connection.execute(
            "INSERT OR IGNORE INTO release_inventory VALUES (?,?,?,?,?,?,?,?,?,?)",
            (
                inventory_id,
                descriptor_type,
                object_id,
                name,
                _text(row.get("contentType")) or _text(row.get("content_type")),
                row.get("size") if isinstance(row.get("size"), int) else None,
                _text(row.get("sha256"))
                or _text(row.get("contentHash"))
                or _text(row.get("hash"))
                or _text(row.get("md5")),
                _json(row),
                _json(provenance),
                raw_record_id,
            ),
        )


def _count(connection: sqlite3.Connection, table: str) -> int:
    return int(connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])


INDEX_NAMES = [
    "idx_raw_kind_source",
    "idx_candidates_source",
    "idx_hydrated_source",
    "idx_compound_cid",
    "idx_compound_inchikey",
    "idx_compound_sid",
    "idx_activity_assay_compound",
    "idx_activity_modality",
    "idx_activity_summary_aeid",
    "idx_assay_annotation_aeid",
    "idx_assay_source_partition",
    "idx_transcriptomic_compound",
    "idx_transcriptomic_accession",
    "idx_relationship_parent",
    "idx_traversal_assay",
    "idx_dataset_identifier",
    "idx_release_inventory_object",
]


class PreapprovalMetadataProvider:
    """Manifest-bound provider adapter with exact pagination and full-set handling."""

    def __init__(
        self,
        registry: PreapprovalProviderReleaseRegistry,
        executor: ProviderTaskExecutor,
        *,
        public_provider_cache_root: Path | None = None,
    ) -> None:
        self.registry = registry
        self.executor = executor
        self.public_provider_cache_root = (
            Path(public_provider_cache_root).resolve()
            if public_provider_cache_root is not None
            else None
        )

    @property
    def artifacts(self) -> LocalArtifactStore:
        return self.executor.artifacts

    def _load_identifiers(
        self,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        *,
        provider_batch_size: int,
    ) -> tuple[list[str], ProviderIdentifierReconciliation | None]:
        reference = request.upstream_identifier_artifact
        if reference is None:
            return [], None
        descriptor, path = self.artifacts.verified_path(reference.artifact_id)
        if descriptor.workflow_id != workflow_id or descriptor.sha256 != reference.sha256:
            raise ValueError("upstream identifier artifact is not bound to this workflow")
        received_identifiers: list[tuple[str, list[str]]] = []
        if path.read_bytes()[:16].startswith(b"SQLite format 3"):
            with sqlite3.connect(path) as connection:
                tables = {
                    row[0]
                    for row in connection.execute(
                        "SELECT name FROM sqlite_master WHERE type='table'"
                    )
                }
                if "compound_index" not in tables:
                    raise ValueError("identifier SQLite artifact lacks compound_index")
                rows = connection.execute(
                    """
                    SELECT COALESCE(compound_identifier, cid, inchikey, sid)
                    FROM compound_index
                    WHERE COALESCE(compound_identifier, cid, inchikey, sid) IS NOT NULL
                    ORDER BY mapping_id
                    """
                )
                received_identifiers = [(str(row[0]), []) for row in rows]
        else:
            with path.open(encoding="utf-8-sig") as handle:
                for line in handle:
                    value = line.strip()
                    parsed = json.loads(value) if value.startswith("{") else value
                    source_partitions: list[str] = []
                    if isinstance(parsed, dict):
                        partitions = parsed.get("source_partitions")
                        if isinstance(partitions, list):
                            source_partitions = sorted(
                                {str(item).strip() for item in partitions if str(item).strip()}
                            )
                        value = str(
                            parsed.get("compound_identifier")
                            or parsed.get("cid")
                            or parsed.get("inchikey")
                            or parsed.get("sid")
                            or ""
                        ).strip()
                    else:
                        value = str(parsed).strip()
                    received_identifiers.append((value, source_partitions))
        if len(received_identifiers) != request.upstream_identifier_count:
            raise ValueError(
                "upstream identifier artifact count does not match its declared exact count"
            )
        normalized: list[str] = []
        seen: set[str] = set()
        invalid: list[str] = []
        duplicate_count = 0
        mapping_rows: list[dict[str, Any]] = []
        normalized_ordinals: dict[str, int] = {}
        for ordinal, (original, source_partitions) in enumerate(received_identifiers, start=1):
            value = original.strip()
            match = re.fullmatch(r"CID:([1-9][0-9]{0,11})", value)
            if match is None and re.fullmatch(r"[1-9][0-9]{0,11}", value):
                match = re.fullmatch(r"([1-9][0-9]{0,11})", value)
            canonical = f"CID:{int(match.group(1))}" if match is not None else None
            if canonical is None:
                invalid.append(value[:120])
                status = "invalid"
            elif canonical in seen:
                duplicate_count += 1
                status = "duplicate"
            else:
                seen.add(canonical)
                normalized.append(canonical)
                normalized_ordinals[canonical] = len(normalized)
                status = "valid"
            normalized_ordinal = (
                normalized_ordinals.get(canonical) if canonical is not None else None
            )
            mapping_rows.append(
                {
                    "ordinal": ordinal,
                    "original_identifier": value,
                    "normalized_identifier": canonical,
                    "normalized_ordinal": normalized_ordinal,
                    "provider_batch_index": (
                        (normalized_ordinal - 1) // provider_batch_size
                        if normalized_ordinal is not None
                        else None
                    ),
                    "provider_batch_position": (
                        (normalized_ordinal - 1) % provider_batch_size
                        if normalized_ordinal is not None
                        else None
                    ),
                    "source_partitions": source_partitions,
                    "status": status,
                }
            )
        mapping = self.artifacts.put_json(
            workflow_id=workflow_id,
            value={
                "source_artifact": reference.model_dump(mode="json"),
                "received_identifier_count": len(received_identifiers),
                "normalized_unique_identifier_count": len(normalized),
                "duplicate_identifier_count": duplicate_count,
                "invalid_identifier_count": len(invalid),
                "provider_batch_size": provider_batch_size,
                "provider_batch_count": math.ceil(len(normalized) / provider_batch_size),
                "rows": mapping_rows,
            },
            artifact_type="provider_identifier_normalization_map",
            logical_name=f"provider-identifier-normalization-{reference.sha256}.json",
            producer="preapproval-provider:identifier-boundary",
            idempotency_key=f"provider-identifier-normalization:{reference.sha256}",
        )
        return normalized, ProviderIdentifierReconciliation(
            received_identifier_count=len(received_identifiers),
            normalized_unique_identifier_count=len(normalized),
            duplicate_identifier_count=duplicate_count,
            invalid_identifier_count=len(invalid),
            invalid_identifiers_preview=invalid[:20],
            provider_batch_size=provider_batch_size,
            provider_batch_count=math.ceil(len(normalized) / provider_batch_size),
            mapping_artifact=_artifact_reference(mapping),
        )

    def _execute_resource(
        self,
        *,
        manifest: PreapprovalProviderReleaseManifest,
        template: ProviderResourceTemplate,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        idempotency_key: str,
        ledger: DiscoveryExecutionRecord,
        page: int,
        cursor: str | None,
        identifiers: list[str],
        suffix: str,
        aeid: int | None = None,
    ) -> ProviderExecutionOutcome:
        page_size = manifest.page_size or manifest.identifier_batch_size or 1
        locator = _render_locator(
            template.locator_template,
            query=request.query,
            page=page,
            page_size=page_size,
            cursor=cursor,
            identifiers=identifiers,
            identifier_path_kind=template.identifier_path_kind,
            aeid=aeid,
        )
        if (urlparse(locator).hostname or "").lower() not in set(manifest.reviewed_hosts):
            raise ValueError("rendered provider locator escaped its reviewed host allowlist")
        resource_id = _safe_token(f"{template.resource_id}-{suffix}")
        resource = ProviderResourceRequest(
            resource_id=resource_id,
            logical_role=template.logical_role,
            operation_id=template.operation_id,
            locator=locator,
            item_kind=template.item_kind,
            required=template.required,
            expected_sha256=template.expected_sha256,
            accepted_mime_types=template.accepted_mime_types,
            maximum_response_bytes=template.maximum_response_bytes,
            approved_request_hosts=manifest.reviewed_hosts,
            cursor_or_file_token=cursor or suffix,
            compression=template.compression,
            required_credential=template.required_credential,
        )
        task = ProviderRetrievalTask.create(
            workflow_id=workflow_id,
            task_id=request.ledger_record.task_id,
            provider=manifest.provider,
            operation=template.operation_id or f"retrieve_preapproval_{manifest.provider}_metadata",
            source_version=manifest.source_version,
            source_locator=manifest.source_locator,
            resources=[resource],
            licence_and_provenance=manifest.licence_and_provenance,
            timeout_seconds=180.0,
            maximum_transport_retries=0,
            idempotency_key=f"{idempotency_key}:{suffix}",
        )
        parser = (
            _authoritative_toxcast_parser
            if manifest.retrieval_mode is ProviderRetrievalMode.TOXCAST_AUTHORITATIVE
            else _compact_parser
        )
        return self.executor.execute(task, ledger, parser=parser)

    def _records_from_execution(
        self, outcome: ProviderExecutionOutcome, operation_id: str | None = None
    ) -> list[dict[str, Any]]:
        records: list[dict[str, Any]] = []
        for result in outcome.page_or_file_results:
            if result.raw_artifact is None:
                continue
            _, path = self.artifacts.verified_path(result.raw_artifact.artifact_id)
            records.extend(_iter_source_records(path, operation_id))
        return records

    @staticmethod
    def _task_cancelled() -> bool:
        context = current_source_request_context()
        return context is not None and context.cancellation.cancelled

    @staticmethod
    def _stable_identifiers(value: Any) -> set[str]:
        stable: set[str] = set()
        stable_keys = {
            "cid",
            "pubchem_cid",
            "inchikey",
            "inchi_key",
            "dtxsid",
            "compound_identifier",
            "perturbagen_id",
            "pert_id",
        }

        def visit(item: Any, key: str | None = None) -> None:
            if isinstance(item, dict):
                for child_key, child in item.items():
                    visit(child, str(child_key).casefold())
                return
            if isinstance(item, list):
                for child in item:
                    visit(child, key)
                return
            if key not in stable_keys or item in (None, ""):
                return
            text = str(item).strip()
            if key in {"cid", "pubchem_cid"} and text.isdigit() and int(text) > 0:
                stable.add(f"CID:{int(text)}")
            elif key in {"inchikey", "inchi_key"} and re.fullmatch(
                r"[A-Z]{14}-[A-Z]{10}-[A-Z]", text.upper()
            ):
                stable.add(f"InChIKey:{text.upper()}")
            elif key == "dtxsid" and re.fullmatch(r"DTXSID[0-9]{7,12}", text.upper()):
                stable.add(text.upper())
            elif key == "compound_identifier":
                if re.fullmatch(r"CID:[1-9][0-9]*", text, re.IGNORECASE):
                    stable.add(f"CID:{int(text.split(':', maxsplit=1)[1])}")
                elif re.fullmatch(r"(?:INCHIKEY:)?[A-Z]{14}-[A-Z]{10}-[A-Z]", text.upper()):
                    stable.add(f"InChIKey:{text.upper().removeprefix('INCHIKEY:')}")
                elif re.fullmatch(r"DTXSID[0-9]{7,12}", text.upper()):
                    stable.add(text.upper())
            elif key in {"perturbagen_id", "pert_id"}:
                stable.add(f"PERT:{text}")

        visit(value)
        return stable

    def _pagination_yield(
        self,
        manifest: PreapprovalProviderReleaseManifest,
        executions: list[ProviderExecutionOutcome],
    ) -> list[ProviderPageYieldObservation]:
        observations: list[ProviderPageYieldObservation] = []
        template_by_role = {item.logical_role: item for item in manifest.resources}
        seen_stable: set[str] = set()
        seen_candidates: set[str] = set()
        for execution in executions:
            for result in execution.page_or_file_results:
                normalized = execution.normalized_resources.get(result.resource_id, {})
                records: list[dict[str, Any]] = []
                if result.raw_artifact is not None:
                    _, path = self.artifacts.verified_path(result.raw_artifact.artifact_id)
                    template = template_by_role[result.logical_role]
                    records = list(_iter_source_records(path, template.operation_id))
                stable = self._stable_identifiers([normalized, records])
                candidate_values = {
                    str(item)
                    for item in [
                        *(
                            value
                            for key in ("candidate_ids", "identifiers", "accessions")
                            for value in (
                                normalized.get(key, [])
                                if isinstance(normalized, dict)
                                and isinstance(normalized.get(key), list)
                                else []
                            )
                        ),
                        *(
                            row.get("source_identifier")
                            for row in records
                            if row.get("source_identifier") not in (None, "")
                        ),
                    ]
                    if isinstance(item, str | int)
                }
                new_stable = stable - seen_stable
                new_candidates = candidate_values - seen_candidates
                seen_stable.update(stable)
                seen_candidates.update(candidate_values)
                raw_rows = result.parsed_record_count
                unique = len(stable | candidate_values)
                observations.append(
                    ProviderPageYieldObservation(
                        resource_id=result.resource_id,
                        operation=result.logical_role,
                        transport_request_count=sum(
                            int(item.request_left_process) for item in result.transport_attempts
                        ),
                        raw_rows=raw_rows,
                        normalized_rows=raw_rows,
                        provider_unique_records=unique,
                        new_stable_compound_identifiers=len(new_stable),
                        new_compact_candidates=len(new_candidates),
                        duplicates=max(raw_rows - unique, 0),
                        joinability_contribution=len(new_stable),
                        continuation_token=(
                            str(normalized.get("next_cursor"))
                            if isinstance(normalized, dict)
                            and normalized.get("next_cursor") not in (None, "")
                            else None
                        ),
                    )
                )
        return observations

    def _cached_toxcast_normalization(self) -> ToxCastNormalizationResult | None:
        if self.public_provider_cache_root is None:
            return None
        cache = ToxCastArchiveCache(
            self.public_provider_cache_root,
            expected_archive_sha256=TOXCAST_ARCHIVE_SHA256,
        )
        return ToxCastActivityNormalizer(cache).load_cached()

    def _matched_public_toxcast_assays(
        self,
        manifest: PreapprovalProviderReleaseManifest,
        executions: list[ProviderExecutionOutcome],
        query: ProviderMetadataQuery,
    ) -> tuple[list[str], dict[str, str]]:
        templates = {item.logical_role: item for item in manifest.resources}
        target_rows: list[dict[str, Any]] = []
        assay_rows: list[dict[str, Any]] = []
        for execution in executions:
            for result in execution.page_or_file_results:
                template = templates[result.logical_role]
                if result.raw_artifact is None:
                    continue
                _, path = self.artifacts.verified_path(result.raw_artifact.artifact_id)
                rows = list(_iter_source_records(path, template.operation_id))
                if template.operation_id == "public_assay_target_mappings":
                    target_rows.extend(rows)
                elif template.operation_id == "public_assay_annotations":
                    assay_rows.extend(rows)

        targets_by_aeid = _toxcast_public_targets_by_aeid(iter(target_rows))
        target_terms = [
            _normalized_search_text(item)
            for item in (
                *query.target_identifiers,
                *query.target_synonyms,
                *([query.biological_target] if query.biological_target else []),
            )
            if _normalized_search_text(item)
        ]
        matched: list[str] = []
        modality_by_aeid: dict[str, str] = {}
        for row in assay_rows:
            assay = _toxcast_public_assay_row(row, targets_by_aeid)
            aeid = str(assay["assay_identifier"]).removeprefix("AEID:")
            haystack = _normalized_search_text(
                " ".join(
                    str(assay.get(field) or "")
                    for field in (
                        "target",
                        "assay_name",
                        "component_name",
                        "endpoint_name",
                        "intended_target_type",
                        "intended_target_family",
                        "biological_process",
                    )
                )
            )
            target_matches = not target_terms or any(
                term in haystack or all(word in haystack.split() for word in term.split())
                for term in target_terms
            )
            modality = _text(assay.get("modality"))
            modality_matches = query.modality is None or modality == query.modality
            if not target_matches or not modality_matches:
                continue
            matched.append(aeid)
            modality_by_aeid[aeid] = modality or "unclassified"
        return sorted(set(matched), key=int), modality_by_aeid

    def _persist_toxcast_public_activity_coverage(
        self,
        *,
        manifest: PreapprovalProviderReleaseManifest,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        idempotency_key: str,
        executions: list[ProviderExecutionOutcome],
        identifiers: list[str],
        normalized: ToxCastNormalizationResult,
    ) -> tuple[ToxCastPublicActivityCoverage, ArtifactReference]:
        if self.public_provider_cache_root is None:
            raise ValueError("public provider cache root is not configured")
        matched_aeids, modality_by_aeid = self._matched_public_toxcast_assays(
            manifest, executions, request.query
        )
        cache = ToxCastArchiveCache(
            self.public_provider_cache_root,
            expected_archive_sha256=TOXCAST_ARCHIVE_SHA256,
        )
        coverage = ToxCastCoverageService(cache).compute(
            normalized,
            ToxCastCoverageQuery(
                biological_target=(
                    request.query.biological_target
                    or "; ".join(request.query.target_synonyms)
                    or "; ".join(request.query.target_identifiers)
                    or "reviewed endpoint query"
                ),
                matched_aeids=matched_aeids,
                modality_by_aeid=modality_by_aeid,
                upstream_identity_set=identifiers,
            ),
        )
        summary = ToxCastPublicActivityCoverage(
            normalized_bundle=normalized,
            coverage=coverage,
            matched_aeids=matched_aeids,
            modality_by_aeid=modality_by_aeid,
        )
        descriptor = self.artifacts.put_json(
            workflow_id=workflow_id,
            value=summary.model_dump(mode="json"),
            artifact_type="toxcast_public_activity_coverage",
            logical_name=(
                "toxcast-public-activity-coverage-"
                f"{deterministic_fingerprint([normalized.bundle_fingerprint, request.query])}.json"
            ),
            producer="preapproval-provider:toxcast",
            idempotency_key=f"{idempotency_key}:toxcast-public-activity-coverage",
            original_source=manifest.source_locator,
        )
        return summary, _artifact_reference(descriptor)

    def _retrieve_public_toxcast(
        self,
        *,
        manifest: PreapprovalProviderReleaseManifest,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        idempotency_key: str,
    ) -> tuple[
        list[ProviderExecutionOutcome],
        DiscoveryExecutionRecord,
        bool,
        bool,
        dict[str, Any],
    ]:
        executions: list[ProviderExecutionOutcome] = []
        ledger = request.ledger_record
        assay_catalogue_count = 0
        public_resources = [
            item
            for item in manifest.resources
            if item.access_mode is ProviderAccessMode.PUBLIC_RELEASE
            and item.expansion_mode is ProviderExpansionMode.STATIC
        ]
        for index, template in enumerate(public_resources, start=1):
            outcome = self._execute_resource(
                manifest=manifest,
                template=template,
                request=request,
                workflow_id=workflow_id,
                idempotency_key=idempotency_key,
                ledger=ledger,
                page=index,
                cursor=None,
                identifiers=[],
                suffix=f"public-{index:04d}",
            )
            executions.append(outcome)
            ledger = outcome.ledger_record
            if (
                template.operation_id == "public_assay_annotations"
                and outcome.completion_proof.completed
            ):
                assay_catalogue_count = sum(
                    item.parsed_record_count for item in outcome.page_or_file_results
                )
        return (
            executions,
            ledger,
            True,
            False,
            {
                "assay_catalogue_count": assay_catalogue_count,
                "matched_aeid_count": 0,
                "summary_aeid_count": 0,
                "detail_aeid_count": 0,
            },
        )

    def _retrieve_authoritative_toxcast(
        self,
        *,
        manifest: PreapprovalProviderReleaseManifest,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        idempotency_key: str,
    ) -> tuple[
        list[ProviderExecutionOutcome],
        DiscoveryExecutionRecord,
        bool,
        bool,
        dict[str, int],
    ]:
        executions: list[ProviderExecutionOutcome] = []
        ledger = request.ledger_record
        assay_rows: list[dict[str, Any]] = []
        for index, template in enumerate(
            (
                item
                for item in manifest.resources
                if item.access_mode is ProviderAccessMode.AUTHENTICATED_API
                and item.expansion_mode is ProviderExpansionMode.STATIC
            ),
            start=1,
        ):
            outcome = self._execute_resource(
                manifest=manifest,
                template=template,
                request=request,
                workflow_id=workflow_id,
                idempotency_key=idempotency_key,
                ledger=ledger,
                page=index,
                cursor=None,
                identifiers=[],
                suffix=f"static-{index:04d}",
            )
            executions.append(outcome)
            ledger = outcome.ledger_record
            if template.operation_id == "all_assays" and outcome.completion_proof.completed:
                assay_rows = self._records_from_execution(outcome)
        if not all(item.completion_proof.completed for item in executions):
            return (
                executions,
                ledger,
                True,
                False,
                {
                    "assay_catalogue_count": len(assay_rows),
                    "matched_aeid_count": 0,
                    "summary_aeid_count": 0,
                    "detail_aeid_count": 0,
                },
            )

        matched_aeids, _ = _toxcast_matched_aeids(assay_rows, request.query)
        summary_aeids = matched_aeids
        detail_aeids = matched_aeids
        if request.validation_scope == "bounded_smoke":
            summary_aeids = matched_aeids[: request.maximum_summary_aeids]
            detail_aeids = matched_aeids[: request.maximum_detail_aeids]
        template_by_expansion: dict[ProviderExpansionMode, ProviderResourceTemplate] = {
            item.expansion_mode: item
            for item in manifest.resources
            if item.access_mode is ProviderAccessMode.AUTHENTICATED_API
            and item.expansion_mode is not ProviderExpansionMode.STATIC
        }
        for aeids, expansion, label in (
            (
                summary_aeids,
                ProviderExpansionMode.MATCHED_AEID_SUMMARY,
                "summary",
            ),
            (
                detail_aeids,
                ProviderExpansionMode.MATCHED_AEID_DETAILS,
                "details",
            ),
        ):
            template = template_by_expansion[expansion]
            for index, aeid in enumerate(aeids, start=1):
                outcome = self._execute_resource(
                    manifest=manifest,
                    template=template,
                    request=request,
                    workflow_id=workflow_id,
                    idempotency_key=idempotency_key,
                    ledger=ledger,
                    page=index,
                    cursor=None,
                    identifiers=[],
                    suffix=f"{label}-aeid-{aeid}",
                    aeid=aeid,
                )
                executions.append(outcome)
                ledger = outcome.ledger_record
                if not outcome.completion_proof.completed:
                    break
            if not all(item.completion_proof.completed for item in executions):
                break
        return (
            executions,
            ledger,
            True,
            False,
            {
                "assay_catalogue_count": len(assay_rows),
                "matched_aeid_count": len(matched_aeids),
                "summary_aeid_count": len(summary_aeids),
                "detail_aeid_count": len(detail_aeids),
            },
        )

    def _candidate_ids(
        self,
        manifest: PreapprovalProviderReleaseManifest,
        executions: list[ProviderExecutionOutcome],
    ) -> list[str]:
        template_by_role = {item.logical_role: item for item in manifest.resources}
        identifiers: list[str] = []
        for execution in executions:
            for result in execution.page_or_file_results:
                template = template_by_role[result.logical_role]
                if (
                    template.expansion_mode is not ProviderExpansionMode.STATIC
                    or result.raw_artifact is None
                ):
                    continue
                _, path = self.artifacts.verified_path(result.raw_artifact.artifact_id)
                for row in _iter_source_records(path, template.operation_id):
                    identifier = _text(row.get("source_identifier"))
                    if identifier and identifier.isdigit():
                        identifiers.append(identifier)
        return list(dict.fromkeys(identifiers))

    def _geo_accessions(
        self,
        manifest: PreapprovalProviderReleaseManifest,
        executions: list[ProviderExecutionOutcome],
    ) -> list[str]:
        template_by_role = {item.logical_role: item for item in manifest.resources}
        accessions: list[str] = []
        for execution in executions:
            for result in execution.page_or_file_results:
                template = template_by_role[result.logical_role]
                if (
                    template.expansion_mode is not ProviderExpansionMode.CANDIDATE_SUMMARY_BATCH
                    or result.raw_artifact is None
                ):
                    continue
                _, path = self.artifacts.verified_path(result.raw_artifact.artifact_id)
                for row in _iter_source_records(path, template.operation_id):
                    accession = str(_row_value(row, "accession", "gse") or "").upper()
                    if re.fullmatch(r"GSE[1-9][0-9]{1,8}", accession):
                        accessions.append(accession)
        return list(dict.fromkeys(accessions))

    def _retrieve_cursor_hydration(
        self,
        *,
        manifest: PreapprovalProviderReleaseManifest,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        idempotency_key: str,
    ) -> tuple[
        list[ProviderExecutionOutcome],
        DiscoveryExecutionRecord,
        bool,
        bool,
        dict[str, int],
    ]:
        executions: list[ProviderExecutionOutcome] = []
        ledger = request.ledger_record
        search_template = next(
            item
            for item in manifest.resources
            if item.expansion_mode is ProviderExpansionMode.STATIC
        )
        cursor: str | None = None
        cursor_exhausted = False
        search_safety_truncated = False
        bounded_shortlist_satisfied = False
        seen_search_candidates: set[str] = set()
        seen_search_stable_identifiers: set[str] = set()
        consecutive_no_yield_pages = 0
        for page in range(1, manifest.maximum_pages + 1):
            outcome = self._execute_resource(
                manifest=manifest,
                template=search_template,
                request=request,
                workflow_id=workflow_id,
                idempotency_key=idempotency_key,
                ledger=ledger,
                page=page,
                cursor=cursor,
                identifiers=[],
                suffix=f"search-page-{page:06d}",
            )
            executions.append(outcome)
            ledger = outcome.ledger_record
            if self._task_cancelled():
                search_safety_truncated = True
                break
            compact = next(iter(outcome.normalized_resources.values()), {})
            page_candidates = set(self._candidate_ids(manifest, [outcome]))
            new_candidates = page_candidates - seen_search_candidates
            page_stable_identifiers = self._stable_identifiers(
                [
                    compact,
                    self._records_from_execution(outcome, search_template.operation_id),
                ]
            )
            new_stable_identifiers = page_stable_identifiers - seen_search_stable_identifiers
            seen_search_candidates.update(page_candidates)
            seen_search_stable_identifiers.update(page_stable_identifiers)
            next_cursor = _text(compact.get("next_cursor"))
            terminal = bool(compact.get("terminal", not next_cursor))
            if terminal and not next_cursor:
                cursor_exhausted = True
                break
            if (
                request.maximum_hydration_candidates is not None
                and len(seen_search_candidates) >= request.maximum_hydration_candidates
            ):
                bounded_shortlist_satisfied = True
                break
            consecutive_no_yield_pages = (
                0 if new_candidates or new_stable_identifiers else consecutive_no_yield_pages + 1
            )
            if consecutive_no_yield_pages >= 2:
                search_safety_truncated = True
                break
            cursor = next_cursor or str(page * (manifest.page_size or 1))
        else:
            search_safety_truncated = True

        discovered = self._candidate_ids(manifest, executions)
        hydration_ids = discovered
        excluded_hydration_ids: list[str] = []
        hydration_safety_truncated = False
        if (
            request.maximum_hydration_candidates is not None
            and len(hydration_ids) > request.maximum_hydration_candidates
        ):
            excluded_hydration_ids = hydration_ids[request.maximum_hydration_candidates :]
            hydration_ids = hydration_ids[: request.maximum_hydration_candidates]

        templates: dict[ProviderExpansionMode, ProviderResourceTemplate] = {
            item.expansion_mode: item
            for item in manifest.resources
            if item.expansion_mode is not ProviderExpansionMode.STATIC
        }
        batch_size = manifest.identifier_batch_size or 1
        summary_template = templates[ProviderExpansionMode.CANDIDATE_SUMMARY_BATCH]
        summary_record_count = 0
        for batch_index, start in enumerate(range(0, len(hydration_ids), batch_size), start=1):
            batch = hydration_ids[start : start + batch_size]
            outcome = self._execute_resource(
                manifest=manifest,
                template=summary_template,
                request=request,
                workflow_id=workflow_id,
                idempotency_key=idempotency_key,
                ledger=ledger,
                page=batch_index,
                cursor=None,
                identifiers=batch,
                suffix=f"summary-batch-{batch_index:06d}",
            )
            executions.append(outcome)
            ledger = outcome.ledger_record
            summary_record_count += sum(
                item.parsed_record_count for item in outcome.page_or_file_results
            )
            if self._task_cancelled():
                hydration_safety_truncated = True
                break

        activity_attempts = 0
        relationship_attempts = 0
        geo_soft_attempts = 0
        resolved_geo_accessions: list[str] = []
        if manifest.provider in {"tox21", "pubchem-bioassay"}:
            for identifier in hydration_ids:
                for expansion, label in (
                    (
                        ProviderExpansionMode.CANDIDATE_COMPOUND_IDENTIFIERS,
                        "compound-identifiers",
                    ),
                    (ProviderExpansionMode.CANDIDATE_ACTIVITY_ROWS, "activity"),
                    (ProviderExpansionMode.CANDIDATE_RELATIONSHIPS, "relationships"),
                ):
                    template = templates[expansion]
                    outcome = self._execute_resource(
                        manifest=manifest,
                        template=template,
                        request=request,
                        workflow_id=workflow_id,
                        idempotency_key=idempotency_key,
                        ledger=ledger,
                        page=1,
                        cursor=None,
                        identifiers=[identifier],
                        suffix=f"{label}-aid-{identifier}",
                    )
                    executions.append(outcome)
                    ledger = outcome.ledger_record
                    if (
                        expansion is ProviderExpansionMode.CANDIDATE_ACTIVITY_ROWS
                        and outcome.completion_proof.completed
                    ):
                        activity_attempts += 1
                    elif (
                        expansion is ProviderExpansionMode.CANDIDATE_RELATIONSHIPS
                        and outcome.completion_proof.completed
                    ):
                        relationship_attempts += 1
                    if self._task_cancelled():
                        hydration_safety_truncated = True
                        break
                if self._task_cancelled():
                    break
        else:
            resolved_geo_accessions = self._geo_accessions(manifest, executions)
            soft_template = templates[ProviderExpansionMode.GEO_SERIES_SOFT]
            for accession in resolved_geo_accessions:
                outcome = self._execute_resource(
                    manifest=manifest,
                    template=soft_template,
                    request=request,
                    workflow_id=workflow_id,
                    idempotency_key=idempotency_key,
                    ledger=ledger,
                    page=1,
                    cursor=None,
                    identifiers=[accession],
                    suffix=f"soft-{accession.lower()}",
                )
                executions.append(outcome)
                ledger = outcome.ledger_record
                geo_soft_attempts += 1
                if self._task_cancelled():
                    hydration_safety_truncated = True
                    break

        return (
            executions,
            ledger,
            cursor_exhausted,
            search_safety_truncated or hydration_safety_truncated,
            {
                "assay_catalogue_count": 0,
                "matched_aeid_count": 0,
                "summary_aeid_count": 0,
                "detail_aeid_count": 0,
                "discovered_candidate_count": len(discovered),
                "hydrated_candidate_count": len(hydration_ids),
                "shortlisted_candidate_ids": list(hydration_ids),
                "excluded_candidate_ids": excluded_hydration_ids,
                "bounded_shortlist_satisfied": bounded_shortlist_satisfied,
                "hydrated_summary_record_count": summary_record_count,
                "activity_availability_attempt_count": activity_attempts,
                "relationship_attempt_count": relationship_attempts,
                "resolved_geo_accession_count": len(resolved_geo_accessions),
                "geo_soft_attempt_count": geo_soft_attempts,
            },
        )

    def _retrieve(
        self,
        *,
        manifest: PreapprovalProviderReleaseManifest,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        idempotency_key: str,
        identifiers: list[str],
    ) -> tuple[
        list[ProviderExecutionOutcome],
        DiscoveryExecutionRecord,
        bool,
        bool,
        dict[str, Any],
    ]:
        if manifest.retrieval_mode is ProviderRetrievalMode.TOXCAST_AUTHORITATIVE:
            if request.access_mode is ProviderAccessMode.PUBLIC_RELEASE:
                return self._retrieve_public_toxcast(
                    manifest=manifest,
                    request=request,
                    workflow_id=workflow_id,
                    idempotency_key=idempotency_key,
                )
            return self._retrieve_authoritative_toxcast(
                manifest=manifest,
                request=request,
                workflow_id=workflow_id,
                idempotency_key=idempotency_key,
            )
        if manifest.retrieval_mode is ProviderRetrievalMode.CURSOR_HYDRATION:
            return self._retrieve_cursor_hydration(
                manifest=manifest,
                request=request,
                workflow_id=workflow_id,
                idempotency_key=idempotency_key,
            )
        executions: list[ProviderExecutionOutcome] = []
        ledger = request.ledger_record
        cursor_exhausted = manifest.retrieval_mode is not ProviderRetrievalMode.CURSOR_PAGES
        safety_truncated = False
        processed_identifier_count = 0
        if manifest.retrieval_mode is ProviderRetrievalMode.MANIFEST_FILES:
            for index, template in enumerate(manifest.resources, start=1):
                outcome = self._execute_resource(
                    manifest=manifest,
                    template=template,
                    request=request,
                    workflow_id=workflow_id,
                    idempotency_key=idempotency_key,
                    ledger=ledger,
                    page=index,
                    cursor=None,
                    identifiers=[],
                    suffix=f"file-{index:04d}",
                )
                executions.append(outcome)
                ledger = outcome.ledger_record
                if self._task_cancelled():
                    safety_truncated = True
                    break
        elif manifest.retrieval_mode is ProviderRetrievalMode.IDENTIFIER_BATCHES:
            if not identifiers:
                if request.upstream_identifier_artifact is None:
                    raise ValueError("identifier-batch provider requires a complete upstream index")
                return (
                    executions,
                    ledger,
                    True,
                    False,
                    {
                        "assay_catalogue_count": 0,
                        "matched_aeid_count": 0,
                        "summary_aeid_count": 0,
                        "detail_aeid_count": 0,
                        "processed_identifier_count": 0,
                    },
                )
            batch_size = manifest.identifier_batch_size or 1
            template = manifest.resources[0]
            for index, start in enumerate(range(0, len(identifiers), batch_size), start=1):
                batch = identifiers[start : start + batch_size]
                outcome = self._execute_resource(
                    manifest=manifest,
                    template=template,
                    request=request,
                    workflow_id=workflow_id,
                    idempotency_key=idempotency_key,
                    ledger=ledger,
                    page=index,
                    cursor=None,
                    identifiers=batch,
                    suffix=f"batch-{index:06d}",
                )
                executions.append(outcome)
                ledger = outcome.ledger_record
                if outcome.completion_proof.completed:
                    processed_identifier_count += len(batch)
                if self._task_cancelled():
                    safety_truncated = True
                    break
        else:
            template = manifest.resources[0]
            cursor: str | None = None
            seen_candidates: set[str] = set()
            seen_stable_identifiers: set[str] = set()
            consecutive_no_yield_pages = 0
            for page in range(1, manifest.maximum_pages + 1):
                outcome = self._execute_resource(
                    manifest=manifest,
                    template=template,
                    request=request,
                    workflow_id=workflow_id,
                    idempotency_key=idempotency_key,
                    ledger=ledger,
                    page=page,
                    cursor=cursor,
                    identifiers=[],
                    suffix=f"page-{page:06d}",
                )
                executions.append(outcome)
                ledger = outcome.ledger_record
                if self._task_cancelled():
                    safety_truncated = True
                    break
                compact = next(iter(outcome.normalized_resources.values()), {})
                page_candidates = set(self._candidate_ids(manifest, [outcome]))
                new_candidates = page_candidates - seen_candidates
                page_stable_identifiers = self._stable_identifiers(
                    [compact, self._records_from_execution(outcome, template.operation_id)]
                )
                new_stable_identifiers = page_stable_identifiers - seen_stable_identifiers
                seen_candidates.update(page_candidates)
                seen_stable_identifiers.update(page_stable_identifiers)
                consecutive_no_yield_pages = (
                    0
                    if new_candidates or new_stable_identifiers
                    else consecutive_no_yield_pages + 1
                )
                next_cursor = _text(compact.get("next_cursor"))
                terminal = bool(compact.get("terminal", not next_cursor))
                if terminal and not next_cursor:
                    cursor_exhausted = True
                    break
                if consecutive_no_yield_pages >= 2:
                    safety_truncated = True
                    break
                cursor = next_cursor or str(page * (manifest.page_size or 1))
            else:
                safety_truncated = True
        return (
            executions,
            ledger,
            cursor_exhausted,
            safety_truncated,
            {
                "assay_catalogue_count": 0,
                "matched_aeid_count": 0,
                "summary_aeid_count": 0,
                "detail_aeid_count": 0,
                "processed_identifier_count": processed_identifier_count,
            },
        )

    @staticmethod
    def _failure_classification(executions: list[ProviderExecutionOutcome]) -> str | None:
        for execution in executions:
            for result in execution.page_or_file_results:
                if result.error_classification:
                    return result.error_classification
        return None

    def _structural_validation_codes(
        self,
        manifest: PreapprovalProviderReleaseManifest,
        executions: list[ProviderExecutionOutcome],
        metrics: dict[str, Any],
    ) -> list[str]:
        template_by_role = {item.logical_role: item for item in manifest.resources}
        scientific_hash_roles: dict[str, set[str]] = {}
        codes: set[str] = set()
        for execution in executions:
            for result in execution.page_or_file_results:
                template = template_by_role[result.logical_role]
                if (
                    result.observed_sha256
                    and template.required
                    and template.resource_scope is ProviderResourceScope.SCIENTIFIC_DATA
                ):
                    scientific_hash_roles.setdefault(result.observed_sha256, set()).add(
                        result.logical_role
                    )
                if (
                    template.resource_scope is ProviderResourceScope.SCIENTIFIC_DATA
                    and "/api/spaces/" in result.locator
                    and result.locator.rstrip("/").endswith("/datasets")
                ):
                    codes.add("RELEASE_LIST_CANNOT_SATISFY_SCIENTIFIC_ROLE")
        if any(len(roles) > 1 for roles in scientific_hash_roles.values()):
            codes.add("ROLE_ENDPOINT_ALIASING_DETECTED")
        if (
            manifest.retrieval_mode is ProviderRetrievalMode.TOXCAST_AUTHORITATIVE
            and metrics["assay_catalogue_count"] <= 5
            and all(item.completion_proof.completed for item in executions)
        ):
            codes.add("ASSAY_CATALOGUE_COMPLETION_NOT_PROVEN")
        if manifest.retrieval_mode is ProviderRetrievalMode.CURSOR_HYDRATION:
            hydrated = metrics.get("hydrated_candidate_count", 0)
            if metrics.get("hydrated_summary_record_count", 0) < hydrated:
                codes.add("CANDIDATE_SUMMARY_HYDRATION_INCOMPLETE")
            if manifest.provider in {"tox21", "pubchem-bioassay"}:
                if metrics.get("activity_availability_attempt_count", 0) != hydrated:
                    codes.add("ACTIVITY_AVAILABILITY_NOT_INSPECTED")
                if metrics.get("relationship_attempt_count", 0) != hydrated:
                    codes.add("ASSAY_RELATIONSHIPS_NOT_INSPECTED")
            if manifest.provider == "ncbi-geo" and metrics.get(
                "geo_soft_attempt_count", 0
            ) != metrics.get("resolved_geo_accession_count", 0):
                codes.add("GEO_SERIES_SOFT_HYDRATION_INCOMPLETE")
        return sorted(codes)

    def _normalize(
        self,
        *,
        manifest: PreapprovalProviderReleaseManifest,
        request: ProviderMetadataExecutionInput,
        workflow_id: str,
        idempotency_key: str,
        executions: list[ProviderExecutionOutcome],
        proof: ProviderDatasetCompletionProof,
    ) -> tuple[ProviderNormalizedDatasetManifest, ArtifactReference, bool]:
        raw_by_id: dict[str, tuple[ArtifactReference, str, ProviderRecordKind, str | None]] = {}
        template_by_role = {item.logical_role: item for item in manifest.resources}
        for execution in executions:
            for result in execution.page_or_file_results:
                if result.raw_artifact is None:
                    continue
                template = template_by_role[result.logical_role]
                raw_by_id[result.raw_artifact.artifact_id] = (
                    result.raw_artifact,
                    result.resource_id,
                    template.record_kind,
                    template.operation_id,
                )
        raw_artifacts = sorted(
            (value[0] for value in raw_by_id.values()), key=lambda x: x.artifact_id
        )
        cache_identity = deterministic_fingerprint(
            {
                "provider": manifest.provider,
                "release": manifest.release_id,
                "access_mode": request.access_mode,
                "task": request.ledger_record.task_id,
                "manifest": manifest.manifest_fingerprint,
                "raw": [(item.artifact_id, item.sha256) for item in raw_artifacts],
                "normalizer": PREAPPROVAL_NORMALIZED_SCHEMA_VERSION,
            }
        )
        logical_name = f"provider-normalization-{cache_identity}.json"
        cached = self.artifacts.find_by_logical_name(workflow_id, logical_name)
        if cached is not None:
            _, content = self.artifacts.get(cached.id)
            normalized = ProviderNormalizedDatasetManifest.model_validate_json(content)
            self.artifacts.verified_path(normalized.sqlite_artifact.artifact_id)
            return normalized, _artifact_reference(cached), False

        with tempfile.TemporaryDirectory(
            prefix=f"provider-{manifest.provider}-",
            dir=self.artifacts.root,
            ignore_cleanup_errors=True,
        ) as temp:
            sqlite_path = Path(temp) / "normalized.sqlite"
            with closing(sqlite3.connect(sqlite_path)) as connection:
                _create_schema(connection)
                assay_by_aeid: dict[str, dict[str, Any]] = {}
                public_targets_by_aeid: dict[str, list[dict[str, Any]]] = {}
                for reference, _resource_id, _kind, operation_id in raw_by_id.values():
                    if operation_id != "public_assay_target_mappings":
                        continue
                    _, raw_path = self.artifacts.verified_path(reference.artifact_id)
                    public_targets_by_aeid = _toxcast_public_targets_by_aeid(
                        _iter_source_records(raw_path, operation_id)
                    )
                for reference, _resource_id, _kind, operation_id in raw_by_id.values():
                    if operation_id not in {"all_assays", "public_assay_annotations"}:
                        continue
                    _, raw_path = self.artifacts.verified_path(reference.artifact_id)
                    for source_row in _iter_source_records(raw_path, operation_id):
                        assay = (
                            _toxcast_public_assay_row(source_row, public_targets_by_aeid)
                            if operation_id == "public_assay_annotations"
                            else _toxcast_assay_row(source_row)
                        )
                        assay_by_aeid[assay["assay_identifier"].removeprefix("AEID:")] = assay
                for reference, resource_id, kind, operation_id in raw_by_id.values():
                    descriptor, raw_path = self.artifacts.verified_path(reference.artifact_id)
                    for ordinal, source_row in enumerate(
                        _iter_source_records(raw_path, operation_id), start=1
                    ):
                        rows = (
                            _toxcast_rows_for_operation(
                                operation_id,
                                source_row,
                                assay_by_aeid,
                                public_targets_by_aeid,
                            )
                            if manifest.retrieval_mode
                            is ProviderRetrievalMode.TOXCAST_AUTHORITATIVE
                            else _provider_rows_for_operation(
                                provider=manifest.provider,
                                operation_id=operation_id,
                                resource_id=resource_id,
                                kind=kind,
                                source_row=source_row,
                                query=request.query,
                            )
                        )
                        for derived_ordinal, row in enumerate(rows, start=1):
                            row_kind = ProviderRecordKind(row.get("record_kind", kind.value))
                            source_identifier = _source_identifier(row, row_kind)
                            raw_fingerprint = deterministic_fingerprint(
                                [reference.sha256, ordinal, derived_ordinal, row]
                            )
                            raw_record_id = f"raw-{raw_fingerprint[:32]}"
                            connection.execute(
                                "INSERT OR IGNORE INTO raw_records VALUES (?,?,?,?,?,?,?,?,?)",
                                (
                                    raw_record_id,
                                    manifest.provider,
                                    manifest.release_id,
                                    resource_id,
                                    row_kind.value,
                                    source_identifier,
                                    _json(row),
                                    reference.artifact_id,
                                    descriptor.sha256,
                                ),
                            )
                            if not (
                                manifest.provider == "toxcast"
                                and row_kind is ProviderRecordKind.ASSAY
                                and row.get("source_partition") == "tox21"
                            ):
                                _insert_candidate_and_hydration(
                                    connection,
                                    provider=manifest.provider,
                                    source_identifier=source_identifier,
                                    kind=row_kind,
                                    row=row,
                                    task_id=request.ledger_record.task_id,
                                    raw_record_id=raw_record_id,
                                )
                            _insert_kind_record(
                                connection,
                                kind=row_kind,
                                source_identifier=source_identifier,
                                row=row,
                                raw_record_id=raw_record_id,
                            )
                if request.upstream_identifier_artifact is not None:
                    requested_identifiers, _ = self._load_identifiers(
                        request,
                        workflow_id,
                        provider_batch_size=manifest.identifier_batch_size,
                    )
                    observed_identifiers = {
                        str(row[0])
                        for row in connection.execute(
                            """
                            SELECT COALESCE(compound_identifier, source_identifier)
                            FROM compound_index
                            """
                        )
                    }
                    fallback_artifact = raw_artifacts[0]
                    for identifier in requested_identifiers:
                        if identifier in observed_identifiers:
                            continue
                        missing_row = {
                            "record_kind": "compound",
                            "source_identifier": identifier,
                            "compound_identifier": identifier,
                            "mapping_status": "unavailable_record",
                            "structure_status": "unavailable_record",
                            "provenance": [
                                "Requested identifier was absent from the reviewed "
                                "provider response."
                            ],
                        }
                        missing_fingerprint = deterministic_fingerprint(
                            [fallback_artifact.sha256, missing_row]
                        )
                        raw_record_id = f"raw-{missing_fingerprint[:32]}"
                        connection.execute(
                            "INSERT OR IGNORE INTO raw_records VALUES (?,?,?,?,?,?,?,?,?)",
                            (
                                raw_record_id,
                                manifest.provider,
                                manifest.release_id,
                                "deterministic-missing-identifier-reconciliation",
                                ProviderRecordKind.COMPOUND.value,
                                identifier,
                                _json(missing_row),
                                fallback_artifact.artifact_id,
                                fallback_artifact.sha256,
                            ),
                        )
                        _insert_candidate_and_hydration(
                            connection,
                            provider=manifest.provider,
                            source_identifier=identifier,
                            kind=ProviderRecordKind.COMPOUND,
                            row=missing_row,
                            task_id=request.ledger_record.task_id,
                            raw_record_id=raw_record_id,
                        )
                        _insert_kind_record(
                            connection,
                            kind=ProviderRecordKind.COMPOUND,
                            source_identifier=identifier,
                            row=missing_row,
                            raw_record_id=raw_record_id,
                        )
                connection.commit()
                counts = ProviderNormalizedCounts(
                    raw_records=_count(connection, "raw_records"),
                    source_candidates=_count(connection, "source_candidates"),
                    hydrated_sources=_count(connection, "hydrated_sources"),
                    compound_mappings=_count(connection, "compound_index"),
                    activity_records=_count(connection, "activity_records"),
                    activity_summaries=_count(connection, "activity_summaries"),
                    transcriptomic_records=_count(connection, "transcriptomic_records"),
                    assay_relationships=_count(connection, "assay_relationships"),
                    assay_traversal_records=_count(connection, "assay_traversal_ledger"),
                    assay_annotations=_count(connection, "assay_annotations"),
                    release_inventory_records=_count(connection, "release_inventory"),
                    publication_dataset_links=_count(connection, "publication_dataset_links"),
                    exact_identity_mappings=int(
                        connection.execute(
                            "SELECT COUNT(*) FROM compound_index WHERE mapping_status='exact'"
                        ).fetchone()[0]
                    ),
                    ambiguous_identity_mappings=int(
                        connection.execute(
                            """
                            SELECT COUNT(*) FROM compound_index
                            WHERE mapping_status IN (
                                'ambiguous', 'multiple_structures', 'mixture',
                                'salt_or_parent', 'stereochemical_difference'
                            )
                            """
                        ).fetchone()[0]
                    ),
                    missing_identity_mappings=int(
                        connection.execute(
                            """
                            SELECT COUNT(*) FROM compound_index
                            WHERE mapping_status IN (
                                'missing', 'unavailable_record', 'transport_failure'
                            )
                            """
                        ).fetchone()[0]
                    ),
                )
                modalities = [
                    str(row[0])
                    for row in connection.execute(
                        """
                        SELECT DISTINCT modality FROM activity_records
                        WHERE modality IS NOT NULL ORDER BY modality
                        """
                    )
                ]
                provider_partitions = {
                    str(row[0]): int(row[1])
                    for row in connection.execute(
                        """
                        SELECT source_partition, COUNT(*) FROM assay_annotations
                        GROUP BY source_partition ORDER BY source_partition
                        """
                    )
                }
            sqlite_descriptor = self.artifacts.put_file(
                workflow_id=workflow_id,
                source_path=sqlite_path,
                mime_type="application/octet-stream",
                artifact_type=f"{manifest.provider}_preapproval_metadata_sqlite",
                logical_name=f"provider-{cache_identity}.sqlite",
                producer=f"preapproval-provider:{manifest.adapter_id}",
                original_source=manifest.source_locator,
                idempotency_key=f"{idempotency_key}:normalized-sqlite",
                reviewed_maximum_bytes=MAXIMUM_PREAPPROVAL_NORMALIZED_SQLITE_BYTES,
            )
        sqlite_reference = _artifact_reference(sqlite_descriptor)
        bundle_fingerprint = deterministic_fingerprint(
            {
                "manifest": manifest.manifest_fingerprint,
                "sqlite_sha256": sqlite_descriptor.sha256,
                "raw_sha256": [item.sha256 for item in raw_artifacts],
                "counts": counts,
                "modalities": modalities,
                "completion": proof,
                "access_mode": request.access_mode,
                "capability_findings": manifest.public_capabilities,
            }
        )
        operation_ids = sorted(
            {
                operation_id
                for execution in executions
                for result in execution.page_or_file_results
                if (operation_id := template_by_role[result.logical_role].operation_id) is not None
            }
        )
        normalized = ProviderNormalizedDatasetManifest(
            provider=manifest.provider,
            release_id=manifest.release_id,
            source_system=manifest.source_system,
            source_version=manifest.source_version,
            adapter_id=manifest.adapter_id,
            adapter_version=manifest.adapter_version,
            release_manifest_fingerprint=manifest.manifest_fingerprint,
            sqlite_artifact=sqlite_reference,
            raw_artifacts=raw_artifacts,
            counts=counts,
            modalities=modalities,
            provider_provenance_partitions=provider_partitions,
            relationship_findings=manifest.unavailable_relationship_fields,
            operation_hashes={
                operation: sorted(
                    {
                        result.observed_sha256
                        for execution in executions
                        for result in execution.page_or_file_results
                        if result.observed_sha256
                        and template_by_role[result.logical_role].operation_id == operation
                    }
                )
                for operation in operation_ids
            },
            index_names=INDEX_NAMES,
            completion_proof=proof,
            bundle_fingerprint=bundle_fingerprint,
            access_mode=request.access_mode,
            capability_findings=(
                manifest.public_capabilities
                if request.access_mode is ProviderAccessMode.PUBLIC_RELEASE
                else []
            ),
        )
        descriptor = self.artifacts.put_json(
            workflow_id=workflow_id,
            value=normalized.model_dump(mode="json"),
            artifact_type="preapproval_provider_normalization_manifest",
            logical_name=logical_name,
            producer=f"preapproval-provider:{manifest.adapter_id}",
            idempotency_key=f"{idempotency_key}:normalization-manifest",
        )
        return normalized, _artifact_reference(descriptor), True

    def _previews(
        self,
        manifest: PreapprovalProviderReleaseManifest,
        request: ProviderMetadataExecutionInput,
        normalized: ProviderNormalizedDatasetManifest,
    ) -> tuple[DiskBackedSourceCandidateSet, DiskBackedHydratedSourceSet]:
        _, path = self.artifacts.verified_path(normalized.sqlite_artifact.artifact_id)
        raw_refs = {item.artifact_id: item for item in normalized.raw_artifacts}
        with sqlite3.connect(path) as connection:
            candidate_rows = connection.execute(
                """
                SELECT c.candidate_id, c.source_identifier, c.evidence_role, c.modality,
                       c.task_id, c.candidate_status, r.source_artifact_id
                FROM source_candidates c JOIN raw_records r ON r.record_id=c.raw_record_id
                ORDER BY c.candidate_id LIMIT ?
                """,
                (PREVIEW_LIMIT,),
            ).fetchall()
            candidates = [
                SourceCandidate(
                    candidate_id=row[0],
                    provider=manifest.provider,
                    source_identifier=row[1],
                    evidence_role=row[2],
                    modality=row[3],
                    discovery_task_id=row[4],
                    discovery_query_provenance=[raw_refs[row[6]]],
                    candidate_status=row[5],
                )
                for row in candidate_rows
            ]
            hydrated_rows = connection.execute(
                """
                SELECT h.hydrated_source_id, h.candidate_id, h.source_identifier,
                       h.evidence_role, h.modality, h.verified_metadata_json,
                       h.compound_index_available, h.label_fields_available, h.organism,
                       h.experimental_context_json, h.completeness_status,
                       h.missing_fields_json, h.exclusion_reason, r.source_artifact_id
                FROM hydrated_sources h JOIN raw_records r ON r.record_id=h.raw_record_id
                ORDER BY h.hydrated_source_id LIMIT ?
                """,
                (PREVIEW_LIMIT,),
            ).fetchall()
            hydrated = [
                HydratedSource(
                    hydrated_source_id=row[0],
                    source_candidate_id=row[1],
                    provider=manifest.provider,
                    source_identifier=row[2],
                    evidence_role=row[3],
                    modality=row[4],
                    verified_metadata=json.loads(row[5]),
                    compound_index_available=bool(row[6]),
                    label_or_activity_fields_available=bool(row[7]),
                    organism=row[8],
                    experimental_context=json.loads(row[9]),
                    source_version=manifest.source_version,
                    licence_and_provenance=manifest.licence_and_provenance,
                    completeness_status=row[10],
                    explicit_missing_fields=json.loads(row[11]),
                    exclusion_reason=row[12],
                    evidence_artifacts=[raw_refs[row[13]]],
                )
                for row in hydrated_rows
            ]
        return (
            DiskBackedSourceCandidateSet(
                artifact=normalized.sqlite_artifact,
                exact_count=normalized.counts.source_candidates,
                preview=candidates,
                preview_is_complete=(normalized.counts.source_candidates <= PREVIEW_LIMIT),
            ),
            DiskBackedHydratedSourceSet(
                artifact=normalized.sqlite_artifact,
                exact_count=normalized.counts.hydrated_sources,
                preview=hydrated,
                preview_is_complete=(normalized.counts.hydrated_sources <= PREVIEW_LIMIT),
            ),
        )

    def _persist_execution_manifest(
        self,
        *,
        workflow_id: str,
        idempotency_key: str,
        provider: str,
        release_id: str,
        executions: list[ProviderExecutionOutcome],
    ) -> ArtifactReference:
        execution_fingerprint = deterministic_fingerprint(
            [item.task.task_fingerprint for item in executions]
        )
        descriptor = self.artifacts.put_json(
            workflow_id=workflow_id,
            value={
                "provider": provider,
                "release_id": release_id,
                "exact_execution_count": len(executions),
                "executions": [item.model_dump(mode="json") for item in executions],
            },
            artifact_type="preapproval_provider_execution_manifest",
            logical_name=(f"provider-executions-{provider}-{execution_fingerprint}.json"),
            producer=f"preapproval-provider:{provider}",
            idempotency_key=f"{idempotency_key}:execution-manifest",
        )
        return _artifact_reference(descriptor)

    def execute(
        self,
        provider: str,
        request: ProviderMetadataExecutionInput,
        invocation: ToolInvocation,
    ) -> ProviderMetadataExecutionOutput:
        if invocation.workflow_id is None:
            raise ValueError("provider metadata retrieval requires a durable workflow ID")
        if request.ledger_record.provider != provider:
            raise ValueError("typed provider tool does not match its discovery ledger record")
        manifest = self.registry.release(provider, request.release_id)
        if request.access_mode is not manifest.default_access_mode and (
            request.access_mode not in manifest.optional_access_modes
        ):
            raise ValueError("requested provider access mode is not reviewed for this release")
        authenticated_api_available = self.executor.credential_boundary.available(
            "epa_comptox_api_key"
        )
        capability_findings = (
            manifest.public_capabilities
            if request.access_mode is ProviderAccessMode.PUBLIC_RELEASE
            else []
        )
        identifiers, identifier_reconciliation = self._load_identifiers(
            request,
            invocation.workflow_id,
            provider_batch_size=manifest.identifier_batch_size,
        )
        executions, ledger, cursor_exhausted, safety_truncated, metrics = self._retrieve(
            manifest=manifest,
            request=request,
            workflow_id=invocation.workflow_id,
            idempotency_key=invocation.idempotency_key or request.ledger_record.task_id,
            identifiers=identifiers,
        )
        cached_toxcast_normalization = (
            self._cached_toxcast_normalization()
            if provider == "toxcast"
            and request.access_mode is ProviderAccessMode.PUBLIC_RELEASE
            and self.public_provider_cache_root is not None
            else None
        )
        completed_resources = sum(int(outcome.completion_proof.completed) for outcome in executions)
        all_resources_complete = completed_resources == len(executions)
        processed_identifier_count = metrics.get("processed_identifier_count", len(identifiers))
        all_identifiers = (
            manifest.retrieval_mode is not ProviderRetrievalMode.IDENTIFIER_BATCHES
            or (
                processed_identifier_count == len(identifiers)
                and (
                    identifier_reconciliation is None
                    or identifier_reconciliation.invalid_identifier_count == 0
                )
            )
        )
        structural_codes = self._structural_validation_codes(manifest, executions, metrics)
        if (
            provider == "toxcast"
            and request.access_mode is ProviderAccessMode.PUBLIC_RELEASE
            and self.public_provider_cache_root is not None
            and cached_toxcast_normalization is None
        ):
            structural_codes = sorted(
                {*structural_codes, "PUBLIC_TOXCAST_ACTIVITY_CACHE_NOT_STAGED"}
            )
        completed = (
            all_resources_complete
            and (cursor_exhausted or bool(metrics.get("bounded_shortlist_satisfied", False)))
            and not safety_truncated
            and all_identifiers
            and not structural_codes
        )
        failure_classification = self._failure_classification(executions)
        if structural_codes:
            reason = "Provider structural validation failed: " + ", ".join(structural_codes) + "."
        elif (
            completed
            and request.access_mode is ProviderAccessMode.PUBLIC_RELEASE
            and capability_findings
        ):
            reason = (
                "Every reviewed public-release task completed. Public capability findings "
                "remain explicit and do not require an authenticated API key."
            )
        elif completed and request.validation_scope == "bounded_smoke":
            reason = (
                "The reviewed bounded validation scope completed; it does not claim full "
                "target discovery or select a source."
            )
        elif completed:
            if metrics.get("bounded_shortlist_satisfied") or metrics.get("excluded_candidate_ids"):
                reason = (
                    "The deterministic bounded assay shortlist completed every mandatory "
                    "structural inspection; any observed but unselected search hits remain "
                    "explicitly excluded and the provider search universe is not claimed "
                    "as exhaustively enumerated."
                )
            else:
                reason = (
                    "All reviewed resources/pages were exhausted and every upstream "
                    "identifier was processed."
                )
        else:
            reason = (
                "Provider task remains incomplete; resources, cursor, or full "
                "identifier coverage were not reconciled."
            )
        shortlist_decision_artifact: ArtifactReference | None = None
        excluded_candidate_ids = list(metrics.get("excluded_candidate_ids", []))
        if metrics.get("shortlisted_candidate_ids") or excluded_candidate_ids:
            shortlist_descriptor = self.artifacts.put_json(
                workflow_id=invocation.workflow_id,
                value={
                    "provider": provider,
                    "release_id": manifest.release_id,
                    "task_id": request.ledger_record.task_id,
                    "ranking_policy": (
                        "Retain the reviewed provider search order within the task's fair "
                        "allocation; do not infer activity from assay titles."
                    ),
                    "shortlisted_candidate_ids": metrics.get("shortlisted_candidate_ids", []),
                    "excluded_candidate_ids": excluded_candidate_ids,
                    "exclusion_reason": (
                        "Excluded by deterministic fair-allocation shortlist before "
                        "mandatory structural inspection."
                    ),
                },
                artifact_type="provider_candidate_shortlist",
                logical_name=(f"provider-candidate-shortlist-{request.ledger_record.task_id}.json"),
                producer=f"preapproval-provider:{provider}",
                idempotency_key=(
                    f"{invocation.idempotency_key or request.ledger_record.task_id}:shortlist"
                ),
            )
            shortlist_decision_artifact = _artifact_reference(shortlist_descriptor)
        proof = ProviderDatasetCompletionProof(
            provider=provider,
            release_id=manifest.release_id,
            task_id=request.ledger_record.task_id,
            declared_resource_count=len(executions),
            completed_resource_count=completed_resources,
            requested_identifier_count=request.upstream_identifier_count,
            processed_identifier_count=processed_identifier_count,
            cursor_exhausted=cursor_exhausted,
            file_manifest_reconciled=all_resources_complete,
            safety_truncated=safety_truncated,
            all_upstream_compounds_processed=all_identifiers,
            completed=completed,
            reason=reason,
            validation_scope=request.validation_scope,
            matched_aeid_count=metrics["matched_aeid_count"],
            summary_aeid_count=metrics["summary_aeid_count"],
            detail_aeid_count=metrics["detail_aeid_count"],
            discovered_candidate_count=metrics.get("discovered_candidate_count", 0),
            hydrated_candidate_count=metrics.get("hydrated_candidate_count", 0),
            shortlisted_candidate_ids=metrics.get("shortlisted_candidate_ids", []),
            excluded_candidate_ids=excluded_candidate_ids[:500],
            excluded_candidate_count=len(excluded_candidate_ids),
            candidate_exclusion_reason=(
                "Excluded by deterministic fair-allocation shortlist before hydration."
                if excluded_candidate_ids
                else None
            ),
            shortlist_decision_artifact=shortlist_decision_artifact,
            activity_availability_attempt_count=metrics.get(
                "activity_availability_attempt_count", 0
            ),
            relationship_attempt_count=metrics.get("relationship_attempt_count", 0),
            resolved_geo_accession_count=metrics.get("resolved_geo_accession_count", 0),
            geo_soft_attempt_count=metrics.get("geo_soft_attempt_count", 0),
            structural_validation_codes=structural_codes,
        )
        if identifier_reconciliation is not None:
            identifier_reconciliation = identifier_reconciliation.model_copy(
                update={
                    "completed_identifier_count": processed_identifier_count,
                    "missing_identifier_count": max(
                        identifier_reconciliation.normalized_unique_identifier_count
                        - processed_identifier_count,
                        0,
                    ),
                }
            )
        ledger = ledger.model_copy(
            update={
                "status": (
                    DiscoveryTaskStatus.COMPLETED if completed else DiscoveryTaskStatus.RUNNING
                ),
                "completion_reason": proof.reason,
                "error_classification": (
                    None
                    if completed
                    else structural_codes[0]
                    if structural_codes
                    else failure_classification or "provider_discovery_incomplete"
                ),
                "local_request_validation_failure_count": (
                    identifier_reconciliation.invalid_identifier_count
                    if identifier_reconciliation is not None
                    else 0
                ),
            }
        )
        execution_manifest = self._persist_execution_manifest(
            workflow_id=invocation.workflow_id,
            idempotency_key=invocation.idempotency_key or request.ledger_record.task_id,
            provider=provider,
            release_id=manifest.release_id,
            executions=executions,
        )
        ledger = ledger.model_copy(update={"source_response_artifacts": [execution_manifest]})
        partial_identity_normalization = provider == "pubchem-compound" and any(
            item.completion_proof.completed for item in executions
        )
        if not completed and not partial_identity_normalization:
            return ProviderMetadataExecutionOutput(
                provider=provider,
                release_id=manifest.release_id,
                ledger_record=ledger,
                execution_manifest_artifact=execution_manifest,
                exact_execution_count=len(executions),
                page_or_file_execution_preview=executions[:PREVIEW_LIMIT],
                execution_preview_is_complete=len(executions) <= PREVIEW_LIMIT,
                completion_proof=proof,
                cache_only_replay=bool(executions)
                and all(not outcome.scientific_source_request_count for outcome in executions),
                access_mode=request.access_mode,
                optional_authenticated_api_available=authenticated_api_available,
                capability_findings=capability_findings,
                pagination_yield=self._pagination_yield(manifest, executions),
                identifier_reconciliation=identifier_reconciliation,
            )
        normalized, manifest_reference, normalization_ran = self._normalize(
            manifest=manifest,
            request=request,
            workflow_id=invocation.workflow_id,
            idempotency_key=invocation.idempotency_key or request.ledger_record.task_id,
            executions=executions,
            proof=proof,
        )
        toxcast_public_activity: ToxCastPublicActivityCoverage | None = None
        toxcast_public_activity_artifact: ArtifactReference | None = None
        if cached_toxcast_normalization is not None:
            toxcast_public_activity, toxcast_public_activity_artifact = (
                self._persist_toxcast_public_activity_coverage(
                    manifest=manifest,
                    request=request,
                    workflow_id=invocation.workflow_id,
                    idempotency_key=(invocation.idempotency_key or request.ledger_record.task_id),
                    executions=executions,
                    identifiers=identifiers,
                    normalized=cached_toxcast_normalization,
                )
            )
        candidates, hydrated = self._previews(manifest, request, normalized)
        source_response_artifacts = [
            execution_manifest,
            manifest_reference,
            normalized.sqlite_artifact,
        ]
        if toxcast_public_activity_artifact is not None:
            source_response_artifacts.append(toxcast_public_activity_artifact)
        ledger = ledger.model_copy(
            update={
                "search_result_count": normalized.counts.source_candidates,
                "unique_candidate_count": normalized.counts.source_candidates,
                "source_response_artifacts": source_response_artifacts,
            }
        )
        return ProviderMetadataExecutionOutput(
            provider=provider,
            release_id=manifest.release_id,
            ledger_record=ledger,
            execution_manifest_artifact=execution_manifest,
            exact_execution_count=len(executions),
            page_or_file_execution_preview=executions[:PREVIEW_LIMIT],
            execution_preview_is_complete=len(executions) <= PREVIEW_LIMIT,
            completion_proof=proof,
            dataset_manifest=normalized,
            dataset_manifest_artifact=manifest_reference,
            candidates=candidates,
            hydrated_sources=hydrated,
            normalization_ran=normalization_ran,
            cache_only_replay=bool(executions)
            and all(not outcome.scientific_source_request_count for outcome in executions),
            access_mode=request.access_mode,
            optional_authenticated_api_available=authenticated_api_available,
            capability_findings=capability_findings,
            toxcast_public_activity=toxcast_public_activity,
            toxcast_public_activity_artifact=toxcast_public_activity_artifact,
            pagination_yield=self._pagination_yield(manifest, executions),
            identifier_reconciliation=(identifier_reconciliation),
        )


PROVIDER_TOOL_NAMES = {
    "toxcast": "retrieve_toxcast_preapproval_metadata",
    "tox21": "retrieve_tox21_preapproval_metadata",
    "pubchem-bioassay": "retrieve_pubchem_bioassay_preapproval_metadata",
    "pubchem-compound": "resolve_pubchem_compound_full_set",
    "ncbi-geo": "retrieve_geo_preapproval_metadata",
    "ncbi-supporting-metadata": "retrieve_supporting_metadata_catalogue",
}


class PreapprovalProviderLayer:
    """Expose distinct typed operations without allowing provider substitution."""

    def __init__(self, provider: PreapprovalMetadataProvider) -> None:
        self.provider = provider

    def execute(
        self,
        provider_id: str,
        request: ProviderMetadataExecutionInput,
        invocation: ToolInvocation,
    ) -> ProviderMetadataExecutionOutput:
        if provider_id not in PROVIDER_TOOL_NAMES:
            raise KeyError(provider_id)
        return self.provider.execute(provider_id, request, invocation)
